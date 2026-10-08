#!/usr/bin/env python3
"""Deterministic generator for the committed Harbourline HR-Policies files.

Writes 17 files to data/sharepoint/Harbourline-Hub/HR-Policies/ and 3 files to
HR-Policies/Restricted/. The 18th HR file (Employee-Handbook-Full.docx) is made
at setup time by generate_oversized_handbook.py.

Usage: python3 tools/generate-data/build_hr_policies.py [--root <repo root>]

Documents are described as plain Python data and rendered to .docx
(python-docx) or .pdf (reportlab). Output is deterministic: fixed document
properties, reportlab invariant mode, normalised zip timestamps for .docx, and a
fixed random seed for the scanned acknowledgement image.
"""

import datetime
import io
import os
import zipfile
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor, Cm
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import (ListFlowable, ListItem, Paragraph,
                                SimpleDocTemplate, Spacer, Table, TableStyle)

COMPANY = "Harbourline Energy Co."
FIXED_DT = datetime.datetime(2025, 6, 1, 9, 0, 0)
ZIP_DT = (2025, 6, 1, 9, 0, 0)
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
FORBIDDEN = (chr(0x2014), chr(0x2013))  # em dash, en dash


def check_text(s):
    for ch in FORBIDDEN:
        if ch in s:
            raise ValueError("Forbidden dash character in text: %r" % s[:80])
    return s


def all_strings(obj):
    """Yield every string inside a nested doc spec."""
    if isinstance(obj, str):
        yield obj
    elif isinstance(obj, dict):
        for v in obj.values():
            yield from all_strings(v)
    elif isinstance(obj, (list, tuple)):
        for v in obj:
            yield from all_strings(v)


def meta_rows(spec):
    rows = [
        ("Document ID", spec["doc_id"]),
        ("Version", spec["version"]),
        ("Effective date", spec["effective"]),
        ("Policy owner", spec["owner"]),
        ("Approved by", spec["approver"]),
        ("Applies to", spec.get("applies", "All Harbourline Energy Co. employees")),
        ("Next scheduled review", spec["next_review"]),
    ]
    if spec.get("classification"):
        rows.append(("Classification", spec["classification"]))
    return rows


# --------------------------------------------------------------------------
# DOCX
# --------------------------------------------------------------------------

def normalise_zip(path):
    """Rewrite a zip (docx) with fixed timestamps and sorted-stable order."""
    with open(path, "rb") as fh:
        data = fh.read()
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            zi = zipfile.ZipInfo(info.filename, date_time=ZIP_DT)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o644 << 16
            dst.writestr(zi, src.read(info.filename))
    with open(path, "wb") as fh:
        fh.write(out.getvalue())


def set_core(doc, spec):
    cp = doc.core_properties
    cp.author = spec.get("author", "Harbourline Human Resources")
    cp.last_modified_by = "Harbourline Human Resources"
    cp.title = spec["title"]
    cp.subject = spec.get("subject", "Human Resources policy")
    cp.keywords = spec.get("keywords", "Harbourline; HR; policy")
    cp.comments = ""
    cp.category = "Policy"
    cp.created = FIXED_DT
    cp.modified = FIXED_DT
    cp.last_printed = FIXED_DT
    cp.revision = 1


def _docx_table(doc, header, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    for i, h in enumerate(header):
        cell = t.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(check_text(h))
        run.bold = True
        run.font.size = Pt(9.5)
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(check_text(str(v)))
            run.font.size = Pt(9.5)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return t


def _docx_blocks(doc, blocks):
    for b in blocks:
        if isinstance(b, str):
            doc.add_paragraph(check_text(b))
        elif b[0] == "bullets":
            for item in b[1]:
                doc.add_paragraph(check_text(item), style="List Bullet")
        elif b[0] == "numbered":
            for n, item in enumerate(b[1], 1):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.8)
                p.add_run(check_text("%s) %s" % (chr(96 + n), item)))
        elif b[0] == "table":
            _docx_table(doc, b[1], b[2], b[3] if len(b) > 3 else None)
        elif b[0] == "sub":
            doc.add_heading(check_text(b[1]), level=2)
        elif b[0] == "note":
            p = doc.add_paragraph()
            r = p.add_run(check_text(b[1]))
            r.italic = True
        else:
            raise ValueError("unknown block %r" % (b[0],))


def render_docx(spec, path):
    for s in all_strings(spec):
        check_text(s)
    doc = Document()
    set_core(doc, spec)
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Cm(2.2)
        hp = sec.header.paragraphs[0]
        top = spec.get("header_banner")
        if top:
            r = hp.add_run(top + "    ")
            r.bold = True
            r.font.color.rgb = RGBColor(0xB0, 0x00, 0x00)
        r = hp.add_run("%s  |  %s  |  Version %s" % (COMPANY, spec["doc_id"], spec["version"]))
        r.font.size = Pt(8)
        fp = sec.footer.paragraphs[0]
        fr = fp.add_run(spec.get("footer", "Printed copies are uncontrolled. The current version is maintained in the Harbourline Hub HR-Policies library."))
        fr.font.size = Pt(8)
    if spec.get("header_banner"):
        p = doc.add_paragraph()
        r = p.add_run(spec["header_banner"])
        r.bold = True
        r.font.size = Pt(14)
        r.font.color.rgb = RGBColor(0xB0, 0x00, 0x00)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph()
    r = p.add_run(COMPANY)
    r.bold = True
    r.font.color.rgb = NAVY
    doc.add_heading(check_text(spec["title"]), level=0)
    if spec.get("subtitle"):
        doc.add_paragraph(check_text(spec["subtitle"])).runs[0].italic = True
    _docx_table(doc, ["Field", "Value"], meta_rows(spec), widths=[5, 11.5])
    for n, (heading, blocks) in enumerate(spec["sections"], 1):
        doc.add_heading("%d. %s" % (n, check_text(heading)), level=1)
        _docx_blocks(doc, blocks)
    doc.add_heading("%d. Revision History" % (len(spec["sections"]) + 1), level=1)
    _docx_table(doc, ["Version", "Date", "Author", "Summary of changes"], spec["revisions"],
                widths=[2, 2.8, 4, 7.7])
    os.makedirs(os.path.dirname(path), exist_ok=True)
    doc.save(path)
    normalise_zip(path)


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------

_ss = getSampleStyleSheet()
P_BODY = ParagraphStyle("body", parent=_ss["BodyText"], fontName="Helvetica", fontSize=10, leading=13.5,
                        spaceAfter=6, alignment=TA_LEFT)
P_H0 = ParagraphStyle("h0", parent=_ss["Title"], fontName="Helvetica-Bold", fontSize=19, leading=23,
                      textColor=colors.HexColor("#1F3A5F"), alignment=TA_LEFT, spaceAfter=4)
P_H1 = ParagraphStyle("h1", parent=_ss["Heading2"], fontName="Helvetica-Bold", fontSize=13, leading=16,
                      textColor=colors.HexColor("#1F3A5F"), spaceBefore=10, spaceAfter=5)
P_H2 = ParagraphStyle("h2", parent=_ss["Heading3"], fontName="Helvetica-Bold", fontSize=11, leading=14,
                      spaceBefore=6, spaceAfter=3)
P_CELL = ParagraphStyle("cell", parent=P_BODY, fontSize=8.8, leading=11, spaceAfter=0)
P_CELLB = ParagraphStyle("cellb", parent=P_CELL, fontName="Helvetica-Bold")
P_SMALL = ParagraphStyle("small", parent=P_BODY, fontSize=9, textColor=colors.HexColor("#444444"))


def _e(s):
    return escape(check_text(str(s)))


def _pdf_table(header, rows, widths=None):
    data = [[Paragraph(_e(h), P_CELLB) for h in header]]
    for r in rows:
        data.append([Paragraph(_e(v), P_CELL) for v in r])
    total = LETTER[0] - 4.4 * cm
    if widths:
        s = float(sum(widths))
        cw = [total * w / s for w in widths]
    else:
        cw = [total / len(header)] * len(header)
    t = Table(data, colWidths=cw, repeatRows=1)
    t.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#8A8A8A")),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DCE6F1")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    return [t, Spacer(1, 8)]


def _pdf_blocks(blocks):
    out = []
    for b in blocks:
        if isinstance(b, str):
            out.append(Paragraph(_e(b), P_BODY))
        elif b[0] == "bullets":
            out.append(ListFlowable([ListItem(Paragraph(_e(i), P_BODY), leftIndent=14) for i in b[1]],
                                    bulletType="bullet", start="•", leftIndent=14))
        elif b[0] == "numbered":
            out.append(ListFlowable([ListItem(Paragraph(_e(i), P_BODY), leftIndent=18) for i in b[1]],
                                    bulletType="a", bulletFormat="%s)", leftIndent=18))
        elif b[0] == "table":
            out.extend(_pdf_table(b[1], b[2], b[3] if len(b) > 3 else None))
        elif b[0] == "sub":
            out.append(Paragraph(_e(b[1]), P_H2))
        elif b[0] == "note":
            out.append(Paragraph("<i>%s</i>" % _e(b[1]), P_BODY))
        else:
            raise ValueError("unknown block %r" % (b[0],))
    return out


def render_pdf(spec, path):
    for s in all_strings(spec):
        check_text(s)
    os.makedirs(os.path.dirname(path), exist_ok=True)
    header_txt = "%s  |  %s  |  Version %s" % (COMPANY, spec["doc_id"], spec["version"])
    banner = spec.get("header_banner")
    footer_txt = spec.get("footer", "Printed copies are uncontrolled. The current version is maintained in the Harbourline Hub HR-Policies library.")

    def on_page(canv, d):
        canv.saveState()
        canv.setFont("Helvetica", 7.5)
        canv.setFillColor(colors.HexColor("#555555"))
        canv.drawString(2.2 * cm, LETTER[1] - 1.3 * cm, header_txt)
        if banner:
            canv.setFont("Helvetica-Bold", 8)
            canv.setFillColor(colors.HexColor("#B00000"))
            canv.drawRightString(LETTER[0] - 2.2 * cm, LETTER[1] - 1.3 * cm, banner)
            canv.setFillColor(colors.HexColor("#555555"))
            canv.setFont("Helvetica", 7.5)
        canv.drawString(2.2 * cm, 1.2 * cm, footer_txt)
        canv.drawRightString(LETTER[0] - 2.2 * cm, 1.2 * cm, "Page %d" % d.page)
        canv.restoreState()

    doc = SimpleDocTemplate(path, pagesize=LETTER, leftMargin=2.2 * cm, rightMargin=2.2 * cm,
                            topMargin=2.0 * cm, bottomMargin=2.0 * cm, title=spec["title"],
                            author=spec.get("author", "Harbourline Human Resources"),
                            subject=spec.get("subject", "Human Resources policy"),
                            creator="Harbourline Document Services", invariant=1)
    story = []
    if banner:
        story.append(Paragraph("<b>%s</b>" % _e(banner),
                               ParagraphStyle("ban", parent=P_BODY, textColor=colors.HexColor("#B00000"), fontSize=12)))
    story.append(Paragraph("<b>%s</b>" % _e(COMPANY), ParagraphStyle("co", parent=P_BODY, textColor=colors.HexColor("#1F3A5F"))))
    story.append(Paragraph(_e(spec["title"]), P_H0))
    if spec.get("subtitle"):
        story.append(Paragraph("<i>%s</i>" % _e(spec["subtitle"]), P_BODY))
    story.extend(_pdf_table(["Field", "Value"], meta_rows(spec), widths=[5, 11.5]))
    for n, (heading, blocks) in enumerate(spec["sections"], 1):
        story.append(Paragraph("%d. %s" % (n, _e(heading)), P_H1))
        story.extend(_pdf_blocks(blocks))
    story.append(Paragraph("%d. Revision History" % (len(spec["sections"]) + 1), P_H1))
    story.extend(_pdf_table(["Version", "Date", "Author", "Summary of changes"], spec["revisions"],
                            widths=[2, 2.8, 4, 7.7]))
    doc.build(story, onFirstPage=on_page, onLaterPages=on_page)


# ==========================================================================
# People used across the HR documents (fictional)
# ==========================================================================
VP_PEOPLE = "Margaret Osei, Vice President, People and Culture"
DIR_REWARDS = "Elena Marchetti, Director, Total Rewards"
DIR_ER = "Hamid Qureshi, Director, Employee Relations"
DIR_HS = "Colin Fraser, Director, Health and Safety"
US_HR = "Jennifer Alvarez, Manager, US Human Resources"
GC = "Andrew Kim, General Counsel"


# ==========================================================================
# 1 and 2. Leave policy, version 3 (docx) and version 4 (pdf)
# ==========================================================================

def leave_policy(v):
    if v == 3:
        tiers = [("Years 1 to 2 (0 to 24 months)", "15 days", "1.25 days"),
                 ("Years 3 to 7", "18 days", "1.50 days"),
                 ("Year 8 and later", "22 days", "1.83 days")]
        carry = 10
        tier_text = ("Regular full-time employees earn annual leave (vacation) according to completed years of "
                     "continuous service: 15 days per year during years 1 to 2, 18 days per year during years 3 "
                     "to 7, and 22 days per year from year 8 onward.")
    else:
        tiers = [("Years 1 to 4 (0 to 48 months)", "15 days", "1.25 days"),
                 ("Years 5 to 9", "20 days", "1.67 days"),
                 ("Year 10 and later", "25 days", "2.08 days")]
        carry = 5
        tier_text = ("Regular full-time employees earn annual leave (vacation) according to completed years of "
                     "continuous service: 15 days per year during years 1 to 4, 20 days per year during years 5 "
                     "to 9, and 25 days per year from year 10 onward.")
    spec = {
        "title": "Leave Policy",
        "subtitle": "Annual leave, carry-over, scheduling and payout",
        "doc_id": "HR-POL-012",
        "version": "%d.0" % v,
        "effective": "2024-01-01" if v == 3 else "2025-04-01",
        "owner": DIR_REWARDS,
        "approver": VP_PEOPLE,
        "next_review": "2025-12-31" if v == 3 else "2027-03-31",
        "applies": "All regular full-time and part-time employees in Canada and the United States",
        "keywords": "leave; vacation; annual leave; carry-over; HR-POL-012",
        "sections": [
            ("Purpose", [
                "This policy sets out how Harbourline Energy Co. (Harbourline) employees earn, schedule and use "
                "paid annual leave, how unused leave is carried over between years, and how leave is paid out "
                "when employment ends. Harbourline provides annual leave so that employees can rest and return "
                "to safety-sensitive utility work refreshed.",
            ]),
            ("Scope", [
                "This policy applies to all regular full-time and regular part-time employees of Harbourline in "
                "Ontario, New York and Ohio. Temporary and seasonal employees receive vacation pay as required by "
                "the applicable employment standards legislation and are otherwise excluded.",
                "Employees covered by a collective agreement receive the greater of the entitlement in this policy "
                "and the entitlement in their collective agreement. Where the collective agreement sets scheduling "
                "rules (for example, seniority-based vacation selection for line crews), those rules prevail.",
            ]),
            ("Definitions", [
                ("table", ["Term", "Meaning"], [
                    ("Annual leave", "Paid time away from work for rest and personal use, also called vacation."),
                    ("Leave year", "The calendar year, January 1 to December 31."),
                    ("Continuous service", "Service from the most recent date of hire, including approved leaves of absence."),
                    ("Carry-over", "Unused annual leave moved from one leave year into the next leave year."),
                    ("Day", "One regularly scheduled working day. For part-time staff, leave is prorated by scheduled hours."),
                ], [4, 12.5]),
            ]),
            ("Policy Statements", [
                ("sub", "%d.1 Annual leave entitlement" % 4),
                tier_text,
                ("table", ["Completed years of continuous service", "Annual leave per leave year", "Monthly accrual"],
                 tiers, [7, 5, 4.5]),
                "Leave accrues monthly from the date of hire. The entitlement tier changes on the first day of the "
                "month following the employee's service anniversary.",
                ("sub", "4.2 Carry-over of unused leave"),
                "Employees may carry over up to %d days of unused annual leave into the next leave year. Carried-over "
                "days must be used by March 31 of the following leave year. Any carried-over days not used by March 31 "
                "are forfeited unless forfeiture would breach employment standards legislation, in which case the "
                "minimum statutory amount is paid out." % carry,
                "Days above the %d-day carry-over limit at December 31 are forfeited. Managers are expected to review "
                "leave balances with their teams each September so that leave can be scheduled before year end." % carry,
                ("sub", "4.3 Statutory minimums"),
                "Nothing in this policy reduces an employee's entitlement under the Employment Standards Act, 2000 "
                "(Ontario) or any other applicable law. In New York and Ohio there is no statutory vacation "
                "entitlement, so this policy is the governing entitlement for US employees.",
                ("sub", "4.4 Payout on termination"),
                "On termination of employment for any reason, accrued and unused annual leave for the current leave "
                "year, plus any unexpired carried-over days, is paid out at the employee's base rate of pay on the "
                "final pay date.",
            ]),
            ("Procedures", [
                ("sub", "5.1 Requesting leave"),
                ("numbered", [
                    "Submit the request in the HR self-service portal (hrportal.harbourline.example).",
                    "Requests of more than 3 consecutive days must be submitted at least 10 business days in advance.",
                    "The manager approves or declines within 5 business days. A declined request must include a reason and an alternative period.",
                    "During a declared Storm Level 3 or higher emergency, operations managers may defer approved leave for "
                    "field and control-room staff. Deferred days are restored and are not counted against the carry-over limit.",
                ]),
                ("sub", "5.2 Year-end processing"),
                "Payroll applies the carry-over limit automatically on the first pay run in January. Employees can see "
                "carried-over days as a separate balance labelled Carry-over in the HR self-service portal.",
            ]),
            ("Roles and Responsibilities", [
                ("table", ["Role", "Responsibility"], [
                    ("Employee", "Plans and requests leave in advance; monitors own balance."),
                    ("People leader", "Approves leave fairly, balances operational coverage, reviews balances each September."),
                    ("Total Rewards", "Owns this policy, interprets entitlement questions, maintains accrual tables."),
                    ("Payroll", "Applies accruals, carry-over limits and termination payouts."),
                ], [4, 12.5]),
            ]),
            ("Related Documents", [
                ("bullets", [
                    "Parental Leave Policy (HR-POL-015)",
                    "Bereavement Leave Policy (HR-POL-017)",
                    "Overtime and On-Call Policy (HR-POL-038)",
                    "FMLA Guidance, United States (HR-POL-045)",
                    "Vacation Policy quick guide (Harbourline Hub, Getting-Started)",
                ]),
            ]),
        ],
    }
    revs = [
        ("1.0", "2019-03-01", "HR Policy Office", "Initial policy."),
        ("2.0", "2021-06-15", "HR Policy Office", "Added US employees in New York and Ohio; added payout rules."),
        ("3.0", "2024-01-01", "Elena Marchetti", "Revised entitlement tiers to 15, 18 and 22 days; carry-over limit set to 10 days, to be used by March 31."),
    ]
    if v == 4:
        revs.append(("4.0", "2025-04-01", "Elena Marchetti",
                     "Revised entitlement tiers to 15, 20 and 25 days (years 1 to 4, 5 to 9, 10 and later); "
                     "carry-over limit reduced to 5 days, to be used by March 31. Replaces version 3.0."))
    spec["revisions"] = revs
    return spec


# ==========================================================================
# 3. Code of Conduct (pdf)
# ==========================================================================

CODE_OF_CONDUCT = {
    "title": "Code of Business Conduct",
    "subtitle": "How we work safely, honestly and fairly",
    "doc_id": "HR-POL-001", "version": "5.1", "effective": "2025-01-15",
    "owner": GC, "approver": "Board of Directors, Harbourline Energy Co.", "next_review": "2026-01-15",
    "applies": "All employees, officers, directors and contractors of Harbourline Energy Co.",
    "sections": [
        ("Purpose", [
            "Harbourline delivers an essential public service. Customers, regulators such as the Ontario Energy "
            "Board and state public service commissions, and our communities rely on us to act with integrity. This "
            "Code sets the minimum standard of conduct expected from everyone who works for or on behalf of Harbourline.",
        ]),
        ("Scope", [
            "The Code applies to all employees, officers and directors in Canada and the United States, and to "
            "contractors and consultants while they perform work for Harbourline. Where local law sets a higher "
            "standard, the higher standard applies.",
        ]),
        ("Definitions", [
            ("table", ["Term", "Meaning"], [
                ("Conflict of interest", "A situation where a personal interest could influence, or appear to influence, a business decision."),
                ("Gift or hospitality", "Anything of value received or given, including meals, tickets, travel and discounts."),
                ("Related party", "A family member, household member or a business in which you or they hold a financial interest."),
                ("Retaliation", "Any adverse action taken against a person for raising a concern in good faith."),
            ], [4, 12.5]),
        ]),
        ("Standards of Conduct", [
            ("sub", "4.1 Safety first"),
            "No task is so urgent that it cannot be done safely. Every employee has the right and the duty to stop "
            "work that is unsafe. Working under the influence of alcohol, cannabis or impairing drugs is prohibited.",
            ("sub", "4.2 Conflicts of interest"),
            "Employees must disclose actual, potential or perceived conflicts of interest on the Conflict of Interest "
            "Disclosure Form (COI-01) within 10 business days of becoming aware of them. Employees may not approve "
            "purchases, contracts or hiring decisions involving a related party.",
            ("sub", "4.3 Gifts and hospitality"),
            "Gifts and hospitality valued above CAD 100 (Canada) or USD 75 (United States) must be declared in the "
            "Gifts and Hospitality Register within 5 business days of receipt. Cash, gift cards and securities may "
            "never be accepted. Employees involved in an active procurement may not accept any gift from a bidder.",
            ("sub", "4.4 Regulatory integrity"),
            "Information submitted to the Ontario Energy Board, NERC, FERC or a state public service commission must "
            "be complete and accurate. Employees must not share non-public transmission or market information with "
            "affiliates in breach of the Affiliate Relationships Code or FERC Standards of Conduct.",
            ("sub", "4.5 Respectful workplace"),
            "Harassment, discrimination and violence are prohibited. See the Workplace Harassment Prevention Policy "
            "(HR-POL-031) for complaint and investigation procedures.",
            ("sub", "4.6 Protecting information and assets"),
            "Company systems, vehicles, tools and data are for business use. Customer personal information is "
            "handled in line with the Privacy Policy and may only be accessed when needed for your role.",
        ]),
        ("Raising Concerns", [
            "Employees may raise concerns with their manager, Human Resources, the Legal department, or the "
            "independent Ethics Hotline, which is available 24 hours a day at 1-888-555-0142 or "
            "ethics.harbourline.example. Reports may be made anonymously.",
            "Harbourline does not tolerate retaliation. Anyone who retaliates against a person who raised a concern "
            "in good faith is subject to discipline up to and including termination.",
        ]),
        ("Annual Attestation", [
            "All employees must complete the annual Code of Conduct training and attestation in the learning system "
            "by February 28 each year. New employees complete it within 30 days of their start date. Signed paper "
            "acknowledgements are scanned and filed in the HR-Policies library.",
        ]),
        ("Roles and Responsibilities", [
            ("table", ["Role", "Responsibility"], [
                ("Every employee", "Follows the Code, completes annual attestation, raises concerns."),
                ("People leaders", "Model the Code, respond to concerns promptly, never retaliate."),
                ("Legal department", "Owns the Code, administers the Ethics Hotline, reports to the Audit Committee quarterly."),
                ("Human Resources", "Tracks completion of attestations and applies discipline consistently."),
            ], [4, 12.5]),
        ]),
        ("Related Documents", [
            ("bullets", ["Workplace Harassment Prevention Policy, Ontario (HR-POL-031)",
                         "Grievance Procedure (HR-POL-070)",
                         "Safety Incident Reporting Standard (HS-POL-005)",
                         "Travel Policy (HR-POL-050)"]),
        ]),
    ],
    "revisions": [
        ("4.0", "2021-01-01", "Legal department", "Full rewrite; added regulatory integrity section."),
        ("5.0", "2024-01-15", "Andrew Kim", "Added US gift threshold of USD 75; new Ethics Hotline provider."),
        ("5.1", "2025-01-15", "Andrew Kim", "Attestation deadline moved to February 28; clarified procurement gift ban."),
    ],
}


# ==========================================================================
# 4. Remote Work Policy (docx)
# ==========================================================================

REMOTE_WORK = {
    "title": "Remote and Hybrid Work Policy",
    "doc_id": "HR-POL-020", "version": "2.0", "effective": "2024-09-01",
    "owner": DIR_ER, "approver": VP_PEOPLE, "next_review": "2026-08-31",
    "applies": "Office-based employees in Canada and the United States",
    "sections": [
        ("Purpose", [
            "This policy describes when Harbourline employees may work away from a Harbourline site, the "
            "expectations that apply while they do, and the support Harbourline provides for a safe and secure "
            "home workspace.",
        ]),
        ("Scope", [
            "The policy applies to office-based roles at the Toronto head office, regional offices and the Albany "
            "and Columbus US offices. Field roles (line, substation, meter and vegetation crews), control-room "
            "operators and warehouse staff are not eligible for remote work because their duties require physical "
            "presence.",
        ]),
        ("Definitions", [
            ("table", ["Term", "Meaning"], [
                ("Hybrid role", "A role that works partly on site and partly remotely under an approved arrangement."),
                ("Fully remote role", "A role approved by a Vice President to work remotely with no regular on-site days."),
                ("Anchor days", "The on-site days set by each department so teams overlap in the office."),
                ("Temporary remote work", "Remote work from a location other than the approved home address."),
            ], [4, 12.5]),
        ]),
        ("Policy Statements", [
            ("sub", "4.1 Hybrid arrangements"),
            "Hybrid employees work on site a minimum of 3 days per week, including their department's anchor days. "
            "Arrangements are recorded in the HR self-service portal and reviewed each year during the performance cycle.",
            ("sub", "4.2 Work location"),
            "Employees must work from their approved home address within their province or state of employment. "
            "Temporary remote work from another province or state is limited to 20 working days per calendar year "
            "and requires manager and Human Resources approval in advance because of payroll tax and employment law "
            "implications. Remote work from outside Canada and the United States is not permitted.",
            ("sub", "4.3 Equipment and allowances"),
            ("table", ["Item", "Canada", "United States"], [
                ("One-time home office stipend", "CAD 500", "USD 400"),
                ("Monthly internet allowance", "CAD 40", "USD 30"),
                ("Laptop, monitor, headset", "Provided by IT", "Provided by IT"),
                ("Ergonomic assessment", "On request, virtual", "On request, virtual"),
            ], [6, 5, 5.5]),
            ("sub", "4.4 Information security"),
            "Company devices must connect through the Harbourline VPN when accessing internal systems. Operational "
            "technology systems, including SCADA and energy management systems, may never be accessed remotely except "
            "through the approved privileged access workstation. Printed customer information must not be taken home.",
            ("sub", "4.5 Health and safety"),
            "Employees are responsible for keeping their home workspace free of hazards. Work-related injuries at the "
            "home workspace must be reported under the Safety Incident Reporting Standard (HS-POL-005).",
        ]),
        ("Procedures", [
            ("numbered", [
                "Employee discusses the proposed arrangement with their manager.",
                "Employee completes the Remote Work Agreement and Home Workspace Checklist in the HR self-service portal.",
                "Manager approves; fully remote roles also need Vice President approval.",
                "IT ships equipment within 10 business days of approval.",
                "Either party may end a hybrid arrangement with 30 days notice.",
            ]),
        ]),
        ("Roles and Responsibilities", [
            ("table", ["Role", "Responsibility"], [
                ("Employee", "Maintains a safe, secure workspace and is reachable during core hours of 10:00 to 15:00 local time."),
                ("Manager", "Approves arrangements consistently and sets anchor days."),
                ("Employee Relations", "Owns the policy and resolves disputes."),
                ("IT Service Desk", "Supplies equipment and VPN access."),
            ], [4, 12.5]),
        ]),
        ("Related Documents", [
            ("bullets", ["Code of Business Conduct (HR-POL-001)", "Travel Policy (HR-POL-050)",
                         "Safety Incident Reporting Standard (HS-POL-005)", "Acceptable Use of IT Standard (IT-STD-004)"]),
        ]),
    ],
    "revisions": [
        ("1.0", "2020-06-01", "HR Policy Office", "Pandemic interim remote work guideline."),
        ("1.1", "2022-03-01", "Hamid Qureshi", "Converted guideline to policy; added US allowances."),
        ("2.0", "2024-09-01", "Hamid Qureshi", "Minimum on-site days raised from 2 to 3; 20-day limit for temporary remote work."),
    ],
}


# ==========================================================================
# 5. Harassment prevention, Ontario (docx)
# ==========================================================================

HARASSMENT = {
    "title": "Workplace Harassment and Violence Prevention Policy (Ontario)",
    "doc_id": "HR-POL-031", "version": "3.2", "effective": "2025-02-01",
    "owner": DIR_ER, "approver": VP_PEOPLE, "next_review": "2026-02-01",
    "applies": "All employees, contractors and visitors at Harbourline workplaces in Ontario",
    "sections": [
        ("Purpose", [
            "Harbourline is committed to a workplace free of harassment, sexual harassment and violence. This "
            "policy and its supporting program meet the employer obligations for workplace harassment and "
            "workplace violence under the Occupational Health and Safety Act (Ontario) (OHSA), including the "
            "requirement to have a written policy, a program to implement it, and to review both at least annually.",
        ]),
        ("Scope", [
            "The policy applies to conduct at all Harbourline workplaces in Ontario, including offices, "
            "substations, service centres, vehicles, customer premises where employees work, work-related "
            "social events, and online communication channels. US employees are covered by the equivalent US "
            "Respectful Workplace Policy.",
        ]),
        ("Definitions", [
            ("table", ["Term", "Meaning"], [
                ("Workplace harassment", "A course of vexatious comment or conduct against a worker that is known or ought reasonably to be known to be unwelcome, including workplace sexual harassment."),
                ("Workplace violence", "The exercise, attempt or threat of physical force against a worker that causes or could cause physical injury."),
                ("Reasonable management action", "Normal supervision such as assigning work, performance feedback and discipline. It is not harassment."),
                ("Workplace Harassment Officer", "The person designated to receive and oversee complaints: Denise Achterberg, Employee Relations."),
            ], [4.5, 12]),
        ]),
        ("Policy Statements", [
            ("bullets", [
                "Harassment and violence by or against any worker are prohibited.",
                "Complaints are investigated in a manner appropriate in the circumstances.",
                "Information about a complaint is kept confidential except where disclosure is needed to investigate, take corrective action, or is required by law.",
                "Reprisal against anyone who makes a complaint or participates in an investigation is prohibited.",
                "Workers may refuse work where workplace violence is likely to endanger them, following the OHSA work refusal process.",
            ]),
        ]),
        ("Procedures", [
            ("sub", "5.1 Reporting"),
            "Workers report incidents to their supervisor, to the Workplace Harassment Officer, or through the Ethics "
            "Hotline (1-888-555-0142). If the supervisor is the alleged harasser, the report goes directly to the "
            "Workplace Harassment Officer. Emergencies involving violence are reported to 911 first.",
            ("sub", "5.2 Investigation timelines"),
            ("table", ["Step", "Harbourline standard"], [
                ("Acknowledge complaint", "Within 2 business days"),
                ("Investigation starts", "Within 5 business days of the complaint"),
                ("Investigation completed", "Target of 90 calendar days"),
                ("Written results to complainant and respondent", "Within 10 business days of the investigation concluding"),
            ], [7, 9.5]),
            "Where the alleged harasser is a Director or above, an external investigator is appointed.",
            ("sub", "5.3 Training"),
            "All Ontario workers complete harassment and violence prevention training within 30 days of hire and a "
            "refresher every 2 years. People leaders complete an additional module on receiving complaints.",
            ("sub", "5.4 Risk assessment"),
            "A workplace violence risk assessment is performed for each work location at least every 3 years and "
            "whenever a site changes significantly. Customer-facing roles, including meter disconnection crews, "
            "receive de-escalation training.",
        ]),
        ("Roles and Responsibilities", [
            ("table", ["Role", "Responsibility"], [
                ("Workers", "Treat others with respect and report incidents."),
                ("Supervisors", "Act on reports immediately and protect confidentiality."),
                ("Workplace Harassment Officer", "Receives complaints, assigns investigators, tracks outcomes."),
                ("Joint Health and Safety Committee", "Is consulted on the policy and program during the annual review."),
            ], [5, 11.5]),
        ]),
        ("Related Documents", [
            ("bullets", ["Code of Business Conduct (HR-POL-001)", "Grievance Procedure (HR-POL-070)",
                         "Workplace Accommodation Policy (HR-POL-033)", "Workplace Harassment Program, Ontario (HR-PRG-031)"]),
        ]),
    ],
    "revisions": [
        ("3.0", "2023-02-01", "Hamid Qureshi", "Annual review; added virtual workplace channels to scope."),
        ("3.1", "2024-02-01", "Hamid Qureshi", "Annual review; no substantive change."),
        ("3.2", "2025-02-01", "Denise Achterberg", "Annual review; investigation start moved from 10 to 5 business days."),
    ],
}


# ==========================================================================
# 6. FMLA guidance, US (docx)
# ==========================================================================

FMLA = {
    "title": "Family and Medical Leave (FMLA) Guidance: United States",
    "doc_id": "HR-POL-045", "version": "1.4", "effective": "2024-07-01",
    "owner": US_HR, "approver": VP_PEOPLE, "next_review": "2026-06-30",
    "applies": "Employees of Harbourline Energy Co. in New York and Ohio",
    "sections": [
        ("Purpose", [
            "This guidance explains how Harbourline administers job-protected leave under the federal Family and "
            "Medical Leave Act (FMLA) and how FMLA leave interacts with state programs and company paid leave.",
        ]),
        ("Scope", [
            "The guidance applies to US employees in New York and Ohio. Canadian employees should refer to the "
            "Parental Leave Policy (HR-POL-015) and the Employment Standards Act, 2000 (Ontario).",
        ]),
        ("Definitions", [
            ("table", ["Term", "Meaning"], [
                ("Eligible employee", "Worked for Harbourline at least 12 months, worked at least 1,250 hours in the 12 months before leave, and works at a location with 50 or more employees within 75 miles."),
                ("Leave year", "Harbourline uses a rolling 12-month period measured backward from the date leave is used."),
                ("Serious health condition", "An illness, injury or condition involving inpatient care or continuing treatment by a health care provider."),
                ("Intermittent leave", "Leave taken in separate blocks of time or on a reduced schedule."),
            ], [4, 12.5]),
        ]),
        ("Entitlement", [
            ("bullets", [
                "Up to 12 workweeks of unpaid, job-protected leave in the leave year for the birth or placement of a child, "
                "the employee's own serious health condition, caring for a spouse, child or parent with a serious health "
                "condition, or a qualifying exigency arising from military service.",
                "Up to 26 workweeks in a single 12-month period to care for a covered servicemember (military caregiver leave).",
                "Group health coverage continues on the same terms during leave; the employee continues to pay their share.",
            ]),
            ("sub", "4.1 Coordination with paid leave"),
            "Harbourline requires accrued annual leave and paid sick time to run concurrently with FMLA leave, except "
            "during paid parental leave under HR-POL-015, which also runs concurrently with FMLA. Short-term disability "
            "benefits run concurrently with FMLA when the absence is for the employee's own condition.",
            ("sub", "4.2 State programs"),
            "New York: New York Paid Family Leave and New York disability benefits run concurrently with FMLA where the "
            "employee is eligible for both. Ohio: Ohio has no separate state family leave program for private employers, "
            "so FMLA and company policy apply.",
        ]),
        ("Procedures", [
            ("numbered", [
                "Give 30 days notice when leave is foreseeable, or notice as soon as practicable when it is not.",
                "Contact the US Leave Administration desk at leave-us@harbourline.example or 1-888-555-0177.",
                "Harbourline issues the eligibility notice and rights and responsibilities notice within 5 business days.",
                "Return medical certification within 15 calendar days of the request.",
                "Harbourline issues the designation notice within 5 business days of receiving complete certification.",
                "Before returning from leave for your own condition, provide a fitness-for-duty certification if your role is safety-sensitive.",
            ]),
        ]),
        ("Roles and Responsibilities", [
            ("table", ["Role", "Responsibility"], [
                ("Employee", "Gives notice, provides certification, reports changes in leave need."),
                ("Manager", "Refers leave requests to US Leave Administration within 2 business days; never asks for diagnosis."),
                ("US Leave Administration", "Determines eligibility, issues notices, tracks leave usage."),
                ("Benefits", "Maintains health coverage and collects employee premiums during leave."),
            ], [4.5, 12]),
        ]),
        ("Related Documents", [
            ("bullets", ["Leave Policy (HR-POL-012)", "Parental Leave Policy (HR-POL-015)",
                         "Workplace Accommodation Policy (HR-POL-033)", "US Department of Labor FMLA employee notice (posted at each US site)"]),
        ]),
    ],
    "revisions": [
        ("1.0", "2021-06-15", "US Human Resources", "Initial guidance for New York and Ohio employees."),
        ("1.3", "2023-01-01", "Jennifer Alvarez", "Added New York Paid Family Leave coordination."),
        ("1.4", "2024-07-01", "Jennifer Alvarez", "Switched leave year to rolling backward method; new leave desk contact."),
    ],
}


# ==========================================================================
# 7. Overtime and On-Call (docx)
# ==========================================================================

OVERTIME = {
    "title": "Overtime, On-Call and Storm Call-Out Policy",
    "doc_id": "HR-POL-038", "version": "4.0", "effective": "2025-01-01",
    "owner": DIR_REWARDS, "approver": VP_PEOPLE, "next_review": "2026-12-31",
    "applies": "Non-union employees in Canada and the United States; see section 2 for union employees",
    "sections": [
        ("Purpose", [
            "Reliable electricity service depends on employees who can respond outside normal hours, especially "
            "during storms. This policy sets the rules for overtime pay, on-call duty, storm call-outs and the rest "
            "periods that protect employee safety.",
        ]),
        ("Scope", [
            "This policy applies to non-union employees. Employees represented by the Harbourline Unit of Local "
            "4410 are governed by the overtime, call-out and standby provisions of their collective agreement, "
            "which prevail over this policy where they differ. The rest period rules in section 4.4 apply to all "
            "employees, union and non-union, because they are safety rules.",
        ]),
        ("Definitions", [
            ("table", ["Term", "Meaning"], [
                ("Overtime-eligible", "Non-exempt (US) or non-managerial hourly (Canada) employees."),
                ("On-call", "Scheduled availability outside normal hours to respond to system events."),
                ("Call-out", "Being called in to work outside scheduled hours."),
                ("Storm level", "Emergency level declared by the System Control Centre, from Level 1 (local) to Level 4 (system-wide)."),
            ], [4, 12.5]),
        ]),
        ("Policy Statements", [
            ("sub", "4.1 Overtime"),
            ("table", ["Jurisdiction", "Overtime threshold", "Rate"], [
                ("Ontario", "Hours worked over 44 in a work week", "1.5 times regular rate"),
                ("New York and Ohio", "Hours worked over 40 in a work week", "1.5 times regular rate"),
                ("All (company rule)", "Hours worked on a statutory or company holiday", "2 times regular rate"),
            ], [4.5, 7, 5]),
            "Overtime must be approved in advance by the supervisor except during emergency response.",
            ("sub", "4.2 On-call stipend"),
            "Employees scheduled on call receive a weekly stipend of CAD 300 (Canada) or USD 250 (United States) "
            "per full on-call week, prorated for partial weeks. On-call employees must acknowledge a page within 15 "
            "minutes and be on site within 60 minutes. Stipends are paid in addition to pay for hours actually worked.",
            ("sub", "4.3 Storm call-out"),
            "An employee called out after leaving work, whether or not on call, is paid a minimum of 4 hours at 2 "
            "times the regular rate, or actual hours worked at 2 times the regular rate if greater. When Storm Level 3 "
            "or higher is declared, all field employees may be required to report, and new vacation approvals are "
            "suspended until the storm level is lowered. A meal allowance of CAD 25 or USD 20 is paid for each "
            "5-hour block after the first 10 hours of continuous work.",
            ("sub", "4.4 Rest periods (all employees)"),
            "No employee may work more than 16 consecutive hours. After 16 hours of work, an employee must receive a "
            "rest period of at least 8 consecutive hours before returning to work. Rest time that falls in regular "
            "scheduled hours is paid at the regular rate. A supervisor may not waive this rule; only the Storm "
            "Incident Commander may approve an exception, in writing, for a documented public safety emergency.",
            ("sub", "4.5 Exempt and managerial employees"),
            "Exempt and managerial employees are not paid overtime. During Storm Level 3 or higher they receive an "
            "emergency response premium of CAD 60 or USD 50 per hour for hours worked beyond 50 in the week.",
        ]),
        ("Procedures", [
            ("numbered", [
                "Supervisors publish the on-call roster at least 14 days in advance in the scheduling system.",
                "Employees record overtime and call-out hours in the timekeeping system by the end of the pay period.",
                "Storm hours are coded with the storm event number issued by the System Control Centre.",
                "Payroll pays overtime, call-out and stipends on the next regular pay date.",
            ]),
        ]),
        ("Roles and Responsibilities", [
            ("table", ["Role", "Responsibility"], [
                ("Employee", "Records time accurately and observes rest periods."),
                ("Supervisor", "Approves overtime, maintains roster, enforces rest periods."),
                ("Storm Incident Commander", "Declares storm levels and approves rest-period exceptions."),
                ("Total Rewards", "Owns stipend and premium rates."),
            ], [4.5, 12]),
        ]),
        ("Related Documents", [
            ("bullets", ["Leave Policy (HR-POL-012)", "Storm Restoration Playbook (Harbourline-Operations, Procedures)",
                         "Collective Agreement, Harbourline Unit of Local 4410", "Safety Incident Reporting Standard (HS-POL-005)"]),
        ]),
    ],
    "revisions": [
        ("3.0", "2022-01-01", "Total Rewards", "Added US overtime thresholds."),
        ("3.1", "2023-07-01", "Elena Marchetti", "On-call stipend increased to CAD 275 and USD 225."),
        ("4.0", "2025-01-01", "Elena Marchetti", "Stipend increased to CAD 300 and USD 250; 16-hour limit and 8-hour rest rule apply to all employees."),
    ],
}


# ==========================================================================
# 8. Travel Policy (docx)
# ==========================================================================

TRAVEL = {
    "title": "Business Travel Policy",
    "doc_id": "HR-POL-050", "version": "3.0", "effective": "2025-03-01",
    "owner": "Nadia Rahman, Controller", "approver": "Rajiv Menon, Chief Financial Officer", "next_review": "2026-02-28",
    "sections": [
        ("Purpose", [
            "This policy sets out how employees plan, book and travel on Harbourline business so that travel is "
            "safe, cost effective and consistent with the expectations of our regulators, who review Harbourline's "
            "operating costs in rate proceedings. Reimbursement of travel costs is governed by the Expense Policy "
            "(HLE-FIN-101); where this policy and HLE-FIN-101 overlap, HLE-FIN-101 governs money matters.",
        ]),
        ("Scope", [
            "The policy applies to all employees travelling on Harbourline business. Storm mutual assistance "
            "deployments follow the Storm Restoration Playbook for crew lodging and use this policy and HLE-FIN-101 "
            "for reimbursement.",
        ]),
        ("Definitions", [
            ("table", ["Term", "Meaning"], [
                ("Travel Desk", "Harbourline's contracted travel agency, reachable at travel.harbourline.example."),
                ("Per diem", "A daily allowance covering meals and incidental costs. Rates are published in the Expense Policy (HLE-FIN-101)."),
                ("Travel day", "The day of departure or the day of return on a trip that includes an overnight stay."),
                ("Fleet vehicle", "A Harbourline-owned or leased pool vehicle booked through Fleet Services."),
            ], [4, 12.5]),
        ]),
        ("Policy Statements", [
            ("sub", "4.1 Booking"),
            "Book all air, rail and hotel travel through the Travel Desk at least 14 days before departure. "
            "Bookings made outside the Travel Desk are reimbursed only up to the Travel Desk fare.",
            ("sub", "4.2 Air travel"),
            "Economy class is the standard for all flights. Business class is permitted only for flights over 6 hours "
            "and only with Vice President approval obtained before booking.",
            ("sub", "4.3 Hotel caps (per night, before taxes)"),
            ("table", ["Location", "Nightly cap"], [
                ("Toronto", "CAD 275"),
                ("Other Ontario locations", "CAD 200"),
                ("New York City", "USD 350"),
                ("Other US locations (including Albany and Ohio)", "USD 225"),
            ], [9, 7.5]),
            "A hotel rate above the cap requires Director approval before booking.",
            ("sub", "4.4 Ground travel"),
            ("table", ["Item", "Rule"], [
                ("Personal vehicle mileage, Ontario", "CAD 0.72 per km"),
                ("Personal vehicle mileage, New York and Ohio", "USD 0.70 per mile"),
                ("Trips over 200 km (125 miles)", "Use a fleet vehicle when one is available"),
            ], [8, 8.5]),
            ("sub", "4.5 Meals and receipts"),
            "Meals and incidentals are covered by the per diem rates in HLE-FIN-101. On travel days, 75 percent of "
            "the per diem is paid. Receipts are required for any expense of CAD 25 or USD 25 or more. Alcohol is not "
            "reimbursable except at approved client or regulator events.",
            ("sub", "4.6 Traveller safety"),
            "Travellers register trips through the Travel Desk so that Harbourline can locate them in an emergency. "
            "Employees must not drive more than 10 hours in a day on business, and must follow the rest rules in "
            "the Overtime, On-Call and Storm Call-Out Policy (HR-POL-038) when travel follows a work shift.",
        ]),
        ("Procedures", [
            ("numbered", [
                "Obtain manager approval in the expense system before booking.",
                "Book through the Travel Desk at least 14 days ahead; request Director approval first if a hotel exceeds the cap.",
                "Book a fleet vehicle through Fleet Services for trips over 200 km (125 miles).",
                "Submit the expense claim with receipts under the Expense Policy (HLE-FIN-101) submission rules.",
            ]),
        ]),
        ("Roles and Responsibilities", [
            ("table", ["Role", "Responsibility"], [
                ("Traveller", "Books through the Travel Desk, travels safely and submits claims."),
                ("Manager", "Approves travel and reviews claims for business purpose."),
                ("Director", "Approves hotel rates above the cap before booking."),
                ("Vice President", "Approves business class for flights over 6 hours."),
                ("Finance, Accounts Payable", "Audits and reimburses claims under HLE-FIN-101."),
            ], [4.5, 12]),
        ]),
        ("Related Documents", [
            ("bullets", ["Expense Policy (HLE-FIN-101, Harbourline Hub, Finance)", "Corporate Card Guidelines (Harbourline Hub, Finance)",
                         "Code of Business Conduct (HR-POL-001)", "Remote and Hybrid Work Policy (HR-POL-020)"]),
        ]),
    ],
    "revisions": [
        ("2.0", "2022-09-01", "Finance", "Introduced Travel Desk."),
        ("2.1", "2024-03-01", "Nadia Rahman", "Updated hotel caps and mileage rates."),
        ("3.0", "2025-03-01", "Nadia Rahman", "Aligned with Expense Policy HLE-FIN-101; fleet vehicle rule for trips over 200 km; business class only over 6 hours."),
    ],
}


# ==========================================================================
# 9. Parental Leave (pdf)
# ==========================================================================

PARENTAL = {
    "title": "Parental Leave Policy",
    "doc_id": "HR-POL-015", "version": "2.1", "effective": "2024-05-01",
    "owner": DIR_REWARDS, "approver": VP_PEOPLE, "next_review": "2026-04-30",
    "sections": [
        ("Purpose", [
            "Harbourline supports employees welcoming a child through birth, adoption or surrogacy. This policy "
            "describes statutory leave, Harbourline's salary top-up and paid leave, and the return-to-work process.",
        ]),
        ("Scope", [
            "The policy applies to regular full-time and part-time employees in Canada and the United States. "
            "Union employees receive the greater of this policy and their collective agreement.",
        ]),
        ("Definitions", [
            ("table", ["Term", "Meaning"], [
                ("Pregnancy leave", "Leave for the birth parent under the Employment Standards Act, 2000 (Ontario)."),
                ("Parental leave", "Leave for any new parent following birth or placement of a child."),
                ("Top-up", "A Harbourline payment that adds to Employment Insurance (EI) benefits."),
                ("Eligible employee", "An employee with at least 6 months of continuous service on the leave start date."),
            ], [4, 12.5]),
        ]),
        ("Canada", [
            ("sub", "4.1 Statutory leave"),
            "Under the Employment Standards Act, 2000 (Ontario), a birth parent may take up to 17 weeks of "
            "pregnancy leave. Parental leave is up to 61 weeks for an employee who took pregnancy leave, and up to "
            "63 weeks for other new parents. Employees apply to Service Canada for EI maternity and parental benefits.",
            ("sub", "4.2 Harbourline top-up"),
            ("table", ["Leave", "Top-up level", "Maximum duration"], [
                ("Pregnancy leave", "93% of weekly base salary, including EI", "17 weeks"),
                ("Parental leave (any parent)", "93% of weekly base salary, including EI", "10 weeks"),
            ], [5, 7, 4.5]),
            "Employees who receive top-up must return to work for at least 6 months after leave, or repay the "
            "top-up on a prorated basis.",
        ]),
        ("United States", [
            "Eligible US employees receive up to 12 weeks of paid parental leave at 100% of base salary within 12 "
            "months of the birth or placement of a child. Paid parental leave runs concurrently with FMLA leave and "
            "with New York Paid Family Leave where applicable. See FMLA Guidance (HR-POL-045).",
        ]),
        ("Procedures", [
            ("numbered", [
                "Notify your manager and Human Resources at least 8 weeks before the planned leave start date, or as soon as possible.",
                "Complete the Leave of Absence request in the HR self-service portal.",
                "Canada: Payroll issues the Record of Employment within 5 calendar days of the leave start.",
                "Return to work: employees may use a phased return of up to 4 weeks at 60% of their normal schedule, paid at 100%.",
            ]),
        ]),
        ("Roles and Responsibilities", [
            ("table", ["Role", "Responsibility"], [
                ("Employee", "Gives notice and applies for government benefits."),
                ("Manager", "Plans coverage and supports return to work."),
                ("Total Rewards", "Administers top-up and paid leave; owns this policy."),
                ("Payroll", "Issues Records of Employment and processes top-up payments."),
            ], [4, 12.5]),
        ]),
        ("Related Documents", [
            ("bullets", ["Leave Policy (HR-POL-012)", "FMLA Guidance, United States (HR-POL-045)",
                         "Workplace Accommodation Policy (HR-POL-033)", "Benefits at a Glance (Harbourline Hub, Getting-Started)"]),
        ]),
    ],
    "revisions": [
        ("1.0", "2020-01-01", "Total Rewards", "Initial policy."),
        ("2.0", "2023-05-01", "Elena Marchetti", "Added US paid parental leave of 12 weeks."),
        ("2.1", "2024-05-01", "Elena Marchetti", "Parental top-up extended from 6 to 10 weeks; phased return added."),
    ],
}


# ==========================================================================
# 10. Bereavement Leave (docx)
# ==========================================================================

BEREAVEMENT = {
    "title": "Bereavement Leave Policy",
    "doc_id": "HR-POL-017", "version": "2.0", "effective": "2024-01-01",
    "owner": DIR_REWARDS, "approver": VP_PEOPLE, "next_review": "2025-12-31",
    "sections": [
        ("Purpose", [
            "Harbourline recognises that the death of a family member is one of the most difficult experiences an "
            "employee can face. This policy provides paid time away to grieve, attend services and manage "
            "family affairs.",
        ]),
        ("Scope", [
            "The policy applies to all regular full-time and part-time employees in Canada and the United States. "
            "Part-time employees receive paid days for days they were scheduled to work.",
        ]),
        ("Definitions", [
            ("table", ["Term", "Meaning"], [
                ("Immediate family", "Spouse or partner, child, stepchild, parent, stepparent, sibling, parent-in-law, and any person who lived in the employee's household."),
                ("Extended family", "Grandparent, grandchild, aunt, uncle, niece, nephew, sibling-in-law, son-in-law, daughter-in-law."),
            ], [4, 12.5]),
        ]),
        ("Policy Statements", [
            ("table", ["Relationship", "Paid days", "Additional unpaid days"], [
                ("Immediate family", "5 working days", "Up to 5"),
                ("Extended family", "2 working days", "Up to 3"),
                ("Travel over 500 km one way to attend services", "2 additional paid days", "Not applicable"),
            ], [6.5, 5, 5]),
            "Paid bereavement days may be taken within 3 months of the death, for example to attend a delayed "
            "service. They need not be consecutive.",
            "In Ontario, the Employment Standards Act, 2000 provides a statutory entitlement to unpaid bereavement "
            "leave. Harbourline's paid days count toward and exceed that entitlement.",
            "Pregnancy loss: an employee who experiences a pregnancy loss, or whose spouse or partner does, "
            "receives the immediate family entitlement.",
        ]),
        ("Procedures", [
            ("numbered", [
                "Tell your manager as soon as practicable. A phone call or text message is sufficient.",
                "Record the absence in the HR self-service portal using the Bereavement code within 10 days of returning.",
                "Documentation is not normally required. Human Resources may request it only where leave exceeds the policy.",
                "The Employee and Family Assistance Program (EFAP) offers grief counselling at 1-888-555-0133.",
            ]),
        ]),
        ("Roles and Responsibilities", [
            ("table", ["Role", "Responsibility"], [
                ("Employee", "Notifies manager and records the absence."),
                ("Manager", "Approves leave with compassion and arranges coverage."),
                ("Total Rewards", "Owns the policy and resolves eligibility questions."),
            ], [4, 12.5]),
        ]),
        ("Related Documents", [
            ("bullets", ["Leave Policy (HR-POL-012)", "Workplace Accommodation Policy (HR-POL-033)",
                         "Benefits at a Glance (Harbourline Hub, Getting-Started)"]),
        ]),
    ],
    "revisions": [
        ("1.0", "2018-01-01", "HR Policy Office", "Initial policy."),
        ("1.1", "2021-06-15", "HR Policy Office", "Extended to US employees."),
        ("2.0", "2024-01-01", "Elena Marchetti", "Immediate family paid days increased from 3 to 5; travel allowance added; pregnancy loss added."),
    ],
}


# ==========================================================================
# 11. Performance Review Process (docx)
# ==========================================================================

PERFORMANCE = {
    "title": "Performance Review Process",
    "doc_id": "HR-POL-060", "version": "3.0", "effective": "2025-01-01",
    "owner": "Samuel Idowu, Director, Talent Management", "approver": VP_PEOPLE, "next_review": "2026-12-31",
    "applies": "Non-union employees. Union employees follow the collective agreement.",
    "sections": [
        ("Purpose", [
            "The performance review process helps employees understand expectations, receive regular feedback, "
            "and be recognised fairly for their contribution to safe and reliable service.",
        ]),
        ("Scope", [
            "The process applies to all non-union employees with at least 3 months of service at the start of the "
            "review period. Employees hired after September 30 are reviewed in the following cycle.",
        ]),
        ("Definitions", [
            ("table", ["Term", "Meaning"], [
                ("TalentLine", "Harbourline's performance management system."),
                ("Calibration", "A meeting where leaders compare proposed ratings to ensure consistency."),
                ("PIP", "Performance Improvement Plan."),
            ], [4, 12.5]),
        ]),
        ("Annual Cycle", [
            ("table", ["Milestone", "Deadline"], [
                ("Goals set and approved in TalentLine", "January 31"),
                ("Mid-year check-in completed", "July 15"),
                ("Employee self-assessment due", "November 30"),
                ("Calibration sessions", "December 1 to December 20"),
                ("Final ratings communicated", "February 15 of the following year"),
                ("Merit increases effective", "April 1"),
            ], [9, 7.5]),
            "Every employee must have at least one safety goal. Goals follow the SMART format.",
        ]),
        ("Rating Scale and Merit", [
            ("table", ["Rating", "Label", "Merit increase guideline"], [
                ("5", "Exceptional", "4.0% to 5.0%"),
                ("4", "Exceeds expectations", "3.0% to 4.0%"),
                ("3", "Fully meets expectations", "2.0% to 3.0%"),
                ("2", "Partially meets expectations", "0% to 1.0%"),
                ("1", "Unsatisfactory", "0%"),
            ], [2.5, 6, 8]),
            "Merit guidelines are adjusted each year to fit the approved merit budget. Employees whose salary is "
            "above the maximum of their grade receive a lump sum instead of a base increase.",
        ]),
        ("Performance Improvement", [
            "An employee rated 1, or rated 2 in two consecutive years, is placed on a Performance Improvement Plan "
            "of 60 days. The PIP sets specific goals, support and check-in dates at least every 2 weeks. "
            "Human Resources must review a PIP before it is issued.",
        ]),
        ("Roles and Responsibilities", [
            ("table", ["Role", "Responsibility"], [
                ("Employee", "Sets goals, seeks feedback, completes self-assessment."),
                ("Manager", "Holds check-ins, writes review, proposes rating."),
                ("HR Business Partner", "Facilitates calibration and reviews PIPs."),
                ("Talent Management", "Owns the process and TalentLine configuration."),
            ], [4.5, 12]),
        ]),
        ("Related Documents", [
            ("bullets", ["Compensation Bands 2025 (HR confidential)", "Grievance Procedure (HR-POL-070)",
                         "Code of Business Conduct (HR-POL-001)"]),
        ]),
    ],
    "revisions": [
        ("2.0", "2021-01-01", "Talent Management", "Moved to TalentLine."),
        ("2.1", "2023-01-01", "Samuel Idowu", "Mandatory safety goal added."),
        ("3.0", "2025-01-01", "Samuel Idowu", "Five-point scale replaces four-point scale; PIP set at 60 days."),
    ],
}


# ==========================================================================
# 12. Safety Incident Reporting (pdf)
# ==========================================================================

SAFETY = {
    "title": "Safety Incident Reporting Standard",
    "doc_id": "HS-POL-005", "version": "6.0", "effective": "2025-06-01",
    "owner": DIR_HS, "approver": "Dana Okafor, Chief Operating Officer", "next_review": "2026-05-31",
    "applies": "All employees and contractors at Harbourline workplaces",
    "sections": [
        ("Purpose", [
            "Prompt and accurate incident reporting lets Harbourline care for injured people, meet legal "
            "notification duties, and learn from events before they cause serious harm. This standard defines "
            "what must be reported, when, and to whom.",
        ]),
        ("Scope", [
            "The standard covers injuries, illnesses, near misses, electrical contacts and flashes, vehicle "
            "incidents, dangerous occurrences and property damage, in Ontario, New York and Ohio.",
        ]),
        ("Definitions", [
            ("table", ["Term", "Meaning"], [
                ("Near miss", "An event that could have caused injury or damage but did not."),
                ("Serious incident", "A fatality, critical injury, hospital admission, electrical contact, arc flash, or a fire or explosion."),
                ("SafeLine", "Harbourline's incident management system."),
                ("Safety Duty Officer", "The on-call safety professional, available 24 hours a day."),
            ], [4, 12.5]),
        ]),
        ("Reporting Timelines", [
            ("table", ["Event", "Internal requirement"], [
                ("Any incident", "Report to supervisor immediately; enter in SafeLine within 24 hours"),
                ("Serious incident", "Notify Safety Duty Officer within 1 hour at 1-888-555-0199"),
                ("Near miss", "Enter in SafeLine within 48 hours"),
                ("Vehicle incident", "Report to supervisor and Fleet within 24 hours"),
            ], [5, 11.5]),
            ("sub", "4.1 Ontario legal notifications"),
            "For a critical injury or fatality, the Ministry of Labour, Immigration, Training and Skills Development "
            "must be notified immediately as required by the Occupational Health and Safety Act, and the scene must "
            "not be disturbed except to save life or prevent further injury. Lost-time or health care injuries are "
            "reported to the WSIB (Form 7) within 3 days of learning of them. The Safety Duty Officer makes these "
            "notifications; supervisors must not contact the regulator directly.",
            ("sub", "4.2 United States legal notifications"),
            "OSHA must be notified of a work-related fatality within 8 hours, and of an in-patient hospitalization, "
            "amputation or loss of an eye within 24 hours. The Safety Duty Officer makes these notifications.",
        ]),
        ("Investigation", [
            "Serious incidents are investigated by a team led by a Health and Safety advisor, with a Joint Health and "
            "Safety Committee worker member in Ontario. The root cause analysis is completed within 10 business days, "
            "and corrective actions are tracked in SafeLine to closure. A Safety Alert is issued company-wide within "
            "48 hours of any electrical contact.",
        ]),
        ("Roles and Responsibilities", [
            ("table", ["Role", "Responsibility"], [
                ("Worker", "Stops work, gets help, reports immediately, preserves the scene."),
                ("Supervisor", "Ensures care, enters SafeLine report, starts initial review."),
                ("Safety Duty Officer", "Makes regulatory notifications and mobilises investigators."),
                ("Director, Health and Safety", "Owns this standard and reports statistics monthly to the executive team."),
            ], [5, 11.5]),
        ]),
        ("Related Documents", [
            ("bullets", ["Lockout/Tagout Safety Manual (Harbourline-Operations, Procedures)",
                         "Outage Response Procedure (Harbourline-Operations, Procedures)",
                         "Workplace Harassment and Violence Prevention Policy, Ontario (HR-POL-031)",
                         "Overtime, On-Call and Storm Call-Out Policy (HR-POL-038)"]),
        ]),
    ],
    "revisions": [
        ("5.0", "2022-06-01", "Health and Safety", "Introduced SafeLine."),
        ("5.1", "2024-01-15", "Colin Fraser", "Added near miss 48-hour rule."),
        ("6.0", "2025-06-01", "Colin Fraser", "Serious incident escalation shortened from 2 hours to 1 hour; Safety Alert rule added."),
    ],
}


# ==========================================================================
# 13. Workplace Accommodation (docx)
# ==========================================================================

ACCOMMODATION = {
    "title": "Workplace Accommodation Policy",
    "doc_id": "HR-POL-033", "version": "2.0", "effective": "2024-10-01",
    "owner": DIR_ER, "approver": VP_PEOPLE, "next_review": "2026-09-30",
    "sections": [
        ("Purpose", [
            "Harbourline accommodates employees and applicants to the point of undue hardship, in line with the "
            "Ontario Human Rights Code, the Accessibility for Ontarians with Disabilities Act (AODA), the Americans "
            "with Disabilities Act (ADA), and New York and Ohio law.",
        ]),
        ("Scope", [
            "Accommodation is available for disability, religion or creed, family status, pregnancy and other "
            "protected grounds, for employees and job applicants in Canada and the United States.",
        ]),
        ("Definitions", [
            ("table", ["Term", "Meaning"], [
                ("Accommodation", "A change to the job, workplace or schedule that removes a barrier related to a protected ground."),
                ("Undue hardship", "Cost or health and safety risk so significant that accommodation cannot reasonably be made, as defined by applicable law."),
                ("Interactive process", "Good-faith collaboration between employee, manager and HR to identify options."),
                ("Individual accommodation plan", "A written plan that records agreed accommodations and review dates."),
            ], [4.5, 12]),
        ]),
        ("Policy Statements", [
            ("bullets", [
                "Requests may be made verbally or in writing; using Form ACC-1 in the HR self-service portal is preferred but not required.",
                "Medical information is held by Health Services, not by the manager. Managers receive only functional limitations.",
                "Safety-sensitive roles (for example, lineworkers and system operators) require a fitness-for-duty assessment where the request affects safety-critical tasks.",
                "Accommodation costs are paid from the central Accommodation Fund. Requests up to CAD 5,000 (or USD 3,700) are approved by the HR Business Partner without Vice President approval.",
            ]),
        ]),
        ("Procedures", [
            ("table", ["Step", "Timeline"], [
                ("Request received and acknowledged", "Within 2 business days"),
                ("Initial interactive process meeting", "Within 10 business days of request"),
                ("Decision communicated in writing", "Within 30 calendar days of request"),
                ("Individual accommodation plan reviewed", "Every 12 months, or sooner if needs change"),
            ], [9, 7.5]),
            "If an accommodation is denied on the basis of undue hardship, the decision must be approved by the "
            "Director, Employee Relations, and the employee may use the Grievance Procedure.",
        ]),
        ("Roles and Responsibilities", [
            ("table", ["Role", "Responsibility"], [
                ("Employee", "Makes the need known and participates in the interactive process."),
                ("Manager", "Refers requests to HR, implements the plan, protects confidentiality."),
                ("HR Business Partner", "Leads the interactive process and approves costs within limit."),
                ("Health Services", "Holds medical information and advises on functional abilities."),
            ], [4.5, 12]),
        ]),
        ("Related Documents", [
            ("bullets", ["Workplace Harassment and Violence Prevention Policy, Ontario (HR-POL-031)",
                         "FMLA Guidance, United States (HR-POL-045)", "Grievance Procedure (HR-POL-070)",
                         "Multi-Year Accessibility Plan (Harbourline Hub)"]),
        ]),
    ],
    "revisions": [
        ("1.0", "2019-01-01", "HR Policy Office", "Initial policy."),
        ("1.2", "2022-10-01", "Hamid Qureshi", "Added US ADA references."),
        ("2.0", "2024-10-01", "Hamid Qureshi", "Central Accommodation Fund created; HRBP approval limit CAD 5,000."),
    ],
}


# ==========================================================================
# 14. Grievance Procedure (docx)
# ==========================================================================

GRIEVANCE = {
    "title": "Grievance Procedure (Non-Union Employees)",
    "doc_id": "HR-POL-070", "version": "2.3", "effective": "2024-06-01",
    "owner": DIR_ER, "approver": VP_PEOPLE, "next_review": "2026-05-31",
    "applies": "Non-union employees in Canada and the United States",
    "sections": [
        ("Purpose", [
            "This procedure gives non-union employees a fair, timely and confidential way to raise and resolve "
            "workplace complaints about the application of Harbourline policies, working conditions or treatment.",
        ]),
        ("Scope", [
            "The procedure applies to non-union employees. Employees represented by the Harbourline Unit of Local "
            "4410 use the grievance and arbitration process in their collective agreement. Complaints of harassment "
            "or discrimination follow HR-POL-031; concerns about fraud or ethics may go to the Ethics Hotline.",
        ]),
        ("Definitions", [
            ("table", ["Term", "Meaning"], [
                ("Grievance", "A written complaint that a policy was misapplied or that the employee was treated unfairly."),
                ("Business day", "Monday to Friday, excluding statutory and company holidays."),
                ("Grievance Review Panel", "The Vice President, People and Culture, plus two Directors not involved in the matter."),
            ], [4.5, 12]),
        ]),
        ("Procedure", [
            ("table", ["Step", "Who", "Employee deadline", "Response deadline"], [
                ("Informal discussion", "Immediate manager", "Encouraged first", "Not applicable"),
                ("Step 1", "Immediate manager", "Within 15 business days of the event", "10 business days"),
                ("Step 2", "HR Business Partner and second-level manager", "Within 10 business days of Step 1 response", "10 business days"),
                ("Step 3", "Grievance Review Panel", "Within 10 business days of Step 2 response", "20 business days; decision is final"),
            ], [3, 5, 4.5, 4]),
            "Employees may be accompanied by a co-worker at Step 2 and Step 3 meetings. Legal counsel does not "
            "attend internal grievance meetings.",
            "Timelines may be extended by mutual written agreement. If Harbourline misses a response deadline, "
            "the employee may move to the next step.",
        ]),
        ("Protections", [
            ("bullets", [
                "No reprisal may be taken against an employee for filing a grievance in good faith.",
                "Grievance records are kept by Employee Relations separately from the personnel file for 7 years.",
                "Use of this procedure does not remove any right under employment standards or human rights legislation.",
            ]),
        ]),
        ("Roles and Responsibilities", [
            ("table", ["Role", "Responsibility"], [
                ("Employee", "Submits grievances in writing using Form GRV-1 within deadlines."),
                ("Manager", "Responds in writing within deadlines."),
                ("HR Business Partner", "Facilitates Step 2 and advises managers."),
                ("Employee Relations", "Owns the procedure and supports the Review Panel."),
            ], [4.5, 12]),
        ]),
        ("Related Documents", [
            ("bullets", ["Code of Business Conduct (HR-POL-001)", "Workplace Harassment and Violence Prevention Policy, Ontario (HR-POL-031)",
                         "Employee Relations Contacts (HR-Policies library)"]),
        ]),
    ],
    "revisions": [
        ("2.0", "2020-06-01", "Employee Relations", "Three-step process introduced."),
        ("2.2", "2023-06-01", "Hamid Qureshi", "Added co-worker accompaniment."),
        ("2.3", "2024-06-01", "Hamid Qureshi", "Step 1 filing window extended from 10 to 15 business days."),
    ],
}


# ==========================================================================
# 16. Compensation Bands 2025 (docx, labeled at setup)
# ==========================================================================

def comp_bands():
    """Return rows: grade, CAD min/mid/max, USD min/mid/max. Deterministic."""
    rows = []
    mid_cad = 48000
    for g in range(1, 13):
        mn, mx = round(mid_cad * 0.8 / 100) * 100, round(mid_cad * 1.2 / 100) * 100
        mid_usd = round(mid_cad * 0.78 / 100) * 100
        umn, umx = round(mid_usd * 0.8 / 100) * 100, round(mid_usd * 1.2 / 100) * 100
        rows.append(("G%d" % g, mn, mid_cad, mx, umn, mid_usd, umx))
        mid_cad = round(mid_cad * 1.14 / 500) * 500
    return rows


def money(n):
    return "{:,}".format(n)


def comp_spec():
    rows = [(g, money(a), money(b), money(c), money(d), money(e), money(f)) for g, a, b, c, d, e, f in comp_bands()]
    return {
        "title": "Compensation Bands 2025",
        "subtitle": "Salary structure for non-union employees, grades G1 to G12",
        "doc_id": "HR-CMP-2025", "version": "1.0", "effective": "2025-04-01",
        "owner": DIR_REWARDS, "approver": VP_PEOPLE, "next_review": "2026-01-31",
        "applies": "Non-union employees in Canada and the United States",
        "classification": "Confidential, HR Only",
        "header_banner": "CONFIDENTIAL - HR Only",
        "footer": "CONFIDENTIAL - HR Only. Do not forward, print or store outside the HR-Policies library.",
        "keywords": "compensation; salary bands; confidential",
        "sections": [
            ("Purpose", [
                "This document publishes the 2025 salary structure used to set and review base salaries for "
                "non-union roles. It is restricted to the Human Resources team and is protected by the "
                "Confidential, HR Only sensitivity label.",
            ]),
            ("Structure Adjustment", [
                "The 2025 structure moved all band midpoints up by 3.2% compared with 2024, based on the market "
                "survey of 22 North American utilities completed in November 2024. The midpoint progression between "
                "grades is approximately 14%, and each band spans 80% to 120% of its midpoint.",
                "The 2025 merit budget for non-union employees is 3.0% of base payroll, with a separate 0.5% pool "
                "for market adjustments of employees below 90% of midpoint.",
            ]),
            ("Salary Bands by Grade", [
                ("table", ["Grade", "CAD minimum", "CAD midpoint", "CAD maximum", "USD minimum", "USD midpoint", "USD maximum"],
                 rows, [1.6, 2.5, 2.5, 2.5, 2.5, 2.5, 2.5]),
                ("note", "USD bands apply to New York and Ohio employees and are set at 78% of the CAD midpoint, "
                         "rounded to the nearest 100, reflecting local market data rather than currency conversion."),
            ]),
            ("Typical Roles by Grade", [
                ("table", ["Grade", "Example roles"], [
                    ("G1 to G2", "Customer service representative, clerk, meter reader"),
                    ("G3 to G4", "Senior clerk, technician trainee, HR coordinator"),
                    ("G5 to G6", "Analyst, planner, GIS specialist, accountant"),
                    ("G7 to G8", "Senior engineer, senior analyst, HR business partner, supervisor"),
                    ("G9 to G10", "Manager, principal engineer, senior manager"),
                    ("G11 to G12", "Director, senior director (executive roles are set by the Board)"),
                ], [3, 13.5]),
            ]),
            ("Rules for Use", [
                ("bullets", [
                    "New hires are normally placed between the band minimum and the midpoint.",
                    "Offers above the midpoint require Director, Total Rewards approval.",
                    "Offers above the maximum are not permitted.",
                    "Promotional increases are normally 8% to 12%, or to the new band minimum if greater.",
                ]),
            ]),
        ],
        "revisions": [
            ("1.0", "2025-04-01", "Elena Marchetti", "2025 structure; midpoints increased 3.2%."),
        ],
    }


# ==========================================================================
# 17. Employee Relations Contacts (docx)
# ==========================================================================

ER_CONTACTS = {
    "title": "Employee Relations and HR Business Partner Contacts",
    "doc_id": "HR-REF-002", "version": "2025.3", "effective": "2025-07-01",
    "owner": DIR_ER, "approver": VP_PEOPLE, "next_review": "2025-10-01",
    "sections": [
        ("Purpose", [
            "This directory lists the Human Resources Business Partner (HRBP) for each region, and the central "
            "Employee Relations contacts. Employees should contact their regional HRBP first for questions about "
            "policies, accommodation, performance or workplace concerns.",
        ]),
        ("Regional HR Business Partners", [
            ("table", ["Region", "Office", "HR Business Partner", "Email", "Phone"], [
                ("Toronto Head Office", "Toronto, ON", "Priya Nandakumar (HR Manager)", "priya.nandakumar@harbourline.example", "416-555-0110"),
                ("Ontario East", "Kingston, ON", "Luc Tremblay", "luc.tremblay@harbourline.example", "613-555-0122"),
                ("Ontario North", "Sudbury, ON", "Aisha Mohammed", "aisha.mohammed@harbourline.example", "705-555-0134"),
                ("Ontario West", "London, ON", "Graham Wilson", "graham.wilson@harbourline.example", "519-555-0146"),
                ("New York", "Albany, NY", "Maria Santos", "maria.santos@harbourline.example", "518-555-0158"),
                ("Ohio", "Columbus, OH", "Derek Johnson", "derek.johnson@harbourline.example", "614-555-0161"),
            ], [3, 2.6, 3.6, 4.6, 2.7]),
        ]),
        ("Central Contacts", [
            ("table", ["Service", "Contact", "Details"], [
                ("Employee Relations inbox", "er@harbourline.example", "Monitored business days 08:00 to 17:00 Eastern; response within 1 business day"),
                ("Workplace Harassment Officer", "Denise Achterberg", "denise.achterberg@harbourline.example, 416-555-0175"),
                ("Director, Employee Relations", "Hamid Qureshi", "hamid.qureshi@harbourline.example"),
                ("US Leave Administration", "leave-us@harbourline.example", "1-888-555-0177"),
                ("Ethics Hotline (independent, anonymous)", "1-888-555-0142", "ethics.harbourline.example, 24 hours a day"),
                ("Employee and Family Assistance Program", "1-888-555-0133", "24 hours a day, confidential"),
                ("HR self-service portal", "hrportal.harbourline.example", "Leave, forms, pay statements"),
            ], [4.5, 4.5, 7.5]),
        ]),
        ("Escalation", [
            "If your regional HRBP is unavailable for more than 2 business days, contact the Employee Relations "
            "inbox. If your concern involves your HRBP, contact the Director, Employee Relations directly.",
            "During a declared Storm Level 3 or higher, HRBPs for Ontario East, North and West rotate an after-hours "
            "HR duty line at 1-888-555-0188 for crew welfare and fatigue issues.",
        ]),
        ("Related Documents", [
            ("bullets", ["Grievance Procedure (HR-POL-070)", "Workplace Harassment and Violence Prevention Policy, Ontario (HR-POL-031)",
                         "Workplace Accommodation Policy (HR-POL-033)"]),
        ]),
    ],
    "revisions": [
        ("2025.1", "2025-01-06", "Employee Relations", "Quarterly update."),
        ("2025.2", "2025-04-01", "Employee Relations", "Ohio HRBP changed to Derek Johnson."),
        ("2025.3", "2025-07-01", "Employee Relations", "Added storm HR duty line."),
    ],
}


# ==========================================================================
# Restricted documents
# ==========================================================================

DISCIPLINARY = {
    "title": "Disciplinary Case Handling Procedure",
    "doc_id": "HR-RST-101", "version": "1.6", "effective": "2025-03-01",
    "owner": DIR_ER, "approver": VP_PEOPLE, "next_review": "2026-02-28",
    "applies": "Human Resources staff and people leaders handling discipline, with HR guidance",
    "classification": "Restricted, HR group only",
    "header_banner": "RESTRICTED - HR Only",
    "sections": [
        ("Purpose", [
            "This procedure tells HR staff how to open, manage, document and close disciplinary cases so that "
            "discipline is fair, consistent and defensible. It is an internal HR procedure and is not published "
            "to employees.",
        ]),
        ("Scope", [
            "Applies to discipline of non-union employees. For union employees, HR follows the collective agreement "
            "and notifies the union steward before any disciplinary meeting.",
        ]),
        ("Case Management", [
            "Every case is opened in the Employee Relations case system with a case number in the format "
            "DC-YYYY-NNN (for example, DC-2025-031). Case files are retained for 7 years after closure.",
            "Before any disciplinary meeting, the HR Business Partner confirms: the facts are documented, the "
            "employee has had a chance to respond, and prior comparable cases have been reviewed for consistency "
            "using the Consistency Log.",
        ]),
        ("Progressive Discipline and Retention", [
            ("table", ["Level", "Approval required", "Active period on file"], [
                ("Documented verbal warning", "Manager with HRBP", "12 months"),
                ("Written warning", "Manager with HRBP", "18 months"),
                ("Final written warning", "Director and HRBP", "24 months"),
                ("Unpaid suspension (1 to 5 days)", "Director, Employee Relations", "24 months"),
                ("Termination for cause", "Vice President, People and Culture, after Legal review", "Permanent"),
            ], [5, 6.5, 5]),
            "Code Red safety violations, including a breach of lockout/tagout, working without required arc flash "
            "PPE, or operating a vehicle while impaired, may proceed directly to a final written warning or "
            "termination without earlier steps.",
            "A warning that has passed its active period may not be relied on for progressive discipline but stays in "
            "the case file.",
        ]),
        ("Meetings", [
            ("bullets", [
                "Give the employee at least 24 hours notice of a disciplinary meeting, except for Code Red matters.",
                "An HR representative attends every meeting above verbal warning.",
                "Employees may bring a co-worker as a support person.",
                "The outcome letter is issued within 3 business days of the meeting using template ER-L4.",
            ]),
        ]),
        ("Related Documents", [
            ("bullets", ["Investigation Protocol (HR-RST-103)", "Grievance Procedure (HR-POL-070)",
                         "Code of Business Conduct (HR-POL-001)"]),
        ]),
    ],
    "revisions": [
        ("1.4", "2023-03-01", "Hamid Qureshi", "Consistency Log introduced."),
        ("1.5", "2024-03-01", "Hamid Qureshi", "Code Red category added."),
        ("1.6", "2025-03-01", "Hamid Qureshi", "Final written warning active period extended from 18 to 24 months."),
    ],
}

EXEC_COMP = {
    "title": "Executive Compensation Review, Fiscal 2025",
    "doc_id": "HR-RST-102", "version": "1.0", "effective": "2025-10-15",
    "owner": VP_PEOPLE, "approver": "Human Resources and Compensation Committee of the Board", "next_review": "2026-10-15",
    "applies": "Board committee members, CEO, VP People and Culture, Director Total Rewards",
    "classification": "Restricted, HR group only",
    "header_banner": "RESTRICTED - HR Only",
    "sections": [
        ("Purpose", [
            "This paper summarises the annual review of executive compensation prepared for the Human Resources and "
            "Compensation Committee meeting on 2025-11-18. It contains personal compensation data and must not be "
            "shared outside the committee and the named HR staff.",
        ]),
        ("Market Comparison", [
            "The executive peer group includes 14 regulated North American utilities with revenue between 0.5 and "
            "2 times Harbourline's. Target total direct compensation is set at the 50th percentile of the peer group. "
            "The independent adviser is Northbridge Compensation Advisors (fictional firm).",
        ]),
        ("Recommended Fiscal 2026 Compensation", [
            ("table", ["Executive", "Current base (CAD)", "Recommended base (CAD)", "STIP target", "LTIP target"], [
                ("Eleanor Voss, President and CEO", "665,000", "685,000", "75% of base", "150% of base"),
                ("Rajiv Menon, Chief Financial Officer", "440,000", "455,000", "50% of base", "90% of base"),
                ("Dana Okafor, Chief Operating Officer", "425,000", "440,000", "50% of base", "90% of base"),
                ("Andrew Kim, General Counsel", "360,000", "372,000", "40% of base", "60% of base"),
                ("Margaret Osei, VP People and Culture", "315,000", "326,000", "40% of base", "60% of base"),
            ], [5.3, 2.8, 3, 2.7, 2.7]),
            "The recommended executive merit budget is 3.5% of executive base payroll.",
        ]),
        ("Short-Term Incentive Plan (STIP) Scorecard, Fiscal 2025", [
            ("table", ["Measure", "Weight", "Result", "Score"], [
                ("Safety: total recordable injury rate", "25%", "0.92 (target 1.00)", "115%"),
                ("Reliability: SAIDI minutes", "25%", "88 (target 95)", "120%"),
                ("Customer satisfaction", "20%", "84% (target 85%)", "95%"),
                ("Net income versus budget", "30%", "101.5%", "105%"),
            ], [7, 2.5, 4, 3]),
            "Weighted STIP payout factor: 109%. The Committee may apply discretion of plus or minus 10 points.",
        ]),
        ("Governance", [
            ("bullets", [
                "Clawback applies to incentive pay for 3 years after payment in case of restatement or misconduct.",
                "CEO share ownership guideline: 3 times base salary within 5 years.",
                "Recommendations require Board approval at the 2025-12-09 meeting.",
            ]),
        ]),
    ],
    "revisions": [
        ("1.0", "2025-10-15", "Margaret Osei", "Paper prepared for the 2025-11-18 committee meeting."),
    ],
}

INVESTIGATION = {
    "title": "Workplace Investigation Protocol",
    "doc_id": "HR-RST-103", "version": "2.0", "effective": "2025-05-01",
    "owner": DIR_ER, "approver": GC, "next_review": "2026-04-30",
    "applies": "Human Resources investigators and approved external investigators",
    "classification": "Restricted, HR group only",
    "header_banner": "RESTRICTED - HR Only",
    "sections": [
        ("Purpose", [
            "This protocol sets the internal standard for planning and conducting workplace investigations into "
            "harassment, discrimination, violence, misconduct and Code of Conduct complaints.",
        ]),
        ("Triage and Assignment", [
            "Every complaint is triaged within 2 business days by the Director, Employee Relations, and given an "
            "investigation number in the format INV-YYYY-NNN. Triage decides whether the matter is investigated "
            "internally, externally, or resolved informally.",
            ("table", ["Situation", "Investigator"], [
                ("Respondent below Director level", "Trained HR investigator from another region"),
                ("Respondent is Director or above", "External investigator engaged through Legal"),
                ("Complaint involves an HR employee", "External investigator engaged through Legal"),
                ("Possible criminal conduct", "Legal and Corporate Security jointly; police may be notified"),
            ], [7, 9.5]),
        ]),
        ("Conduct of the Investigation", [
            ("bullets", [
                "Target completion for an internal investigation is 45 calendar days from assignment, well inside the 90-day outer target published in HR-POL-031. Any extension is approved by the Director, Employee Relations, and the parties are informed.",
                "The standard of proof is the balance of probabilities.",
                "Interviews are conducted by two people (investigator and note taker). Audio recording is not permitted.",
                "Witness statements are reviewed and signed by the witness within 3 business days of the interview.",
                "Interim measures (schedule changes, paid administrative leave) may be applied during the investigation and do not imply a finding.",
            ]),
        ]),
        ("Evidence and Records", [
            "All evidence is stored in the restricted ER-Vault library with access limited to the investigator, the "
            "Director, Employee Relations, and Legal. Investigation files are retained for 7 years after closure, "
            "or longer if litigation is pending. Email searches require written approval from the General Counsel.",
        ]),
        ("Reporting", [
            "The investigator issues a confidential report with findings (substantiated, not substantiated, or "
            "inconclusive) within 5 business days of the last interview. The complainant and respondent receive a "
            "written summary of results, not the full report.",
        ]),
        ("Related Documents", [
            ("bullets", ["Disciplinary Case Handling Procedure (HR-RST-101)",
                         "Workplace Harassment and Violence Prevention Policy, Ontario (HR-POL-031)",
                         "Code of Business Conduct (HR-POL-001)"]),
        ]),
    ],
    "revisions": [
        ("1.0", "2022-05-01", "Employee Relations", "Initial protocol."),
        ("1.1", "2023-11-01", "Hamid Qureshi", "ER-Vault introduced for evidence storage."),
        ("2.0", "2025-05-01", "Hamid Qureshi", "Target completion reduced from 60 to 45 days; audio recording prohibited."),
    ],
}


# ==========================================================================
# 15. Signed Policy Acknowledgement Scan (image-only pdf)
# ==========================================================================

FONT_DIRS = ["/usr/share/fonts/truetype/dejavu", "/usr/share/fonts/dejavu", "/Library/Fonts",
             "C:/Windows/Fonts"]


def load_font(names, size):
    from PIL import ImageFont
    for d in FONT_DIRS:
        for n in names:
            p = os.path.join(d, n)
            if os.path.exists(p):
                return ImageFont.truetype(p, size)
    return ImageFont.load_default(size=size)


ACK_LINES = [
    ("title", "HARBOURLINE ENERGY CO."),
    ("title2", "Employee Policy Acknowledgement Form"),
    ("small", "Form HR-F-009  (Rev. 2024-11)          Acknowledgement reference: ACK-2025-0457"),
    ("gap", ""),
    ("label", "Employee name:  Daniel Kowalczyk"),
    ("label", "Employee ID:  HL-20931          Department:  Distribution Operations, Ontario East"),
    ("label", "Job title:  Powerline Technician          Work location:  Kingston Service Centre"),
    ("gap", ""),
    ("body", "I acknowledge that I have received, read and understood the following Harbourline policies:"),
    ("check", "Code of Business Conduct (HR-POL-001), version 5.1"),
    ("check", "Workplace Harassment and Violence Prevention Policy, Ontario (HR-POL-031), version 3.2"),
    ("check", "Safety Incident Reporting Standard (HS-POL-005), version 6.0"),
    ("check", "Overtime, On-Call and Storm Call-Out Policy (HR-POL-038), version 4.0"),
    ("gap", ""),
    ("body", "I understand that I must follow these policies as a condition of employment, that policies may be"),
    ("body", "updated, and that I must re-acknowledge them every 24 months or whenever HR notifies me of a"),
    ("body", "material change. Questions should be directed to my HR Business Partner."),
    ("gap", ""),
    ("body", "Next re-acknowledgement due: March 2027."),
    ("gap", ""),
    ("gap", ""),
    ("sig", "Employee signature:"),
    ("label", "Date signed:  2025-03-14"),
    ("gap", ""),
    ("label", "Witnessed by (HR):  Luc Tremblay, HR Business Partner, Ontario East"),
    ("sig2", "HR signature:"),
    ("gap", ""),
    ("small", "Original retained in employee file. Scanned copy filed in HR-Policies library. Scan batch SCN-0314-07."),
]


def render_ack_image():
    import random
    from PIL import Image, ImageDraw, ImageFilter
    rnd = random.Random(20250314)
    W, H = 1275, 1650  # US Letter at 150 dpi
    img = Image.new("L", (W, H), 250)
    d = ImageDraw.Draw(img)
    f_title = load_font(["DejaVuSerif-Bold.ttf", "DejaVuSans-Bold.ttf", "timesbd.ttf"], 34)
    f_title2 = load_font(["DejaVuSans-Bold.ttf", "arialbd.ttf"], 28)
    f_body = load_font(["DejaVuSans.ttf", "arial.ttf"], 21)
    f_small = load_font(["DejaVuSans.ttf", "arial.ttf"], 17)
    y = 110
    d.rectangle([90, 80, W - 90, H - 80], outline=60, width=2)
    for kind, text in ACK_LINES:
        if kind == "title":
            d.text((W // 2, y), text, font=f_title, fill=20, anchor="ma")
            y += 52
        elif kind == "title2":
            d.text((W // 2, y), text, font=f_title2, fill=20, anchor="ma")
            y += 50
            d.line([130, y, W - 130, y], fill=40, width=3)
            y += 20
        elif kind == "small":
            d.text((130, y), text, font=f_small, fill=45)
            y += 34
        elif kind == "gap":
            y += 22
        elif kind == "label" or kind == "body":
            d.text((130, y), text, font=f_body, fill=25)
            y += 36
        elif kind == "check":
            d.rectangle([150, y + 3, 170, y + 23], outline=25, width=2)
            # hand-drawn tick
            d.line([(152, y + 12), (159, y + 21), (174, y - 2)], fill=15, width=3)
            d.text((190, y), text, font=f_body, fill=25)
            y += 36
        elif kind in ("sig", "sig2"):
            d.text((130, y + 40), text, font=f_body, fill=25)
            x0 = 380
            d.line([x0, y + 65, x0 + 470, y + 65], fill=30, width=2)
            # scribbled signature: a jittered cursive-like polyline
            pts = []
            n = 70 if kind == "sig" else 45
            amp = 26 if kind == "sig" else 18
            for i in range(n):
                t = i / float(n - 1)
                import math
                xx = x0 + 20 + t * (380 if kind == "sig" else 260)
                yy = y + 45 - amp * math.sin(t * 19.0 + rnd.random() * 0.6) * (0.6 + 0.4 * math.cos(t * 5.0))
                pts.append((xx, yy + rnd.uniform(-3, 3)))
            d.line(pts, fill=10, width=3, joint="curve")
            d.line([(x0 + 30, y + 58), (x0 + 330, y + 48)], fill=10, width=2)
            y += 95
    # scanner artefacts: speckle noise, faint vertical streak, blur, rotation
    px = img.load()
    for _ in range(9000):
        x, yv = rnd.randrange(W), rnd.randrange(H)
        px[x, yv] = rnd.choice((90, 120, 160, 200))
    for yv in range(H):
        px[1010, yv] = min(px[1010, yv], 225)
    img = img.filter(ImageFilter.GaussianBlur(0.7))
    img = img.rotate(-0.9, resample=Image.BICUBIC, expand=False, fillcolor=235)
    return img


def render_scan_pdf(path):
    """Write a one-page PDF whose only content is a JPEG of the scanned form.

    The PDF is assembled by hand (no reportlab canvas) so that it contains no
    text operators and no font resources at all, like a scanner output.
    """
    img = render_ack_image()
    buf = io.BytesIO()
    img.save(buf, format="JPEG", quality=70, optimize=False)
    jpg = buf.getvalue()
    w, h = img.size
    pw, ph = 612, 792  # US Letter in points
    content = ("q %d 0 0 %d 0 0 cm /Im1 Do Q" % (pw, ph)).encode("ascii")
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        ("<< /Type /Page /Parent 2 0 R /MediaBox [0 0 %d %d] /Resources << /XObject << /Im1 4 0 R >> >> "
         "/Contents 5 0 R >>" % (pw, ph)).encode("ascii"),
        ("<< /Type /XObject /Subtype /Image /Width %d /Height %d /ColorSpace /DeviceGray /BitsPerComponent 8 "
         "/Filter /DCTDecode /Length %d >>\nstream\n" % (w, h, len(jpg))).encode("ascii") + jpg + b"\nendstream",
        ("<< /Length %d >>\nstream\n" % len(content)).encode("ascii") + content + b"\nendstream",
        b"<< /Producer (Scan to SharePoint) /Creator (Office MFP Scanner) /Title (Scan 2025-03-14 SCN-0314-07) >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = []
    for i, body in enumerate(objs, 1):
        offsets.append(out.tell())
        out.write(("%d 0 obj\n" % i).encode("ascii") + body + b"\nendobj\n")
    xref = out.tell()
    out.write(("xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)).encode("ascii"))
    for off in offsets:
        out.write(("%010d 00000 n \n" % off).encode("ascii"))
    out.write(("trailer\n<< /Size %d /Root 1 0 R /Info 6 0 R >>\nstartxref\n%d\n%%%%EOF\n"
               % (len(objs) + 1, xref)).encode("ascii"))
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "wb") as fh:
        fh.write(out.getvalue())


# ==========================================================================
# Main
# ==========================================================================

def build(root):
    hr = os.path.join(root, "data", "sharepoint", "Harbourline-Hub", "HR-Policies")
    rs = os.path.join(hr, "Restricted")
    jobs = [
        (render_docx, leave_policy(3), os.path.join(hr, "Leave-Policy-v3-2024.docx")),
        (render_pdf, leave_policy(4), os.path.join(hr, "Leave-Policy-v4-2025.pdf")),
        (render_pdf, CODE_OF_CONDUCT, os.path.join(hr, "Code-of-Conduct.pdf")),
        (render_docx, REMOTE_WORK, os.path.join(hr, "Remote-Work-Policy.docx")),
        (render_docx, HARASSMENT, os.path.join(hr, "Harassment-Prevention-Ontario.docx")),
        (render_docx, FMLA, os.path.join(hr, "FMLA-Guidance-US.docx")),
        (render_docx, OVERTIME, os.path.join(hr, "Overtime-and-On-Call.docx")),
        (render_docx, TRAVEL, os.path.join(hr, "Travel-Policy.docx")),
        (render_pdf, PARENTAL, os.path.join(hr, "Parental-Leave.pdf")),
        (render_docx, BEREAVEMENT, os.path.join(hr, "Bereavement-Leave.docx")),
        (render_docx, PERFORMANCE, os.path.join(hr, "Performance-Review-Process.docx")),
        (render_pdf, SAFETY, os.path.join(hr, "Safety-Incident-Reporting.pdf")),
        (render_docx, ACCOMMODATION, os.path.join(hr, "Workplace-Accommodation.docx")),
        (render_docx, GRIEVANCE, os.path.join(hr, "Grievance-Procedure.docx")),
        (render_docx, comp_spec(), os.path.join(hr, "Compensation-Bands-2025.docx")),
        (render_docx, ER_CONTACTS, os.path.join(hr, "Employee-Relations-Contacts.docx")),
        (render_docx, DISCIPLINARY, os.path.join(rs, "Disciplinary-Case-Handling.docx")),
        (render_docx, EXEC_COMP, os.path.join(rs, "Executive-Compensation-Review.docx")),
        (render_pdf, INVESTIGATION, os.path.join(rs, "Investigation-Protocol.pdf")),
    ]
    for fn, spec, path in jobs:
        fn(spec, path)
        print("wrote %s" % os.path.relpath(path, root))
    scan = os.path.join(hr, "Signed-Policy-Acknowledgement-Scan.pdf")
    render_scan_pdf(scan)
    print("wrote %s" % os.path.relpath(scan, root))


def main():
    import argparse
    here = os.path.dirname(os.path.abspath(__file__))
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--root", default=os.path.abspath(os.path.join(here, "..", "..")),
                    help="Repository root (default: two levels above this script)")
    ap.add_argument("--print-bands", action="store_true", help="Print the compensation band table and exit")
    a = ap.parse_args()
    if a.print_bands:
        for r in comp_bands():
            print("| %s | %s | %s | %s | %s | %s | %s |" % tuple([r[0]] + [money(x) for x in r[1:]]))
        return
    build(a.root)


if __name__ == "__main__":
    main()
