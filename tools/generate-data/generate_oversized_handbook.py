#!/usr/bin/env python3
"""Generate the oversized Employee-Handbook-Full.docx for Lab 3 (file size caveat).

The handbook is real, readable text (20 chapters built from templates) plus
embedded photo-style PNG images that make the file large. The default target is
9 MB, above the 7 MB limit that applies when the tenant has no Microsoft 365
Copilot license (reference/limits.md CS-K01) and far below the 200 MB limit
that applies with tenant graph grounding (CS-K02).

Chapter 14 holds one fact that exists in no other course file:
  "Harbourline Service Recognition Awards: employees reaching 25 years of
   service receive 3 extra days of leave (one time) and a crystal award."
Labs use it to show that an excluded file cannot answer the question.

The output is gitignored and created at setup time. Deterministic (fixed seed).

Usage:
  python3 tools/generate-data/generate_oversized_handbook.py [--out PATH] [--target-mb 9]
"""

import argparse
import io
import math
import os
import random
import sys

from docx import Document
from docx.shared import Cm, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from build_hr_policies import check_text, load_font, normalise_zip, set_core  # noqa: E402

DEFAULT_OUT = os.path.join(HERE, "..", "..", "data", "sharepoint", "Harbourline-Hub", "HR-Policies",
                           "Employee-Handbook-Full.docx")
SEED = 20250901

UNIQUE_FACT = ("Harbourline Service Recognition Awards: employees who reach 25 years of continuous service "
               "receive a one-time award of 3 extra days of paid leave and a crystal award, presented at the "
               "annual Service Recognition Dinner in October. The extra days must be used within 12 months of "
               "the anniversary.")

# (chapter title, [section titles], image caption)
CHAPTERS = [
    ("Welcome to Harbourline", ["Our history", "Our mission and values", "How we are organised", "Our service territory"],
     "Figure 1.1: Harbourline service territory overview"),
    ("Employment Basics", ["Employment categories", "Probationary period", "Hours of work", "Pay schedule"],
     "Figure 2.1: Pay calendar illustration"),
    ("Workplace Safety", ["Safety commitments", "Personal protective equipment", "Stop work authority", "Fatigue management"],
     "Figure 3.1: Line crew performing a job briefing"),
    ("Respectful Workplace", ["Our expectations", "Diversity and inclusion", "Harassment and violence prevention", "Speaking up"],
     "Figure 4.1: Employee resource group event"),
    ("Time Away from Work", ["Annual leave", "Statutory and company holidays", "Sick time", "Other leaves"],
     "Figure 5.1: Leave request workflow"),
    ("Benefits Overview", ["Health and dental", "Retirement savings", "Employee and Family Assistance Program", "Wellness"],
     "Figure 6.1: Benefits enrolment timeline"),
    ("Learning and Development", ["Technical training", "Apprenticeship programs", "Tuition assistance", "Leadership programs"],
     "Figure 7.1: Apprentice lineworkers at the training yard"),
    ("Performance and Career Growth", ["Setting goals", "Feedback and check-ins", "Internal job postings", "Mentoring"],
     "Figure 8.1: Annual performance cycle"),
    ("Technology and Information Security", ["Acceptable use", "Passwords and multifactor authentication", "Phishing", "Operational technology"],
     "Figure 9.1: Control room operator workstation"),
    ("Customers and Communities", ["Customer service standards", "Working on customer property", "Community investment", "Media enquiries"],
     "Figure 10.1: Community tree planting day"),
    ("Storm and Emergency Response", ["Storm levels", "Your role in restoration", "Mutual assistance", "Family preparedness"],
     "Figure 11.1: Crews restoring service after an ice storm"),
    ("Vehicles and Equipment", ["Fleet vehicles", "Driver qualifications", "Equipment care", "Reporting damage"],
     "Figure 12.1: Bucket truck pre-trip inspection"),
    ("Environment and Sustainability", ["Environmental commitments", "Spill response", "Wildlife protection", "Reducing our footprint"],
     "Figure 13.1: Substation oil containment"),
    ("Recognition and Service Awards", ["Everyday recognition", "Harbourline Service Recognition Awards", "Safety excellence awards", "Nominating a colleague"],
     "Figure 14.1: Service Recognition Dinner"),
    ("Health and Wellbeing", ["Mental health support", "Physical wellbeing", "Return to work support", "Substance use"],
     "Figure 15.1: Wellness fair"),
    ("Ethics and Compliance", ["Code of Business Conduct", "Conflicts of interest", "Regulatory compliance", "Records retention"],
     "Figure 16.1: Compliance training dashboard"),
    ("Working in the United States", ["New York operations", "Ohio operations", "US benefits differences", "US leave programs"],
     "Figure 17.1: Albany operations centre"),
    ("Pay and Expenses", ["Payroll questions", "Direct deposit", "Travel and expenses", "Corporate cards"],
     "Figure 18.1: Expense claim process"),
    ("Leaving Harbourline", ["Resignation", "Retirement", "Exit interviews", "Returning equipment"],
     "Figure 19.1: Retirement celebration"),
    ("Contacts and Resources", ["HR Business Partners", "Help desks", "Hotlines", "Policy library"],
     "Figure 20.1: Harbourline Hub home page"),
]

OPENERS = [
    "This section explains {topic_l} at Harbourline and what it means for you in your day-to-day work.",
    "Every employee should understand {topic_l}, whether they work in an office, a service centre or in the field.",
    "Harbourline's approach to {topic_l} reflects our commitment to safe, reliable and affordable electricity service.",
    "{topic} is an area where the expectations of our customers, our regulators and our colleagues come together.",
]
MIDDLES = [
    "In Ontario, Harbourline follows the requirements of the Ontario Energy Board and provincial employment and "
    "safety legislation. In New York and Ohio, the equivalent state and federal rules apply, and the US Human "
    "Resources team can explain where practices differ.",
    "Your people leader is your first point of contact for questions. If your leader is not available, your regional "
    "HR Business Partner can help, and the HR self-service portal at hrportal.harbourline.example holds the forms you need.",
    "The detailed rules are set out in the related policy in the HR-Policies library. This handbook summarises them "
    "in plain language; where the handbook and a policy differ, the policy governs.",
    "We review our practices every year with input from employees, the Joint Health and Safety Committees and, for "
    "represented employees, the Harbourline Unit of Local 4410.",
    "Managers are expected to explain these practices during onboarding and to revisit them at team meetings at least "
    "once a year, so that new and experienced employees share the same understanding.",
    "Field employees can access this information on their rugged tablets through the Harbourline Hub, even when "
    "working from a truck or a remote substation with limited connectivity.",
]
CLOSERS = [
    "If something is unclear, ask. Asking questions early prevents mistakes later.",
    "Thank you for helping make Harbourline a great place to work and a reliable service for our customers.",
    "Remember that safety comes first in every decision, including decisions about {topic_l}.",
    "Suggestions for improving how we handle {topic_l} are welcome through the Harbourline Hub feedback form.",
]

# Chapter-specific paragraphs, keyed by (chapter number, section title).
SPECIAL = {
    (1, "Our history"): [
        "Harbourline Energy Co. began in 1911 as the Harbour Light and Power Company, supplying street lighting to "
        "the Toronto waterfront. Through mergers with rural distribution utilities it grew into a regulated electric "
        "utility serving customers across Ontario, and in 2016 it acquired distribution businesses in New York and Ohio.",
        "Today about 2,000 employees operate and maintain substations, overhead and underground lines, and the "
        "systems that keep them running, under oversight from the Ontario Energy Board, NERC, FERC and state public "
        "service commissions.",
    ],
    (5, "Annual leave"): [
        "Your annual leave entitlement, carry-over and scheduling rules are set out in the Leave Policy (HR-POL-012) "
        "in the HR-Policies library. Always check the most recent version of that policy. This handbook does not "
        "repeat the entitlement table because it changes from time to time.",
    ],
    (14, "Harbourline Service Recognition Awards"): [
        "The Harbourline Service Recognition Awards celebrate long service at milestones of 5, 10, 15, 20, 25, 30 "
        "and 35 years. Milestones from 5 to 20 years are recognised with a certificate signed by the President and "
        "Chief Executive Officer and a choice of gift from the recognition catalogue.",
        UNIQUE_FACT,
        "Employees reaching 30 and 35 years receive the same crystal award, a catalogue gift and an invitation for a "
        "guest to attend the Service Recognition Dinner. Awards are coordinated by the Total Rewards team, who "
        "confirm eligibility from continuous service dates in the HR system each August.",
    ],
}


def chapter_text(rnd, ch_no, topic):
    key = (ch_no, topic)
    paras = []
    if key in SPECIAL:
        paras.extend(SPECIAL[key])
    fmt = {"topic": topic, "topic_l": topic[0].lower() + topic[1:]}
    paras.append(rnd.choice(OPENERS).format(**fmt) + " " + rnd.choice(MIDDLES))
    mids = rnd.sample(MIDDLES, 2)
    paras.append(mids[0] + " " + mids[1])
    paras.append(rnd.choice(CLOSERS).format(**fmt))
    return [check_text(p) for p in paras]


def make_image(rnd, w, h, caption_seed, kind):
    """Photo-style image: gradient sky, landscape shapes, utility poles, and sensor noise.

    Per-pixel noise keeps the PNG large (it does not compress well), which is
    the point of this generator.
    """
    from PIL import Image, ImageDraw
    img = Image.new("RGB", (w, h))
    d = ImageDraw.Draw(img)
    top = (rnd.randint(90, 140), rnd.randint(140, 190), rnd.randint(200, 240))
    bot = (rnd.randint(200, 240), rnd.randint(200, 230), rnd.randint(170, 210))
    horizon = int(h * rnd.uniform(0.55, 0.7))
    for y in range(h):
        t = y / float(h)
        c = tuple(int(top[i] + (bot[i] - top[i]) * t) for i in range(3))
        d.line([0, y, w, y], fill=c)
    ground = (rnd.randint(60, 100), rnd.randint(100, 140), rnd.randint(50, 80))
    pts = [(0, horizon)]
    for x in range(0, w + 40, 40):
        pts.append((x, horizon + int(18 * math.sin(x / 90.0 + caption_seed))))
    pts += [(w, h), (0, h)]
    d.polygon(pts, fill=ground)
    if kind % 3 == 0:
        # utility poles and conductors
        xs = [int(w * f) for f in (0.12, 0.42, 0.72, 0.98)]
        for x in xs:
            d.rectangle([x - 5, horizon - 230, x + 5, horizon + 10], fill=(80, 60, 40))
            d.rectangle([x - 60, horizon - 225, x + 60, horizon - 215], fill=(80, 60, 40))
        for off in (-55, 0, 55):
            for a, b in zip(xs, xs[1:]):
                d.arc([a + off, horizon - 260, b + off, horizon - 190], 20, 160, fill=(30, 30, 30), width=2)
    elif kind % 3 == 1:
        # substation: transformers and bus structure
        for i in range(4):
            x = int(w * (0.15 + i * 0.2))
            d.rectangle([x, horizon - 120, x + 90, horizon + 20], fill=(150, 150, 140), outline=(60, 60, 60), width=3)
            for j in range(3):
                d.rectangle([x + 12 + j * 26, horizon - 170, x + 22 + j * 26, horizon - 120], fill=(120, 90, 60))
        d.line([int(w * 0.1), horizon - 190, int(w * 0.9), horizon - 190], fill=(90, 90, 90), width=6)
    else:
        # trucks and people silhouettes
        for i in range(3):
            x = int(w * (0.1 + i * 0.3))
            d.rectangle([x, horizon - 60, x + 170, horizon + 30], fill=(235, 200, 40), outline=(40, 40, 40), width=3)
            d.rectangle([x + 170, horizon - 30, x + 230, horizon + 30], fill=(235, 200, 40), outline=(40, 40, 40), width=3)
            d.ellipse([x + 20, horizon + 15, x + 60, horizon + 55], fill=(30, 30, 30))
            d.ellipse([x + 170, horizon + 15, x + 210, horizon + 55], fill=(30, 30, 30))
            d.line([x + 60, horizon - 60, x + 140, horizon - 200], fill=(200, 200, 200), width=8)
    # sensor noise (high entropy, keeps PNG size up)
    noise = bytes(rnd.getrandbits(8) for _ in range(w * h * 3))
    nimg = Image.frombytes("RGB", (w, h), noise)
    img = Image.blend(img, nimg, 0.22)
    buf = io.BytesIO()
    img.save(buf, format="PNG", optimize=False, compress_level=6)
    return buf.getvalue()


def build(out, target_bytes):
    rnd = random.Random(SEED)
    doc = Document()
    set_core(doc, {"title": "Harbourline Employee Handbook (Full Edition)",
                   "subject": "Employee handbook", "keywords": "handbook; employees; service awards"})
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)
    sec = doc.sections[0]
    sec.header.paragraphs[0].add_run("Harbourline Energy Co.  |  HR-HBK-001  |  Employee Handbook, Full Edition 2025").font.size = Pt(8)
    sec.footer.paragraphs[0].add_run("This handbook summarises Harbourline policies. Where it differs from a policy, the policy governs.").font.size = Pt(8)

    doc.add_paragraph("Harbourline Energy Co.").runs[0].bold = True
    doc.add_heading("Employee Handbook, Full Edition", level=0)
    doc.add_paragraph("Document ID HR-HBK-001, edition 2025, effective 2025-09-01. Owner: Margaret Osei, "
                      "Vice President, People and Culture.")
    doc.add_heading("Contents", level=1)
    for n, (title, _, _) in enumerate(CHAPTERS, 1):
        doc.add_paragraph("Chapter %d: %s" % (n, title))

    # Image budget: about 85% of the target across one image per chapter; any shortfall is
    # topped up with appendix photos. Bytes per pixel is calibrated from a sample image.
    img_w = 1400
    bpp = len(make_image(random.Random(1), img_w, 200, 0, 0)) / float(img_w * 200)
    per_img = int(target_bytes * 0.85 / len(CHAPTERS))
    img_h = max(120, int(per_img / (img_w * bpp)))
    for n, (title, sections, caption) in enumerate(CHAPTERS, 1):
        doc.add_page_break()
        doc.add_heading("Chapter %d: %s" % (n, title), level=1)
        for s_no, stitle in enumerate(sections, 1):
            doc.add_heading("%d.%d %s" % (n, s_no, stitle), level=2)
            for p in chapter_text(rnd, n, stitle):
                doc.add_paragraph(p)
            if s_no == 2:
                doc.add_picture(io.BytesIO(make_image(rnd, img_w, img_h, n, n)), width=Cm(16))
                cp = doc.add_paragraph(caption)
                cp.runs[0].italic = True

    def save():
        doc.save(out)
        normalise_zip(out)
        return os.path.getsize(out)

    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    size = save()
    extra = 0
    if size < target_bytes:
        doc.add_page_break()
        doc.add_heading("Appendix A: Photo Gallery", level=1)
        doc.add_paragraph("Photographs from Harbourline field operations, training days and community events.")
    while size < target_bytes:
        extra += 1
        need = target_bytes - size
        h = max(80, min(img_h, int(need / (img_w * bpp)) + 40))
        doc.add_picture(io.BytesIO(make_image(rnd, img_w, h, 100 + extra, extra)), width=Cm(16))
        doc.add_paragraph("Figure A.%d: Field operations photograph" % extra).runs[0].italic = True
        size = save()
    return size


def main():
    ap = argparse.ArgumentParser(description="Generate the oversized Harbourline employee handbook (.docx).")
    ap.add_argument("--out", default=os.path.abspath(DEFAULT_OUT), help="Output .docx path")
    ap.add_argument("--target-mb", type=float, default=9.0, help="Minimum file size in MB (default 9)")
    a = ap.parse_args()
    if a.target_mb <= 7:
        print("warning: target %.1f MB does not exceed the 7 MB limit in CS-K01" % a.target_mb, file=sys.stderr)
    target = int(a.target_mb * 1024 * 1024)
    size = build(a.out, target)
    print("wrote %s" % a.out)
    print("final size: %d bytes (%.2f MB)" % (size, size / 1024.0 / 1024.0))


if __name__ == "__main__":
    main()
