#!/usr/bin/env python3
"""Build the committed Harbourline-Operations course data.

Writes:
  data/sharepoint/Harbourline-Operations/Procedures/
      Outage-Response-Procedure.docx
      Storm-Restoration-Playbook.pdf
      Lockout-Tagout-Safety-Manual.pdf
      Substation-Switching-Procedure.docx
      Vegetation-Management.docx
      Crew-Briefing-Deck.pptx          (key values only inside images)
      Transformer-Equipment-Specs.xlsx (merged headers, units row, merged row labels)
  data/sharepoint/lab-09-drops/
      Field-Report-2026-10-01-Kingston.docx
      Field-Report-2026-10-02-Watertown.docx   (asset ID deliberately blank)
      Field-Report-2026-10-03-Ashtabula.docx   (severity Critical)

Deterministic: fixed random seed, fixed document dates. Re-runnable.
Usage: python3 build_operations.py [--root <repo root>]
"""

import argparse
import datetime as dt
import io
import os
import random

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Inches, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor as PptxRGB
from pptx.util import Emu, Inches as PInches, Pt as PPt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (PageBreak, Paragraph, SimpleDocTemplate, Spacer,
                                Table, TableStyle)

SEED = 20261001
FIXED_DT = dt.datetime(2026, 9, 15, 9, 0, 0)
COMPANY = "Harbourline Energy Co."
FONT_DIR = "/usr/share/fonts/truetype/dejavu"

# ---------------------------------------------------------------------------
# Shared block renderer for docx
# ---------------------------------------------------------------------------
# Block types:
#   ("h1", text) ("h2", text) ("p", text) ("bullets", [..]) ("steps", [..])
#   ("table", [headers], [[row], ...]) ("note", text) ("pagebreak",)


def _docx_base():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(1)
        s.top_margin = s.bottom_margin = Inches(0.9)
    cp = doc.core_properties
    cp.author = COMPANY
    cp.last_modified_by = COMPANY
    cp.created = FIXED_DT
    cp.modified = FIXED_DT
    cp.revision = 1
    return doc


def _docx_table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(10)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(str(v))
            r.font.size = Pt(10)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    doc.add_paragraph()
    return t


def build_docx(path, title, meta, blocks, revisions):
    doc = _docx_base()
    doc.core_properties.title = title
    doc.core_properties.subject = meta.get("Document ID", "")
    # header / footer
    sec = doc.sections[0]
    hp = sec.header.paragraphs[0]
    hp.text = f"{COMPANY}  |  {meta.get('Document ID', '')}  |  {title}"
    hp.runs[0].font.size = Pt(8)
    fp = sec.footer.paragraphs[0]
    fp.text = "Controlled document. Printed copies are uncontrolled. Verify the current revision on the Harbourline-Operations site."
    fp.runs[0].font.size = Pt(8)

    p = doc.add_paragraph()
    r = p.add_run(COMPANY)
    r.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(0x0B, 0x3D, 0x5C)
    doc.add_heading(title, 0)
    _docx_table(doc, ["Field", "Value"], [[k, v] for k, v in meta.items()], widths=[1.8, 4.7])
    render_docx_blocks(doc, blocks)
    doc.add_heading("Revision History", 1)
    _docx_table(doc, ["Revision", "Date", "Author", "Summary of changes"], revisions,
                widths=[0.8, 1.1, 1.6, 3.0])
    doc.save(path)


def render_docx_blocks(doc, blocks):
    for b in blocks:
        kind = b[0]
        if kind == "h1":
            doc.add_heading(b[1], 1)
        elif kind == "h2":
            doc.add_heading(b[1], 2)
        elif kind == "p":
            doc.add_paragraph(b[1])
        elif kind == "note":
            p = doc.add_paragraph()
            r = p.add_run(b[1])
            r.italic = True
        elif kind == "bullets":
            for item in b[1]:
                doc.add_paragraph(item, style="List Bullet")
        elif kind == "steps":
            for i, item in enumerate(b[1], 1):
                p = doc.add_paragraph()
                r = p.add_run(f"Step {i}. ")
                r.bold = True
                p.add_run(item)
                p.paragraph_format.left_indent = Inches(0.25)
        elif kind == "table":
            _docx_table(doc, b[1], b[2])
        elif kind == "pagebreak":
            doc.add_page_break()
        else:
            raise ValueError(kind)


# ---------------------------------------------------------------------------
# Shared block renderer for PDF (reportlab)
# ---------------------------------------------------------------------------

def build_pdf(path, title, meta, blocks, revisions):
    styles = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=styles["BodyText"], fontName="Helvetica",
                          fontSize=10, leading=13.5, spaceAfter=6)
    h1 = ParagraphStyle("h1", parent=styles["Heading1"], fontSize=14, spaceBefore=10,
                        spaceAfter=6, textColor=colors.HexColor("#0B3D5C"))
    h2 = ParagraphStyle("h2", parent=styles["Heading2"], fontSize=11.5, spaceBefore=8,
                        spaceAfter=4, textColor=colors.HexColor("#0B3D5C"))
    ttl = ParagraphStyle("ttl", parent=styles["Title"], fontSize=20, leading=24)
    cell = ParagraphStyle("cell", parent=body, fontSize=8.8, leading=11, spaceAfter=0)
    cellb = ParagraphStyle("cellb", parent=cell, fontName="Helvetica-Bold")
    note = ParagraphStyle("note", parent=body, fontName="Helvetica-Oblique")
    bullet = ParagraphStyle("bullet", parent=body, leftIndent=14, bulletIndent=4)
    co = ParagraphStyle("co", parent=body, fontName="Helvetica-Bold", fontSize=12,
                        textColor=colors.HexColor("#0B3D5C"), alignment=TA_CENTER)

    def tbl(headers, rows, col_widths=None):
        data = [[Paragraph(h, cellb) for h in headers]]
        data += [[Paragraph(str(v), cell) for v in r] for r in rows]
        t = Table(data, colWidths=col_widths, repeatRows=1)
        t.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#7A8A99")),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DCE6F0")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        return t

    story = [Paragraph(COMPANY, co), Spacer(1, 6), Paragraph(title, ttl), Spacer(1, 6),
             tbl(["Field", "Value"], [[k, v] for k, v in meta.items()],
                 [1.8 * inch, 4.7 * inch]), Spacer(1, 10)]
    for b in blocks:
        kind = b[0]
        if kind == "h1":
            story.append(Paragraph(b[1], h1))
        elif kind == "h2":
            story.append(Paragraph(b[1], h2))
        elif kind == "p":
            story.append(Paragraph(b[1], body))
        elif kind == "note":
            story.append(Paragraph(b[1], note))
        elif kind == "bullets":
            for item in b[1]:
                story.append(Paragraph(item, bullet, bulletText="•"))
        elif kind == "steps":
            for i, item in enumerate(b[1], 1):
                story.append(Paragraph(f"<b>Step {i}.</b> {item}", bullet))
        elif kind == "table":
            story.append(tbl(b[1], b[2]))
            story.append(Spacer(1, 8))
        elif kind == "pagebreak":
            story.append(PageBreak())
    story.append(Paragraph("Revision History", h1))
    story.append(tbl(["Revision", "Date", "Author", "Summary of changes"], revisions,
                     [0.8 * inch, 1.0 * inch, 1.6 * inch, 3.1 * inch]))

    doc_id = meta.get("Document ID", "")

    def on_page(canvas, d):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#555555"))
        canvas.drawString(0.9 * inch, 10.6 * inch, f"{COMPANY}  |  {doc_id}  |  {title}")
        canvas.drawString(0.9 * inch, 0.5 * inch,
                          "Controlled document. Printed copies are uncontrolled.")
        canvas.drawRightString(7.6 * inch, 0.5 * inch, f"Page {d.page}")
        canvas.restoreState()

    doc = SimpleDocTemplate(path, pagesize=letter, leftMargin=0.9 * inch,
                            rightMargin=0.9 * inch, topMargin=0.8 * inch,
                            bottomMargin=0.8 * inch, title=title, author=COMPANY,
                            subject=doc_id, creator=COMPANY, invariant=1)
    doc.build(story, onFirstPage=on_page, onLaterPages=on_page)


# ---------------------------------------------------------------------------
# 1. Outage Response Procedure (docx)
# ---------------------------------------------------------------------------

def outage_response(path):
    meta = {
        "Document ID": "OPS-PRO-001",
        "Revision": "7",
        "Effective date": "March 1, 2026",
        "Next review": "March 1, 2027",
        "Document owner": "Renata Okafor, Director, Distribution Operations",
        "Approved by": "Graham Lindqvist, Vice President, Operations",
        "Applies to": "All Harbourline distribution regions in Ontario, New York and Ohio",
    }
    blocks = [
        ("h1", "1. Purpose"),
        ("p", "This procedure defines how Harbourline Energy Co. classifies, escalates, communicates and restores "
              "unplanned electric service interruptions on its distribution system. It sets the outage severity "
              "levels L1 to L4, the notification timelines for each level, the responsibilities of the Distribution "
              "System Operator, the Storm Coordinator and Customer Communications, and the order in which customers "
              "and facilities are restored."),
        ("p", "The procedure supports Harbourline's obligations to the Ontario Energy Board (OEB), the New York State "
              "Public Service Commission (NY PSC) and the Public Utilities Commission of Ohio (PUCO), and supports "
              "event reporting obligations under NERC reliability standards for the transmission facilities "
              "Harbourline operates."),
        ("h1", "2. Scope"),
        ("p", "This procedure applies to every unplanned interruption on Harbourline distribution feeders, laterals "
              "and secondary services, and to transmission supply interruptions that affect Harbourline customers. "
              "It does not apply to planned outages, which follow OPS-PRO-004 Planned Outage Notification."),
        ("table", ["Operating region", "Control centre", "Approximate customers served"], [
            ["Eastern Ontario (Kingston, Napanee, Brockville)", "Kingston Distribution Control Centre", "148,000"],
            ["Central Ontario (Barrie, Orillia, Midland)", "Barrie Distribution Control Centre", "162,000"],
            ["Durham (Whitby, Oshawa, Port Perry)", "Barrie Distribution Control Centre", "101,000"],
            ["New York North Country (Watertown, Canton)", "Watertown Distribution Control Centre", "71,000"],
            ["Western New York (Jamestown, Dunkirk)", "Watertown Distribution Control Centre", "57,000"],
            ["Northeast Ohio (Ashtabula, Painesville)", "Mansfield Distribution Control Centre", "52,000"],
            ["North Central Ohio (Mansfield, Ashland)", "Mansfield Distribution Control Centre", "44,000"],
        ]),
        ("h1", "3. Definitions"),
        ("bullets", [
            "Outage Management System (OMS): the GridView OMS application that predicts outage devices from customer "
            "calls and smart meter last gasp messages and tracks every outage event from creation to closure.",
            "Customers out: the count of metered service points without supply as reported by OMS at a point in time.",
            "Estimated Time of Restoration (ETR): the time by which Harbourline expects supply to be restored to a "
            "customer, published on the outage map and in customer notifications.",
            "Critical facility: a hospital, long term care home, water or wastewater treatment plant or pumping "
            "station, 911 call centre, police, fire or EMS station, or other site listed in the Critical Facilities "
            "Register maintained by each region.",
            "Life support customer: a residential customer registered with Harbourline as relying on electrically "
            "powered medical equipment.",
            "Emergency Operations Centre (EOC): a Regional EOC is staffed in the affected region's control centre; the "
            "Corporate EOC is staffed at the Toronto head office, 200 Front Street West.",
        ]),
        ("h1", "4. Roles and Responsibilities"),
        ("table", ["Role", "Primary responsibilities", "Backup"], [
            ["Distribution System Operator (DSO)",
             "Monitors SCADA and OMS 24/7. Confirms the outage, declares the initial severity level, performs remote "
             "switching to isolate faults and restore unfaulted sections, dispatches the first responder, and makes "
             "the first internal notification.",
             "Senior DSO on shift"],
            ["Storm Coordinator",
             "On call 24/7 on a weekly rotation. Owns resource allocation for L2 and above: assigns crews, calls out "
             "additional line and forestry crews, requests contractor and mutual assistance resources, sets "
             "restoration sequence within the priority order in Section 7, and confirms or changes the severity level.",
             "Regional Operations Manager"],
            ["Customer Communications",
             "Publishes outage map updates, social media posts, news releases and ETR messages; contacts critical "
             "facilities and life support customers; briefs the contact centre with approved scripts.",
             "Corporate Communications Duty Manager"],
            ["Regional Operations Manager",
             "Leads the Regional EOC at L3 and above, approves overtime extensions and contractor spend.",
             "Director, Distribution Operations"],
            ["Regulatory Affairs",
             "Determines whether an event triggers a reporting obligation to the OEB, NY PSC or PUCO, and files it.",
             "Manager, Regulatory Compliance"],
        ]),
        ("h1", "5. Outage Severity Levels"),
        ("p", "The DSO declares the initial level as soon as an outage is confirmed. The level is based on the peak "
              "number of customers out system wide at the time of assessment, or on any of the other triggers in "
              "the table, whichever gives the higher level. Only the Storm Coordinator can lower a level."),
        ("table", ["Level", "Name", "Customers out", "Other triggers", "Declared by"], [
            ["L1", "Minor", "1 to 499", "Single transformer, tap or lateral; ETR 4 hours or less", "DSO"],
            ["L2", "Elevated", "500 to 4,999",
             "Any critical facility out; any feeder breaker lockout; ETR more than 4 hours", "DSO"],
            ["L3", "Major", "5,000 to 49,999",
             "Three or more feeder lockouts in one region; more than 1,000 customers with projected outage over 24 hours",
             "DSO, confirmed by Storm Coordinator"],
            ["L4", "Emergency", "50,000 or more",
             "Loss of a transformer station or substation supply serving more than 20,000 customers; government "
             "emergency declaration in a service territory",
             "Storm Coordinator, confirmed by VP Operations"],
        ]),
        ("h1", "6. Notification Timelines"),
        ("p", "Timelines are measured from the time the level is declared. Every notification is logged in the OMS "
              "event record with the time, the person notified and the method (phone, Teams, email)."),
        ("table", ["Level", "Internal notification", "Customer communications", "Regulatory"], [
            ["L1", "DSO logs event in OMS within 15 minutes. No escalation.",
             "Outage map updates automatically within 15 minutes.", "None."],
            ["L2", "DSO notifies the on call Storm Coordinator within 30 minutes.",
             "Customer Communications posts on social media within 30 minutes and refreshes every 2 hours. Critical "
             "facility contact within 30 minutes.", "None unless a critical facility is out more than 8 hours."],
            ["L3", "Storm Coordinator notifies the Director, Distribution Operations and the VP Operations within 60 "
                   "minutes. Regional EOC activated within 2 hours.",
             "News release within 2 hours, updates every 4 hours. Life support customers contacted within 6 hours.",
             "Regulatory Affairs notified within 4 hours to assess reporting obligations."],
            ["L4", "VP Operations notifies the CEO within 1 hour. Corporate EOC activated within 1 hour.",
             "News release within 1 hour, updates every 3 hours. Life support customers contacted within 4 hours.",
             "Regulatory Affairs notified within 1 hour. Transmission Operations assesses NERC EOP-004 event "
             "reporting within 24 hours."],
        ]),
        ("h1", "7. Restoration Priority Order"),
        ("p", "Within each region, restoration work is sequenced in the following order. The Storm Coordinator may "
              "depart from the order only to remove an immediate hazard to life, and must record the reason in OMS."),
        ("steps", [
            "Public safety hazards: energized wires down, fires, and requests from police or fire services to make safe.",
            "Hospitals and long term care homes.",
            "Water and wastewater treatment plants and water pumping stations.",
            "911 call centres, police, fire and EMS stations.",
            "Registered life support customers.",
            "Transmission lines, transformer stations and substations, then feeder mainlines that restore the largest "
            "number of customers per crew hour.",
            "Telecommunications hubs, fuel distribution sites, and designated warming or cooling centres.",
            "Taps and laterals, largest customer count first.",
            "Individual services and secondary connections.",
        ]),
        ("h1", "8. Response Workflow"),
        ("h2", "8.1 Detection and confirmation"),
        ("p", "OMS creates an outage event from smart meter last gasp messages, SCADA breaker operations or customer "
              "calls. The DSO confirms the event by SCADA indication, meter ping or a first responder report, and "
              "declares the level. Unconfirmed predicted outages older than 30 minutes are reviewed by the senior DSO."),
        ("h2", "8.2 First response"),
        ("p", "The DSO dispatches the nearest trouble crew. The first responder makes the site safe, patrols the "
              "affected section, reports damage in the mobile workforce application and gives the DSO an initial "
              "assessment within 60 minutes of arrival."),
        ("h2", "8.3 Isolation and partial restoration"),
        ("p", "The DSO uses remote switching and field switching to isolate the faulted section and restore "
              "unfaulted sections. All field switching follows OPS-PRO-006 Substation and Distribution Switching "
              "Procedure. Work on isolated equipment requires a lockout under SAF-MAN-002."),
        ("h2", "8.4 Resource escalation"),
        ("p", "At L2 and above the Storm Coordinator reviews resource needs every 2 hours. At L3 and above the Storm "
              "Coordinator follows OPS-PLB-004 Storm Restoration Playbook for staging, contractor call out and mutual "
              "assistance."),
        ("h2", "8.5 Close out"),
        ("p", "An event is closed in OMS only when every customer in the event is confirmed restored by meter ping "
              "or field confirmation, and the cause code, equipment and crew time are recorded."),
        ("h1", "9. Estimated Time of Restoration Standards"),
        ("bullets", [
            "L1 and L2: an initial ETR is published within 60 minutes of the first responder's assessment.",
            "L3: a regional ETR is published within 8 hours of the level declaration.",
            "L4: a global ETR (for example, 90 percent of customers restored by a stated time) is published within 12 "
            "hours of the level declaration.",
            "An ETR that will be missed is updated at least 30 minutes before it expires.",
        ]),
        ("h1", "10. Escalation and De-escalation"),
        ("p", "The level is raised immediately when any trigger for a higher level is met. The Storm Coordinator may "
              "lower the level when customers out have been below the threshold of the current level for 2 "
              "consecutive hours and no new triggers are active. EOCs are stood down by the person who activated them."),
        ("h1", "11. Post-Event Review"),
        ("bullets", [
            "L2: crew debrief recorded in OMS within 10 business days.",
            "L3 and L4: an After Action Report is issued within 30 calendar days by the Regional Operations Manager "
            "(L3) or the Director, Distribution Operations (L4).",
            "Any missed notification timeline is reported to the Director, Distribution Operations within 5 business days.",
        ]),
        ("h1", "12. Records"),
        ("p", "OMS event records, notification logs and After Action Reports are retained for 7 years."),
        ("h1", "13. Related Documents"),
        ("bullets", [
            "OPS-PLB-004 Storm Restoration Playbook",
            "OPS-PRO-006 Substation and Distribution Switching Procedure",
            "SAF-MAN-002 Lockout and Tagout Safety Manual",
            "OPS-PRO-004 Planned Outage Notification",
        ]),
    ]
    revisions = [
        ["5", "2023-02-10", "R. Okafor", "Added Ohio regions after service territory transfer."],
        ["6", "2024-11-04", "R. Okafor", "Added life support customer contact timelines."],
        ["7", "2026-03-01", "R. Okafor", "L2 threshold changed from 250 to 500 customers; added Durham region."],
    ]
    build_docx(path, "Outage Response Procedure", meta, blocks, revisions)


# ---------------------------------------------------------------------------
# 2. Storm Restoration Playbook (pdf)
# ---------------------------------------------------------------------------

STAGING = [
    # region, primary, address, capacity, alternate
    ["Eastern Ontario", "Kingston Service Centre", "1450 Sydenham Road, Kingston ON", "40 crews",
     "Napanee Fairgrounds, 4 York Road East lot, Napanee ON"],
    ["Central Ontario", "Barrie Operations Yard", "85 Welham Road, Barrie ON", "45 crews",
     "Orillia Recreation Centre lot, 255 West Street South, Orillia ON"],
    ["Durham", "Whitby Service Centre", "1600 Thickson Road North, Whitby ON", "35 crews",
     "Port Perry Fairgrounds, 5 Simcoe Street, Port Perry ON"],
    ["New York North Country", "Watertown Service Center", "310 Coffeen Street, Watertown NY", "30 crews",
     "St. Lawrence County Fairgrounds, Gouverneur NY"],
    ["Western New York", "Jamestown Operations Yard", "92 Harrison Street, Jamestown NY", "25 crews",
     "Chautauqua County Fairgrounds, Dunkirk NY"],
    ["Northeast Ohio", "Ashtabula Service Center", "2200 Lake Avenue, Ashtabula OH", "25 crews",
     "Lake County Fairgrounds, Painesville OH"],
    ["North Central Ohio", "Mansfield Operations Center", "700 Park Avenue West, Mansfield OH", "30 crews",
     "Ashland County Fairgrounds, Ashland OH"],
]


def storm_playbook(path):
    meta = {
        "Document ID": "OPS-PLB-004",
        "Revision": "5",
        "Effective date": "May 15, 2026",
        "Next review": "May 15, 2027",
        "Document owner": "Dale Hutchins, Storm Coordinator Lead",
        "Approved by": "Renata Okafor, Director, Distribution Operations",
        "Applies to": "Distribution Operations, Forestry, Fleet, Supply Chain, Safety, and all contractor and "
                      "mutual assistance crews working on the Harbourline system",
    }
    blocks = [
        ("h1", "1. Purpose and Scope"),
        ("p", "This playbook describes how Harbourline prepares for, responds to and recovers from storms that cause "
              "an L3 Major or L4 Emergency outage under OPS-PRO-001 Outage Response Procedure. It covers storm "
              "readiness levels, the storm organization, staging areas in each region, mutual assistance, crew work "
              "and rest rules, damage assessment, logistics, and demobilization."),
        ("p", "The playbook applies in Ontario, New York and Ohio. Where a collective agreement, provincial or state "
              "law is more restrictive than this playbook, the more restrictive rule applies."),
        ("h1", "2. Storm Readiness Levels"),
        ("p", "The Storm Coordinator Lead reviews the Environment and Climate Change Canada and US National Weather "
              "Service forecasts at 06:00 and 15:00 every day and sets the readiness level for each region."),
        ("table", ["Readiness level", "Forecast triggers (any one)", "Actions"], [
            ["Green (normal)", "No trigger met", "Normal on call roster."],
            ["Amber (prepare)",
             "Wind gusts of 80 km/h (50 mph) or more; ice accretion of 6 mm (0.25 in) or more; wet snow of 20 cm "
             "(8 in) or more; severe thunderstorm watch",
             "Hold crews on standby at shift end; top up fuel at staging areas; confirm contractor availability; "
             "confirm hotel blocks."],
            ["Red (pre-stage)",
             "Wind gusts of 100 km/h (62 mph) or more; ice accretion of 12 mm (0.5 in) or more; wet snow of 30 cm "
             "(12 in) or more; tornado watch",
             "Open staging areas; pre-position crews; place mutual assistance group on alert; activate the Regional "
             "EOC 6 hours before forecast onset."],
        ]),
        ("h1", "3. Storm Organization"),
        ("table", ["Position", "Responsibilities"], [
            ["Storm Coordinator Lead", "Overall resource strategy across regions, mutual assistance requests, "
                                       "readiness levels."],
            ["Regional Storm Lead", "Runs restoration in one region, assigns crews to work packages, owns the "
                                    "regional ETR."],
            ["Logistics Chief", "Staging areas, lodging, meals, fuel, material, laundry and sanitation."],
            ["Damage Assessment Lead", "Assigns patrol teams, consolidates damage reports into work packages."],
            ["Safety Officer", "Daily safety briefings, incident investigation, work and rest compliance checks."],
            ["Mutual Assistance Liaison", "Receives incoming crews, runs onboarding, tracks arrival and release."],
        ]),
        ("h1", "4. Staging Areas by Region"),
        ("p", "Each region has a primary staging area and an alternate. The Logistics Chief opens the primary site "
              "unless it is unsafe, inaccessible or already at capacity. The active sites for the current storm "
              "season are confirmed in the seasonal crew briefing."),
        ("table", ["Region", "Primary staging area", "Address", "Capacity", "Alternate staging area"], STAGING),
        ("p", "Every staging area must provide: a laydown area for poles and transformers, a fuel point, a secure "
              "parking area for at least 20 bucket trucks, portable lighting, sanitation, a check-in tent, and a "
              "meal service point able to serve a crew within 30 minutes of arrival."),
        ("h1", "5. Mutual Assistance"),
        ("h2", "5.1 Membership"),
        ("p", "Harbourline is a member of the Lakes and Northeast Mutual Assistance Group (LNMAG), which includes "
              "utilities in Ontario, Quebec, New York, Pennsylvania, Ohio and Michigan. Requests and offers are made "
              "on the LNMAG conference calls, which are held at 10:00 and 16:00 Eastern during an active event."),
        ("h2", "5.2 When to request"),
        ("bullets", [
            "An L4 Emergency is declared, or",
            "The projected restoration time exceeds 48 hours for more than 25 percent of affected customers, or",
            "Available Harbourline and contractor crews are fully committed and the regional ETR exceeds 72 hours.",
        ]),
        ("h2", "5.3 Request process"),
        ("steps", [
            "The Storm Coordinator Lead prepares a request stating the number of line crews, forestry crews and "
            "damage assessors, the reporting staging area, and the expected duration.",
            "The VP Operations approves requests for more than 20 crews or with an estimated cost over CAD 500,000.",
            "The Storm Coordinator Lead presents the request on the next LNMAG call and confirms offers in writing.",
            "The Mutual Assistance Liaison sends each responding utility the staging address, check-in time, lodging "
            "details and the onboarding package.",
            "For crews crossing the Canada and US border, the Liaison sends the border documentation packet at least "
            "12 hours before departure and notifies the Harbourline customs broker.",
        ]),
        ("h2", "5.4 Onboarding and release"),
        ("p", "Incoming crews complete a 90 minute onboarding at the staging area: Harbourline switching and lockout "
              "rules, system voltages and grounding practices, radio channels, and the work and rest rules in Section "
              "6. Crews receive a Harbourline escort (bird dog) for every 4 crews. Crews are released in order of "
              "home travel distance, longest distance first, unless the responding utility recalls them earlier. "
              "Costs are billed at actual cost under the LNMAG agreement and must be supported by daily time sheets."),
        ("h1", "6. Crew Work and Rest Rules"),
        ("p", "Fatigue is a leading contributor to storm injuries. These rules apply to Harbourline employees, "
              "contractors and mutual assistance crews alike."),
        ("table", ["Rule", "Requirement"], [
            ["Maximum continuous work", "16 consecutive hours, after which the worker must take at least 8 "
                                        "consecutive hours of rest."],
            ["Maximum in any 24 hours", "16 hours worked in any rolling 24 hour period."],
            ["Travel time", "Travel to and from lodging and staging counts as work time."],
            ["Extension", "Up to 18 hours only to complete a task that removes an immediate public safety hazard, "
                          "approved by the Regional Storm Lead and logged with the reason."],
            ["Extended events", "After 14 consecutive days of storm duty, a mandatory 24 hour rest period."],
            ["Meals", "A meal break at least every 5 hours of work."],
            ["Fatigue check", "The crew leader completes the fatigue checklist at the start of each shift and "
                              "after any extension."],
        ]),
        ("h1", "7. Damage Assessment"),
        ("p", "Damage assessment patrols start as soon as winds drop below 60 km/h (37 mph) or the Safety Officer "
              "declares conditions safe. Target: 100 percent of affected three phase mainline patrolled within 12 "
              "hours and 100 percent of laterals within 36 hours. Assessors record broken poles, wires down, "
              "damaged transformers and trees on lines in the mobile damage assessment form, with a photo and GPS "
              "point for every item."),
        ("h1", "8. Logistics"),
        ("bullets", [
            "Lodging: one bed per worker, within 60 minutes drive of the assigned staging area where possible.",
            "Meals: breakfast and dinner at the staging area, packed lunches issued at morning check-out.",
            "Fuel: staging area fuel points are topped up to at least 75 percent capacity at the start of each shift.",
            "Material: each staging area holds a storm kit of 30 poles, 20 pole-mount transformers, 5 km of "
            "conductor and hardware for 200 spans, replenished nightly from the regional warehouse.",
        ]),
        ("h1", "9. Safety During Restoration"),
        ("bullets", [
            "Treat every downed conductor as energized until tested and grounded.",
            "Watch for backfeed from customer generators; test before touch and ground on both sides of the work.",
            "All isolation for work follows SAF-MAN-002 Lockout and Tagout Safety Manual; storm conditions do not "
            "relax any lockout step.",
            "A daily tailboard briefing is held before crews leave the staging area and whenever the job changes.",
        ]),
        ("h1", "10. Storm Cost Tracking"),
        ("p", "The Storm Coordinator Lead opens a storm work order for each event in the format STM-YYYY-NNN, for "
              "example STM-2026-017. All labour, contractor, mutual assistance, material, fleet, lodging and meal "
              "costs are charged to that work order. Storm costs may be subject to recovery requests to the OEB, NY "
              "PSC or PUCO, so every charge must be supported by source documents."),
        ("h1", "11. Demobilization and After Action"),
        ("p", "Staging areas close when the region falls to L2 or lower and no crews remain assigned. The Logistics "
              "Chief returns unused storm kit material within 5 business days. The Storm Coordinator Lead issues the "
              "storm After Action Report within 30 calendar days, covering timeline, peak customers out, crews by "
              "source, injuries and near misses, ETR accuracy, and improvement actions."),
    ]
    revisions = [
        ["3", "2023-06-01", "D. Hutchins", "Added Ohio staging areas."],
        ["4", "2025-05-20", "D. Hutchins", "Added border documentation packet requirement."],
        ["5", "2026-05-15", "D. Hutchins", "Added Durham staging; clarified 18 hour extension approval."],
    ]
    build_pdf(path, "Storm Restoration Playbook", meta, blocks, revisions)


# ---------------------------------------------------------------------------
# 3. Lockout and Tagout Safety Manual (pdf)
# ---------------------------------------------------------------------------

MAD_TABLE = [
    ["0.050 to 0.300", "Avoid contact", "Avoid contact"],
    ["0.301 to 0.750", "0.35 m (1 ft 2 in)", "1.0 m (3 ft 4 in)"],
    ["0.751 to 15.0", "0.70 m (2 ft 4 in)", "3.0 m (10 ft 0 in)"],
    ["15.1 to 36.0", "0.85 m (2 ft 10 in)", "3.0 m (10 ft 0 in)"],
    ["36.1 to 46.0", "1.00 m (3 ft 4 in)", "3.0 m (10 ft 0 in)"],
    ["46.1 to 72.5", "1.20 m (4 ft 0 in)", "3.3 m (10 ft 10 in)"],
    ["72.6 to 121", "1.40 m (4 ft 7 in)", "3.8 m (12 ft 6 in)"],
    ["138 to 145", "1.60 m (5 ft 3 in)", "4.0 m (13 ft 2 in)"],
    ["161 to 169", "1.80 m (5 ft 11 in)", "4.2 m (13 ft 9 in)"],
    ["230 to 242", "2.40 m (7 ft 11 in)", "4.9 m (16 ft 1 in)"],
    ["345 to 362", "3.80 m (12 ft 6 in)", "6.0 m (19 ft 8 in)"],
    ["500 to 550", "5.60 m (18 ft 5 in)", "7.6 m (25 ft 0 in)"],
]


def loto_manual(path):
    meta = {
        "Document ID": "SAF-MAN-002",
        "Revision": "9",
        "Effective date": "January 15, 2026",
        "Next review": "January 15, 2027",
        "Document owner": "Aisha Coombs, Manager, Health and Safety",
        "Approved by": "Graham Lindqvist, Vice President, Operations",
        "Applies to": "All employees and contractors who isolate, work on, or work near electrical and mechanical "
                      "equipment owned or operated by Harbourline",
    }
    blocks = [
        ("h1", "1. Purpose"),
        ("p", "This manual sets the minimum requirements for isolating energy sources, applying locks and tags, and "
              "verifying a zero energy state before work, so that no worker is exposed to the unexpected "
              "energization or release of stored energy. It also sets Harbourline's minimum approach distances to "
              "exposed energized parts."),
        ("h1", "2. Scope and Regulatory Basis"),
        ("p", "The manual applies to all generation, transmission, substation, distribution and facility equipment, "
              "including mechanical, hydraulic, pneumatic and thermal energy sources. It implements the Ontario "
              "Occupational Health and Safety Act and its regulations for electrical utility work in Ontario, and "
              "OSHA 29 CFR 1910.147 and 1910.269 in New York and Ohio. Where these differ, the more protective "
              "requirement applies in all three jurisdictions."),
        ("h1", "3. Definitions"),
        ("bullets", [
            "Authorized Operator: a qualified person designated by the Distribution System Operator or station "
            "operations to perform isolation switching and issue permits.",
            "Permit Holder: the qualified worker in charge of the work who accepts the permit and holds it until the "
            "work is complete.",
            "Isolation point: a device that physically prevents the transmission or release of energy, such as a "
            "disconnect switch, breaker racked out, valve, or removed fuse.",
            "Zero energy state: the condition in which all energy sources are isolated, stored energy is released or "
            "restrained, and absence of voltage is verified.",
            "Working grounds: temporary protective grounds applied to create an equipotential zone at the work location.",
        ]),
        ("h1", "4. Responsibilities"),
        ("table", ["Role", "Responsibilities"], [
            ["Authorized Operator", "Prepares the isolation list, performs switching, applies operations locks and "
                                    "tags, and issues the permit."],
            ["Permit Holder", "Verifies isolation, applies working grounds, briefs the crew, controls the group "
                              "lockbox, and surrenders the permit."],
            ["Each worker", "Applies a personal lock to the group lockbox before starting work and removes only "
                            "their own lock."],
            ["Supervisor", "Ensures workers are trained, audits permits, and authorizes absent worker lock removal "
                           "under Section 11."],
            ["Health and Safety", "Maintains this manual, the training program, and the annual audit schedule."],
        ]),
        ("h1", "5. Lockout Devices and Tags"),
        ("table", ["Device", "Colour", "Who applies it"], [
            ["Operations lock", "Yellow", "Authorized Operator, at each isolation point"],
            ["Personal lock", "Red", "Each worker, on the group lockbox or directly at the isolation point"],
            ["Group lockbox", "Blue", "Permit Holder, holds the keys to all operations locks"],
            ["Contractor personal lock", "Orange", "Contractor workers, issued at site orientation"],
            ["Danger tag", "Red and white", "Attached to every operations lock and personal lock"],
        ]),
        ("p", "Each personal lock is keyed individually, marked with the worker's name and employee number, and has "
              "one key held only by that worker. Tags must show the permit number, the name of the person who "
              "applied it, the date and time, and the reason for isolation."),
        ("h1", "6. Permit System"),
        ("h2", "6.1 Permit numbering format"),
        ("p", "Every lockout is controlled by a Work Protection Permit. Permit numbers use the format "
              "LOTO-RRR-YYYY-NNNNN, where RRR is the three letter region code, YYYY is the year of issue, and NNNNN "
              "is a five digit sequence number that restarts at 00001 on January 1 in each region. Example: "
              "LOTO-EON-2026-00417 is the 417th permit issued in Eastern Ontario in 2026."),
        ("table", ["Region code", "Region"], [
            ["EON", "Eastern Ontario"], ["CON", "Central Ontario"], ["DUR", "Durham"],
            ["NYN", "New York North Country"], ["NYW", "Western New York"],
            ["OHN", "Northeast Ohio"], ["OHC", "North Central Ohio"], ["TRN", "Transmission (all jurisdictions)"],
        ]),
        ("h2", "6.2 Permit validity"),
        ("bullets", [
            "A permit is valid until the work is complete and the permit is surrendered, up to a maximum of 7 "
            "calendar days. Work lasting longer requires a new permit with a new number.",
            "A permit may be transferred at shift change only under Section 9.",
            "Permit records are kept in the Work Protection register for 7 years.",
        ]),
        ("h1", "7. Lockout Procedure"),
        ("steps", [
            "Prepare: the Authorized Operator identifies every energy source using the single line diagram and the "
            "equipment isolation list, and records each isolation point on the permit.",
            "Notify: inform the DSO or station operator and all affected workers that the equipment will be isolated.",
            "Shut down: stop the equipment using normal operating controls.",
            "Isolate: open every isolation point. Rack out breakers, open disconnects and remove fuses where applicable.",
            "Lock and tag: apply a yellow operations lock and danger tag to each isolation point and place all "
            "operations lock keys in the blue group lockbox.",
            "Issue the permit: the Authorized Operator and the Permit Holder review the permit together and both sign it.",
            "Release stored energy: discharge capacitors, bleed hydraulic and pneumatic pressure, block springs and "
            "raised parts.",
            "Verify absence of voltage: test with a rated voltage detector, proving the detector on a known live "
            "source before and after the test.",
            "Apply working grounds: install temporary protective grounds on all sides of the work location.",
            "Apply personal locks: every worker applies a red personal lock to the group lockbox and signs the "
            "permit before starting work.",
        ]),
        ("h2", "7.1 Restoring equipment"),
        ("p", "When work is complete, each worker removes their own personal lock and signs off. The Permit Holder "
              "confirms tools and personnel are clear, removes working grounds, and surrenders the permit to the "
              "Authorized Operator, who removes the operations locks and returns the equipment to service."),
        ("h1", "8. Group Lockout"),
        ("p", "For crews of more than one worker, the group lockbox method is mandatory. The Permit Holder is the "
              "first to apply and the last to remove a personal lock. Contractors apply orange personal locks to the "
              "same group lockbox and are listed by name on the permit."),
        ("h1", "9. Shift Change Transfer"),
        ("p", "The outgoing and incoming Permit Holders review the isolation points and grounds in person or by "
              "recorded phone call with the Authorized Operator. The incoming Permit Holder applies a personal lock "
              "before the outgoing Permit Holder removes theirs, so the lockbox is never without a Permit Holder lock."),
        ("h1", "10. Minimum Approach Distances"),
        ("p", "No person may approach, or take any conductive object closer than, the distances below to exposed "
              "energized parts unless the part is isolated and grounded under this manual, or the worker is a "
              "qualified worker using approved insulated tools and cover up. Harbourline company values meet or "
              "exceed the applicable regulations. Voltages are nominal phase to phase."),
        ("table", ["Nominal voltage (kV)", "Qualified worker MAD", "Unqualified person MAD"], MAD_TABLE),
        ("note", "Common Harbourline distribution voltages: 4.16 kV, 8.32 kV, 12.47 kV (Ohio), 13.2 kV (New York), "
                 "13.8 kV, 27.6 kV and 44 kV (Ontario). Common transmission voltages: 115 kV and 230 kV."),
        ("h1", "11. Removal of a Lock by Someone Other Than the Owner"),
        ("steps", [
            "The supervisor makes at least three documented attempts to contact the lock owner, including a call to "
            "their personal phone.",
            "The supervisor confirms the owner is not at the work site.",
            "The Regional Operations Manager approves the removal in writing on the permit.",
            "The lock is cut in the presence of a second qualified worker.",
            "The owner is informed before returning to work and signs the permit record.",
        ]),
        ("h1", "12. Training and Audits"),
        ("bullets", [
            "Initial training before first isolation duty, and refresher training every 12 months.",
            "Supervisors audit at least 2 permits per crew per quarter in the field.",
            "Health and Safety audits each region's Work Protection register annually.",
            "Any lockout failure or near miss is reported as a safety incident within 24 hours.",
        ]),
    ]
    revisions = [
        ["7", "2022-01-10", "A. Coombs", "Added contractor orange personal locks."],
        ["8", "2024-02-01", "A. Coombs", "Permit numbering changed to include three letter region code."],
        ["9", "2026-01-15", "A. Coombs", "Added Durham (DUR) region code; updated minimum approach distance table."],
    ]
    build_pdf(path, "Lockout and Tagout Safety Manual", meta, blocks, revisions)


# ---------------------------------------------------------------------------
# 4. Substation Switching Procedure (docx)
# ---------------------------------------------------------------------------

def substation_switching(path):
    meta = {
        "Document ID": "OPS-PRO-006",
        "Revision": "4",
        "Effective date": "June 1, 2025",
        "Next review": "June 1, 2027",
        "Document owner": "Victor Anand, Manager, System Control",
        "Approved by": "Renata Okafor, Director, Distribution Operations",
        "Applies to": "Distribution System Operators, station operators and qualified switching persons in all regions",
    }
    blocks = [
        ("h1", "1. Purpose"),
        ("p", "This procedure sets the rules for preparing, approving, executing and closing switching orders in "
              "Harbourline transformer stations, distribution substations and on distribution feeders, so that "
              "equipment is operated in the correct sequence and workers are protected."),
        ("h1", "2. Scope"),
        ("p", "It applies to all planned and emergency switching on equipment rated 4.16 kV and above that is under "
              "the control of a Harbourline control centre. Transmission switching ordered by the Independent "
              "Electricity System Operator (Ontario) or by the New York or PJM system operators is performed under "
              "those entities' instructions and recorded under this procedure."),
        ("h1", "3. Roles"),
        ("table", ["Role", "Responsibilities"], [
            ["Distribution System Operator (DSO)", "Writes or approves the switching order, issues each step, "
                                                   "records execution times, and holds control authority."],
            ["Qualified Switching Person (QSP)", "Executes each step in the field exactly as issued and repeats it "
                                                 "back before operating."],
            ["Independent Checker", "A second DSO who reviews every planned switching order before approval."],
            ["Station Operator", "Performs switching in staffed stations and verifies local indications."],
        ]),
        ("h1", "4. Switching Order Numbering"),
        ("p", "Switching orders are numbered SO-SSS-YYMMDD-NN, where SSS is the three letter station code, YYMMDD "
              "is the planned execution date, and NN is a two digit sequence for that station and day. Example: "
              "SO-KNG-261015-03 is the third switching order for Kingston TS on October 15, 2026."),
        ("table", ["Station code", "Station", "Voltage transformation"], [
            ["KNG", "Kingston TS", "115 kV to 44 kV and 27.6 kV"],
            ["NPN", "Napanee DS", "44 kV to 27.6 kV"],
            ["BAR", "Barrie TS", "230 kV to 44 kV"],
            ["ORL", "Orillia DS", "44 kV to 13.8 kV"],
            ["WHT", "Whitby TS", "230 kV to 27.6 kV"],
            ["WTN", "Watertown Substation", "115 kV to 13.2 kV"],
            ["JMS", "Jamestown Substation", "115 kV to 13.2 kV"],
            ["ASH", "Ashtabula Substation", "138 kV to 12.47 kV"],
            ["MNS", "Mansfield Substation", "138 kV to 12.47 kV"],
        ]),
        ("h1", "5. Preparing a Switching Order"),
        ("bullets", [
            "Planned switching orders are written from the current operating diagram and submitted at least 3 "
            "business days before execution.",
            "Every order lists each device by its unique operating designation (for example, KNG T2 LV breaker "
            "52T2L), the operation (open, close, rack out, lock, tag, ground), and the expected indication.",
            "The Independent Checker verifies the order against the operating diagram and signs it.",
            "Emergency switching orders may be written and executed without the 3 day lead time, but must still be "
            "reviewed by a second DSO before execution when one is available.",
        ]),
        ("h1", "6. Execution Rules"),
        ("h2", "6.1 Three-way communication"),
        ("p", "Every instruction uses three-way communication: the DSO states the step, the QSP repeats it back word "
              "for word, and the DSO confirms \"That is correct\" before the QSP operates. The NATO phonetic "
              "alphabet is used for letters in device designations."),
        ("h2", "6.2 One step at a time"),
        ("p", "Only one step is issued at a time. The QSP reports completion and the observed indication, and the "
              "DSO records the time before issuing the next step. Steps are never executed out of sequence."),
        ("h2", "6.3 Validity and interruptions"),
        ("bullets", [
            "A switching order is valid for 12 hours from the time the first step is issued. After 12 hours the "
            "remaining steps must be re-verified by the DSO against the operating diagram.",
            "If execution is interrupted for more than 30 minutes, the DSO and QSP re-verify the status of every "
            "device operated so far before continuing.",
            "If an indication does not match the expected indication, the QSP stops, makes the area safe and "
            "reports to the DSO. Switching resumes only after the DSO has resolved the discrepancy.",
        ]),
        ("h2", "6.4 Interlocks"),
        ("p", "Interlocks must never be bypassed without a written interlock bypass authorization from the Manager, "
              "System Control, recorded on the switching order."),
        ("h1", "7. Typical Sequence: Isolating a Station Power Transformer"),
        ("steps", [
            "Transfer load to the adjacent transformer by closing the bus tie breaker, confirming the load pick up.",
            "Open the transformer low voltage breaker and confirm open indication locally and on SCADA.",
            "Open the transformer high voltage circuit switcher and confirm open indication.",
            "Open and lock the high voltage and low voltage disconnect switches.",
            "Rack out the low voltage breaker to the disconnected position.",
            "Apply operations locks and danger tags under SAF-MAN-002 and issue the Work Protection Permit.",
            "Test for absence of voltage and apply working grounds on both sides of the transformer.",
        ]),
        ("h1", "8. Closing a Switching Order"),
        ("p", "The DSO closes the order when all steps are complete, the operating diagram is updated, and any "
              "permits have been surrendered. Completed orders are retained for 7 years."),
        ("h1", "9. Related Documents"),
        ("bullets", ["OPS-PRO-001 Outage Response Procedure", "SAF-MAN-002 Lockout and Tagout Safety Manual"]),
    ]
    revisions = [
        ["3", "2023-04-12", "V. Anand", "Added Ohio station codes."],
        ["4", "2025-06-01", "V. Anand", "Switching order validity changed from 8 hours to 12 hours."],
    ]
    build_docx(path, "Substation and Distribution Switching Procedure", meta, blocks, revisions)


# ---------------------------------------------------------------------------
# 5. Vegetation Management (docx)
# ---------------------------------------------------------------------------

def vegetation(path):
    meta = {
        "Document ID": "OPS-PRO-009",
        "Revision": "6",
        "Effective date": "April 1, 2026",
        "Next review": "April 1, 2027",
        "Document owner": "Colleen Yarrow, Manager, Forestry and Vegetation Management",
        "Approved by": "Renata Okafor, Director, Distribution Operations",
        "Applies to": "Forestry staff, line clearance contractors and planners in all regions",
    }
    blocks = [
        ("h1", "1. Purpose"),
        ("p", "Trees contacting power lines are the leading cause of outage minutes on the Harbourline distribution "
              "system. This procedure sets the trim cycles, clearance distances, hazard tree practices and customer "
              "notification rules that Harbourline and its line clearance contractors follow."),
        ("h1", "2. Scope"),
        ("p", "The procedure applies to all overhead transmission, sub-transmission and distribution lines in "
              "Ontario, New York and Ohio. Transmission lines operated at 200 kV and above are also managed to meet "
              "NERC FAC-003 requirements, which take precedence where they are more stringent."),
        ("h1", "3. Trim Cycles by Circuit Type"),
        ("p", "The trim cycle is the planned interval between full maintenance trims on the same circuit. Cycles are "
              "set by circuit type and area."),
        ("table", ["Circuit type", "Typical voltage", "Urban cycle", "Rural cycle", "Patrol"], [
            ["Transmission right of way", "115 kV, 138 kV, 230 kV", "5 years", "5 years",
             "Aerial patrol twice a year; LiDAR survey every 2 years"],
            ["Sub-transmission", "44 kV (ON), 34.5 kV (NY, OH)", "4 years", "4 years", "Ground patrol annually"],
            ["Three phase distribution mainline", "27.6 kV, 13.8 kV, 13.2 kV, 12.47 kV", "3 years", "4 years",
             "Ground patrol annually"],
            ["Single phase lateral", "16 kV, 8 kV, 7.6 kV, 7.2 kV", "5 years", "6 years",
             "Visual check at each trim"],
            ["Secondary and service drops", "120/240 V, 347/600 V", "On request", "On request",
             "Reactive only"],
        ]),
        ("h2", "3.1 Cycle acceleration"),
        ("p", "Any feeder in the worst performing 5 percent by tree caused customer interruptions over the last 3 "
              "years is trimmed in the next work plan year regardless of its cycle position."),
        ("h1", "4. Clearance Distances at Time of Trim"),
        ("p", "Clearances are measured from the nearest conductor at maximum sag. The side clearance applies "
              "horizontally; overhang is removed above the conductor to the height shown."),
        ("table", ["Circuit type", "Side clearance", "Below conductor", "Overhang removal above conductor"], [
            ["Transmission 230 kV", "9.0 m (30 ft)", "Full ROW floor clearing", "Remove all overhang"],
            ["Transmission 115 kV and 138 kV", "7.5 m (25 ft)", "Full ROW floor clearing", "Remove all overhang"],
            ["Sub-transmission 34.5 kV and 44 kV", "4.5 m (15 ft)", "3.0 m (10 ft)", "Remove all overhang"],
            ["Three phase mainline", "3.0 m (10 ft)", "3.0 m (10 ft)", "4.5 m (15 ft)"],
            ["Single phase lateral", "2.5 m (8 ft)", "2.5 m (8 ft)", "3.0 m (10 ft)"],
            ["Secondary and services", "0.5 m (1.5 ft) and no strain on the conductor", "Not specified",
             "Not specified"],
        ]),
        ("h1", "5. Hazard Trees"),
        ("p", "A hazard tree is a dead, dying, leaning, cracked or otherwise defective tree, inside or outside the "
              "trim zone, that could strike a conductor if it or part of it fails. Crews record hazard trees in the "
              "mobile vegetation application with a photo and a risk score from 1 (low) to 5 (imminent)."),
        ("bullets", [
            "Risk score 5: removed or made safe within 48 hours; the DSO is notified immediately.",
            "Risk score 4: removed within 30 calendar days.",
            "Risk score 3 or lower: scheduled with the next cycle trim.",
            "Ash trees killed by emerald ash borer within striking distance of three phase lines are treated as risk "
            "score 4.",
        ]),
        ("h1", "6. Customer and Landowner Notification"),
        ("table", ["Jurisdiction", "Notice before routine work", "Method"], [
            ["Ontario", "At least 14 calendar days", "Door hanger and letter"],
            ["New York", "At least 10 calendar days", "Door hanger; letter to absentee owners"],
            ["Ohio", "At least 7 calendar days", "Door hanger and phone call"],
        ]),
        ("p", "No advance notice is required for emergency work to remove an immediate hazard or restore service. "
              "Landowners who refuse access are referred to the Forestry Supervisor, who documents the refusal and "
              "any resulting reliability risk."),
        ("h1", "7. Work Methods"),
        ("bullets", [
            "Directional pruning to the branch collar following ANSI A300 practices; no topping or rounding over.",
            "Herbicide use on rights of way only by licensed applicators and only with products approved by the "
            "applicable provincial or state authority.",
            "Only line clearance qualified arborists work within 3.0 m (10 ft) of energized distribution conductors.",
            "Wood chips are left on site only with the owner's consent.",
        ]),
        ("h1", "8. Performance Measures"),
        ("table", ["Measure", "2026 target"], [
            ["Tree caused customer interruptions per 100 km of overhead line", "18 or fewer"],
            ["Planned trim kilometres completed", "Ontario 4,200 km; New York 1,150 km; Ohio 980 km"],
            ["Risk score 5 hazard trees made safe within 48 hours", "100 percent"],
        ]),
    ]
    revisions = [
        ["5", "2024-03-15", "C. Yarrow", "Added emerald ash borer guidance."],
        ["6", "2026-04-01", "C. Yarrow", "Rural lateral cycle extended from 5 to 6 years; updated 2026 targets."],
    ]
    build_docx(path, "Vegetation Management Procedure", meta, blocks, revisions)


# ---------------------------------------------------------------------------
# 6. Crew Briefing Deck (pptx) with image-only key values
# ---------------------------------------------------------------------------

# This season's staging areas (differ from the playbook for two regions).
SEASON_STAGING = [
    ("Eastern Ontario", "Napanee Fairgrounds, 4 York Road East", "Kingston yard closed for repaving"),
    ("Central Ontario", "Barrie Operations Yard, 85 Welham Road", "Primary"),
    ("Durham", "Whitby Service Centre, 1600 Thickson Road North", "Primary"),
    ("NY North Country", "Watertown Service Center, 310 Coffeen Street", "Primary"),
    ("Western NY", "Jamestown Operations Yard, 92 Harrison Street", "Primary"),
    ("Northeast Ohio", "Lake County Fairgrounds, Painesville", "Ashtabula yard shared with substation rebuild"),
    ("North Central Ohio", "Mansfield Operations Center, 700 Park Avenue West", "Primary"),
]
BRIDGE_NUMBER = "1-844-555-0162"
BRIDGE_PIN = "4471 902"
SAFETY_NUMBERS = [
    ("10 m (33 ft)", "Minimum public standoff from any downed conductor"),
    ("45 minutes", "Maximum interval between lone worker check-ins"),
    ("1-866-555-0199", "Harbourline Safety Incident Hotline, 24/7"),
]


def _font(size, bold=False):
    name = "DejaVuSans-Bold.ttf" if bold else "DejaVuSans.ttf"
    return ImageFont.truetype(os.path.join(FONT_DIR, name), size)


def _png(img):
    buf = io.BytesIO()
    img.save(buf, format="PNG", optimize=False)
    buf.seek(0)
    return buf


def img_staging():
    W, H = 1800, 900
    im = Image.new("RGB", (W, H), (255, 255, 255))
    d = ImageDraw.Draw(im)
    d.rectangle([0, 0, W, 90], fill=(11, 61, 92))
    d.text((30, 22), "2026 to 2027 WINTER SEASON: ACTIVE STAGING AREAS", font=_font(40, True), fill="white")
    cols = [30, 460, 1240]
    y = 115
    for x, h in zip(cols, ["Region", "Report to", "Note"]):
        d.text((x, y), h, font=_font(30, True), fill=(11, 61, 92))
    y += 50
    d.line([20, y, W - 20, y], fill=(11, 61, 92), width=3)
    y += 15
    for i, (reg, site, note) in enumerate(SEASON_STAGING):
        if i % 2 == 0:
            d.rectangle([20, y - 8, W - 20, y + 82], fill=(232, 240, 247))
        d.text((cols[0], y), reg, font=_font(28, True), fill=(20, 20, 20))
        d.text((cols[1], y), site, font=_font(25), fill=(20, 20, 20))
        d.text((cols[2], y + 2), note, font=_font(20), fill=(90, 90, 90))
        y += 95
    return _png(im)


def img_bridge():
    W, H = 1600, 700
    im = Image.new("RGB", (W, H), (255, 248, 225))
    d = ImageDraw.Draw(im)
    d.rectangle([10, 10, W - 10, H - 10], outline=(200, 120, 0), width=8)
    d.text((60, 50), "STORM CALL-OUT PHONE BRIDGE", font=_font(52, True), fill=(120, 60, 0))
    d.text((60, 170), "Dial:", font=_font(40), fill=(40, 40, 40))
    d.text((60, 225), BRIDGE_NUMBER, font=_font(110, True), fill=(10, 10, 10))
    d.text((60, 400), "Conference ID:", font=_font(40), fill=(40, 40, 40))
    d.text((60, 455), BRIDGE_PIN + " #", font=_font(80, True), fill=(10, 10, 10))
    d.text((60, 600), "Join within 30 minutes of the call-out text. State crew ID and ETA to staging.",
           font=_font(30), fill=(60, 60, 60))
    return _png(im)


def img_safety():
    W, H = 1800, 800
    im = Image.new("RGB", (W, H), (255, 255, 255))
    d = ImageDraw.Draw(im)
    d.text((40, 20), "THE THREE CRITICAL SAFETY NUMBERS", font=_font(48, True), fill=(170, 20, 20))
    bw = (W - 80 - 2 * 40) // 3
    for i, (num, label) in enumerate(SAFETY_NUMBERS):
        x0 = 40 + i * (bw + 40)
        d.rounded_rectangle([x0, 120, x0 + bw, 760], radius=30, fill=(170, 20, 20))
        d.text((x0 + 30, 150), str(i + 1), font=_font(70, True), fill=(255, 210, 210))
        f = _font(64 if len(num) < 14 else 50, True)
        d.text((x0 + 30, 300), num, font=f, fill="white")
        # wrap label
        words, lines, cur = label.split(), [], ""
        for w in words:
            t = (cur + " " + w).strip()
            if d.textlength(t, font=_font(34)) > bw - 60:
                lines.append(cur)
                cur = w
            else:
                cur = t
        lines.append(cur)
        for j, ln in enumerate(lines):
            d.text((x0 + 30, 460 + j * 48), ln, font=_font(34), fill="white")
    return _png(im)


def crew_deck(path):
    prs = Presentation()
    prs.slide_width = PInches(13.333)
    prs.slide_height = PInches(7.5)
    prs.core_properties.author = COMPANY
    prs.core_properties.title = "Winter Storm Season Crew Briefing 2026 to 2027"
    prs.core_properties.created = FIXED_DT
    prs.core_properties.modified = FIXED_DT
    prs.core_properties.last_modified_by = COMPANY
    blank = prs.slide_layouts[6]
    title_layout = prs.slide_layouts[0]

    def add_title(slide, text):
        tb = slide.shapes.add_textbox(PInches(0.5), PInches(0.3), PInches(12.3), PInches(0.9))
        p = tb.text_frame.paragraphs[0]
        p.text = text
        p.runs[0].font.size = PPt(32)
        p.runs[0].font.bold = True
        p.runs[0].font.color.rgb = PptxRGB(0x0B, 0x3D, 0x5C)
        tb.name = "Title"

    def add_bullets(slide, items, top=1.4, height=5.5, size=20):
        tb = slide.shapes.add_textbox(PInches(0.7), PInches(top), PInches(12.0), PInches(height))
        tf = tb.text_frame
        tf.word_wrap = True
        for i, it in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = "• " + it
            p.runs[0].font.size = PPt(size)
            p.space_after = PPt(8)

    def notes(slide, text):
        slide.notes_slide.notes_text_frame.text = text

    def add_pic(slide, buf, left, top, width, name):
        pic = slide.shapes.add_picture(buf, PInches(left), PInches(top), width=PInches(width))
        pic.name = name
        # neutral alt text that does not contain key values
        pic._element.nvPicPr.cNvPr.set("descr", name)
        return pic

    # 1 Title
    s = prs.slides.add_slide(title_layout)
    s.shapes.title.text = "Winter Storm Season Crew Briefing 2026 to 2027"
    s.placeholders[1].text = "Distribution Operations, all regions\nPresented by Dale Hutchins, Storm Coordinator Lead"
    for ph, top, h in ((s.shapes.title, 2.2, 1.5), (s.placeholders[1], 3.9, 1.4)):
        ph.left, ph.top, ph.width, ph.height = PInches(1.0), PInches(top), PInches(11.33), PInches(h)
    notes(s, "Welcome crews. This briefing is mandatory for all line, forestry and damage assessment staff before "
             "November 1. Attendance is recorded in the learning system.")

    # 2 Agenda
    s = prs.slides.add_slide(blank)
    add_title(s, "Agenda")
    add_bullets(s, ["Season at a glance", "Where to report: this season's staging areas", "How call-outs work",
                    "The three critical safety numbers", "Work and rest rules", "Restoration priorities",
                    "Mutual assistance crews", "Questions"])
    notes(s, "Keep this under 40 minutes. Leave time for questions at the end.")

    # 3 Season at a glance
    s = prs.slides.add_slide(blank)
    add_title(s, "Season at a glance")
    add_bullets(s, ["Winter storm season runs November 1, 2026 to March 31, 2027",
                    "Forecast review twice daily; readiness levels Green, Amber, Red",
                    "Last season: 11 storm events, 3 reached L3 Major, 1 reached L4 Emergency",
                    "Focus this year: fewer strain injuries, faster damage assessment"])
    notes(s, "Remind crews that readiness levels come from the Storm Restoration Playbook. Last season's L4 event "
             "was the February ice storm in Eastern Ontario.")

    # 4 Staging (image only)
    s = prs.slides.add_slide(blank)
    add_title(s, "Where to report this season")
    add_pic(s, img_staging(), 0.9, 1.3, 11.5, "Staging area table")
    notes(s, "Walk through the table on the slide region by region. Two regions are not using their usual yard "
             "this season, so point those out and make sure crews in those regions know where to go.")

    # 5 Call-out (image only for number)
    s = prs.slides.add_slide(blank)
    add_title(s, "How call-outs work")
    add_bullets(s, ["You receive a call-out text from the storm dispatch system",
                    "Join the phone bridge shown on the right",
                    "Confirm crew ID, members and ETA to staging",
                    "No answer within 30 minutes: dispatch moves to the next crew on the list"],
                top=1.5, height=5, size=18)
    for shp in list(s.shapes):
        if shp.name != "Title" and shp.shape_type is not None and shp.has_text_frame:
            shp.width = PInches(5.6)
    add_pic(s, img_bridge(), 6.6, 1.6, 6.3, "Call-out bridge card")
    notes(s, "Everyone should save the bridge details from this slide in their phone today. Do not rely on "
             "searching for it during an event.")

    # 6 Safety numbers (image only)
    s = prs.slides.add_slide(blank)
    add_title(s, "The three critical safety numbers")
    add_pic(s, img_safety(), 0.9, 1.4, 11.5, "Critical safety numbers graphic")
    notes(s, "Ask each crew leader to repeat the three numbers back. These are checked at every storm tailboard.")

    # 7 Work and rest
    s = prs.slides.add_slide(blank)
    add_title(s, "Work and rest rules")
    add_bullets(s, ["Maximum 16 consecutive hours, then at least 8 consecutive hours of rest",
                    "Travel to and from lodging counts as work time",
                    "Extension to 18 hours only to remove an immediate public safety hazard, approved and logged",
                    "Meal break at least every 5 hours",
                    "Fatigue checklist at the start of every shift"])
    notes(s, "These rules match the Storm Restoration Playbook, section 6. They apply to contractors and mutual "
             "assistance crews too.")

    # 8 Restoration priorities
    s = prs.slides.add_slide(blank)
    add_title(s, "Restoration priorities")
    add_bullets(s, ["1. Public safety hazards and wires down",
                    "2. Hospitals and long term care homes",
                    "3. Water and wastewater treatment and pumping",
                    "4. 911 centres, police, fire and EMS",
                    "5. Life support customers",
                    "6. Stations and feeder mainlines, then taps, then individual services"], size=18)
    notes(s, "From the Outage Response Procedure, section 7. The Storm Coordinator sets the sequence within "
             "this order.")

    # 9 Mutual assistance
    s = prs.slides.add_slide(blank)
    add_title(s, "Working with mutual assistance crews")
    add_bullets(s, ["Incoming crews complete a 90 minute onboarding at staging",
                    "One Harbourline escort for every 4 incoming crews",
                    "Share our grounding and lockout rules at every tailboard",
                    "Released longest travel distance first"])
    notes(s, "Be welcoming. Most incoming crews have not worked on 27.6 kV or 44 kV systems before.")

    # 10 Questions
    s = prs.slides.add_slide(blank)
    add_title(s, "Questions")
    add_bullets(s, ["Contact your Regional Storm Lead",
                    "Playbook and procedures are on the Harbourline-Operations site, Procedures library"])
    notes(s, "Close by reminding crews to check their call-out contact details in the storm dispatch system.")

    # The default python-pptx master uses an en dash bullet glyph; replace it (no dashes rule).
    for master in prs.slide_masters:
        for el in master._element.iter():
            if el.tag.endswith("}buChar") and el.get("char") in ("\u2013", "\u2014"):
                el.set("char", "\u2022")
        for layout in master.slide_layouts:
            for el in layout._element.iter():
                if el.tag.endswith("}buChar") and el.get("char") in ("\u2013", "\u2014"):
                    el.set("char", "\u2022")
    prs.save(path)


# ---------------------------------------------------------------------------
# 7. Transformer Equipment Specs (xlsx)
# ---------------------------------------------------------------------------

MFRS = {
    "NF": "Northfield Electric",
    "CG": "Castlegate Power Products",
    "LTW": "Lakeview Transformer Works",
    "BRN": "Brennock Industries",
}
DEPOTS = ["Kingston ON", "Barrie ON", "Whitby ON", "Watertown NY", "Jamestown NY", "Ashtabula OH", "Mansfield OH"]

THIN = Side(style="thin", color="7A8A99")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HDR_FILL = PatternFill("solid", fgColor="0B3D5C")
GRP_FILL = PatternFill("solid", fgColor="DCE6F0")
UNIT_FILL = PatternFill("solid", fgColor="F2F2F2")


def _sheet_frame(ws, title, groups, headers, units, widths):
    """Row 1 title (merged across), row 2 group headers (merged spans), row 3 headers, row 4 units."""
    ncol = len(headers)
    ws.cell(row=1, column=1, value=title)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    ws.cell(row=1, column=1).font = Font(bold=True, size=14, color="FFFFFF")
    ws.cell(row=1, column=1).fill = HDR_FILL
    ws.cell(row=1, column=1).alignment = Alignment(horizontal="center")
    col = 1
    for name, span in groups:
        ws.cell(row=2, column=col, value=name)
        if span > 1:
            ws.merge_cells(start_row=2, start_column=col, end_row=2, end_column=col + span - 1)
        c = ws.cell(row=2, column=col)
        c.font = Font(bold=True)
        c.fill = GRP_FILL
        c.alignment = Alignment(horizontal="center")
        col += span
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=3, column=i, value=h)
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="center", wrap_text=True)
        c.fill = GRP_FILL
    for i, u in enumerate(units, 1):
        c = ws.cell(row=4, column=i, value=u if u else None)
        c.font = Font(italic=True, size=9)
        c.fill = UNIT_FILL
        c.alignment = Alignment(horizontal="center")
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for r in range(1, 5):
        for i in range(1, ncol + 1):
            ws.cell(row=r, column=i).border = BORDER
    ws.freeze_panes = "C5"


def _write_grouped(ws, start_row, grouped_rows):
    """grouped_rows: list of (group_label, [row values excluding col A]). Column A merged per group."""
    r = start_row
    for label, rows in grouped_rows:
        first = r
        for vals in rows:
            for i, v in enumerate(vals, 2):
                c = ws.cell(row=r, column=i, value=v)
                c.border = BORDER
            ws.cell(row=r, column=1).border = BORDER
            r += 1
        ws.cell(row=first, column=1, value=label)
        if r - 1 > first:
            ws.merge_cells(start_row=first, start_column=1, end_row=r - 1, end_column=1)
        c = ws.cell(row=first, column=1)
        c.font = Font(bold=True)
        c.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
    return r


# Planted lookup facts
PLANT_POLE = ["NF-PM25-16B", "Northfield Electric", "1", 25, "27.6GrdY/16.0", "120/240", 1.9, 125, 68, 215,
              "Pole, direct hanger", "ON", "Approved"]
PLANT_PAD = ["CG-PAD500-3P-27", "Castlegate Power Products", "3", 500, "27.6GrdY/16.0", "347/600", 4.6, 150, 1135,
             2860, "Pad, loop feed, dead front", "ON", "Approved"]
PLANT_SUB = ["BRN-SPT-33-115", "Brennock Industries", "3", "20/26.7/33.3", "115", "27.6", 8.9, 550, 18400, 61200,
             "ONAN/ONAF/ONAF", "LTC, 33 steps", "ON", "Approved"]


def transformer_specs(path):
    rnd = random.Random(SEED)
    wb = Workbook()

    # ---- Sheet 1: Distribution Pole-Mount
    ws = wb.active
    ws.title = "Distribution Pole-Mount"
    headers = ["Voltage Class", "Model No.", "Manufacturer", "Phase", "Rating", "Primary Voltage",
               "Secondary Voltage", "Impedance", "BIL", "Oil Volume", "Weight", "Mounting",
               "Approved Regions", "Status"]
    units = ["", "", "", "", "kVA", "kV", "V", "%Z", "kV", "L", "kg", "", "", ""]
    _sheet_frame(ws, "Harbourline Energy Co. Distribution Pole-Mount Transformers: Approved Equipment List "
                     "(Rev 2026-07)",
                 [("Identification", 4), ("Electrical Ratings", 5), ("Physical", 3), ("Approval", 2)],
                 headers, units, [18, 18, 26, 7, 9, 16, 12, 10, 8, 11, 10, 22, 16, 12])
    classes = [
        ("Ontario 27.6 kV class", "27.6GrdY/16.0", "16", ["ON"], 125),
        ("Ontario 13.8 kV class", "13.8GrdY/8.0", "8", ["ON"], 95),
        ("New York 13.2 kV class", "13.2GrdY/7.62", "7N", ["NY"], 95),
        ("Ohio 12.47 kV class", "12.47GrdY/7.2", "7", ["OH"], 95),
    ]
    sizes = [10, 15, 25, 37.5, 50, 75, 100, 167]
    oil = {10: 42, 15: 52, 25: 68, 37.5: 90, 50: 110, 75: 150, 100: 185, 167: 280}
    wt = {10: 120, 15: 150, 25: 215, 37.5: 280, 50: 340, 75: 450, 100: 560, 167: 820}
    grouped = []
    for label, pv, code, regions, bil in classes:
        rows = []
        for sz in sizes:
            for mk in (["NF", "LTW"] if sz in (25, 50, 100) else ["NF"]):
                szs = str(sz).replace(".", "")
                model = f"{mk}-PM{szs}-{code}{'A' if mk == 'NF' else 'L'}"
                imp = round(1.5 + sz / 100 + rnd.uniform(0, 0.4), 1)
                rows.append([model, MFRS[mk], "1", sz, pv, "120/240", imp, bil,
                             oil[sz] + rnd.randint(-4, 4), wt[sz] + rnd.randint(-10, 10),
                             "Pole, direct hanger" if sz <= 50 else "Pole, crossarm bracket",
                             ", ".join(regions), "Approved" if rnd.random() > 0.12 else "Phase out 2027"])
            if label.startswith("Ontario 27.6") and sz == 25:
                rows.append(list(PLANT_POLE))
        grouped.append((label, rows))
    _write_grouped(ws, 5, grouped)

    # ---- Sheet 2: Pad-Mount
    ws = wb.create_sheet("Pad-Mount")
    _sheet_frame(ws, "Harbourline Energy Co. Pad-Mount Transformers: Approved Equipment List (Rev 2026-07)",
                 [("Identification", 4), ("Electrical Ratings", 5), ("Physical", 3), ("Approval", 2)],
                 headers, units, [20, 20, 26, 7, 9, 16, 12, 10, 8, 11, 10, 26, 16, 12])
    grouped = []
    pad_classes = [
        ("Single phase residential", "1", [25, 50, 75, 100, 167], ["120/240"]),
        ("Three phase commercial, Ontario", "3", [75, 150, 300, 500, 750, 1000, 1500, 2500], ["347/600", "120/208"]),
        ("Three phase commercial, US", "3", [75, 150, 300, 500, 750, 1000, 1500], ["277/480", "120/208"]),
    ]
    for label, ph, szs, secs in pad_classes:
        rows = []
        for sz in szs:
            for sec in secs:
                for pv, reg, code in ([("27.6GrdY/16.0", "ON", "27"), ("13.8GrdY/8.0", "ON", "13")]
                                      if "Ontario" in label or ph == "1" else
                                      [("13.2GrdY/7.62", "NY", "13N"), ("12.47GrdY/7.2", "OH", "12")]):
                    if ph == "1" and code not in ("27", "12"):
                        continue
                    if ph == "3" and sec in ("120/208",) and sz > 750:
                        continue
                    mk = "CG" if sz >= 300 else "LTW"
                    model = f"{mk}-PAD{sz}-{ph}P-{code}{'' if sec in ('120/240', '347/600', '277/480') else 'W'}"
                    if model == PLANT_PAD[0]:
                        continue
                    base_oil = int(sz * 1.9 + 120)
                    rows.append([model, MFRS[mk], ph, sz, pv, sec,
                                 round(1.8 + min(sz, 2500) / 400 + rnd.uniform(0, 0.5), 2 if ph == "3" else 1),
                                 150 if code == "27" else 95, base_oil + rnd.randint(-15, 15),
                                 int(base_oil * 2.6) + rnd.randint(-30, 30),
                                 "Pad, loop feed, dead front" if sz <= 1000 else "Pad, radial feed, dead front",
                                 reg, "Approved" if rnd.random() > 0.1 else "Under evaluation"])
                    if model == "CG-PAD500-3P-13" and label.endswith("Ontario"):
                        rows.append(list(PLANT_PAD))
        grouped.append((label, rows))
    _write_grouped(ws, 5, grouped)

    # ---- Sheet 3: Substation Power
    ws = wb.create_sheet("Substation Power")
    sheaders = ["Station Class", "Model No.", "Manufacturer", "Phase", "Rating", "HV Winding", "LV Winding",
                "Impedance", "BIL (HV)", "Oil Volume", "Total Weight", "Cooling", "Tap Changer",
                "Approved Regions", "Status"]
    sunits = ["", "", "", "", "MVA", "kV", "kV", "%Z", "kV", "L", "kg", "", "", "", ""]
    _sheet_frame(ws, "Harbourline Energy Co. Substation Power Transformers: Approved Equipment List (Rev 2026-07)",
                 [("Identification", 4), ("Electrical Ratings", 5), ("Physical", 4), ("Approval", 2)],
                 sheaders, sunits, [22, 20, 26, 7, 14, 10, 10, 10, 9, 11, 12, 18, 16, 16, 12])
    grouped = []
    sub_classes = [
        ("Ontario transformer stations", [("230", "44"), ("230", "27.6"), ("115", "44"), ("115", "27.6")], "ON"),
        ("Ontario distribution stations", [("44", "27.6"), ("44", "13.8"), ("44", "4.16")], "ON"),
        ("New York substations", [("115", "13.2"), ("115", "34.5"), ("34.5", "13.2")], "NY"),
        ("Ohio substations", [("138", "12.47"), ("138", "34.5"), ("69", "12.47")], "OH"),
    ]
    ratings = [("10/13.3/16.7", 10), ("15/20/25", 15), ("20/26.7/33.3", 20), ("50/66.7/83.3", 50)]
    for label, pairs, reg in sub_classes:
        rows = []
        for hv, lv in pairs:
            for rt, base in ratings:
                hvf = float(hv)
                if hvf <= 44 and base > 20:
                    continue
                if hvf >= 230 and base < 20:
                    continue
                mk = "BRN" if base >= 20 else rnd.choice(["BRN", "LTW"])
                model = f"{mk}-SPT-{int(base * 5 / 3)}-{hv.replace('.', '')}{lv.replace('.', '')[:3]}"
                bil = {230: 900, 138: 650, 115: 550, 69: 350, 44: 250, 34.5: 200}[hvf]
                oil_l = int(base * 780 + hvf * 12 + rnd.randint(-300, 300))
                rows.append([model, MFRS[mk], "3", rt, hv, lv, round(7.0 + base / 20 + rnd.uniform(0, 1.2), 1),
                             bil, oil_l, int(oil_l * 3.3 + rnd.randint(-500, 500)),
                             "ONAN/ONAF/ONAF", "LTC, 33 steps" if hvf >= 69 else "DETC, 5 positions",
                             reg, "Approved"])
                if label.startswith("Ontario transformer") and hv == "115" and lv == "27.6" and base == 20:
                    rows.append(list(PLANT_SUB))
        grouped.append((label, rows))
    _write_grouped(ws, 5, grouped)

    # ---- Sheet 4: Spares Inventory
    ws = wb.create_sheet("Spares Inventory")
    iheaders = ["Category", "Model No."]
    iunits = ["", ""]
    groups = [("Item", 2)]
    for dep in DEPOTS:
        iheaders += ["On Hand", "Reorder Point"]
        iunits += ["units", "units"]
        groups.append((dep, 2))
    iheaders += ["Total On Hand"]
    iunits += ["units"]
    groups.append(("System", 1))
    _sheet_frame(ws, "Harbourline Energy Co. Transformer Spares Inventory by Depot (as of September 1, 2026)",
                 groups, iheaders, iunits, [20, 20] + [9] * (2 * len(DEPOTS)) + [11])
    inv_models = [
        ("Pole-mount", [r[0] for r in wb["Distribution Pole-Mount"].iter_rows(min_row=5, min_col=2, max_col=2,
                                                                               values_only=True)]),
        ("Pad-mount", [r[0] for r in wb["Pad-Mount"].iter_rows(min_row=5, min_col=2, max_col=2, values_only=True)]),
        ("Substation power", [r[0] for r in wb["Substation Power"].iter_rows(min_row=5, min_col=2, max_col=2,
                                                                             values_only=True)]),
    ]
    region_of = {}
    for sheet in ("Distribution Pole-Mount", "Pad-Mount", "Substation Power"):
        sh = wb[sheet]
        rcol = 13 if sheet != "Substation Power" else 14
        for row in sh.iter_rows(min_row=5, values_only=True):
            region_of[row[1]] = row[rcol - 1]
    dep_region = ["ON", "ON", "ON", "NY", "NY", "OH", "OH"]
    grouped = []
    for cat, models in inv_models:
        chosen = [m for m in models if m]
        # keep a stable subset so the sheet stays in the 30 to 80 row range
        step = 2 if cat != "Substation power" else 3
        subset = chosen[::step]
        for planted in (PLANT_POLE[0], PLANT_PAD[0], PLANT_SUB[0]):
            if planted in chosen and planted not in subset:
                subset.append(planted)
        rows = []
        for m in subset:
            vals = [m]
            total = 0
            for di, dep in enumerate(DEPOTS):
                if region_of.get(m) != dep_region[di]:
                    oh, rp = 0, 0
                elif cat == "Substation power":
                    oh, rp = rnd.choice([0, 0, 1]), 0
                else:
                    oh, rp = rnd.randint(0, 14), rnd.choice([2, 3, 4, 5])
                if m == PLANT_PAD[0]:
                    oh, rp = {"Kingston ON": (5, 2), "Barrie ON": (3, 2), "Whitby ON": (1, 2)}.get(dep, (0, 0))
                if m == PLANT_POLE[0]:
                    oh, rp = {"Kingston ON": (12, 6), "Barrie ON": (9, 6), "Whitby ON": (4, 6)}.get(dep, (0, 0))
                if m == PLANT_SUB[0]:
                    oh, rp = {"Kingston ON": (1, 1)}.get(dep, (0, 0))
                vals += [oh, rp]
                total += oh
            vals.append(total)
            rows.append(vals)
        grouped.append((cat, rows))
    _write_grouped(ws, 5, grouped)

    wb.properties.creator = COMPANY
    wb.properties.created = FIXED_DT
    wb.properties.modified = FIXED_DT
    wb.save(path)


# ---------------------------------------------------------------------------
# 8. Lab 9 field report drops (docx)
# ---------------------------------------------------------------------------

FIELD_REPORTS = [
    {
        "file": "Field-Report-2026-10-01-Kingston.docx",
        "Report ID": "FR-EON-2026-1187",
        "Date": "2026-10-01",
        "Region": "Eastern Ontario",
        "Crew ID": "CREW-ON-03",
        "Crew leader": "Devon Achebe",
        "Asset ID": "TX-ON-10423",
        "Asset description": "25 kVA pole-mount transformer, model NF-PM25-16B, pole P-30912, Bath Road near "
                             "Collins Bay Road, Kingston ON",
        "Issue": "Oil weeping from the tank lid gasket, dark stain on the pole below the tank. No customer outage.",
        "Severity": "Moderate",
        "Recommended action": "Schedule planned replacement within 10 business days and add spill tray inspection "
                              "to the work order.",
        "narrative": "Found during routine pole inspection patrol on feeder KNG-M4. Tank temperature normal by "
                     "infrared scan. Oil level sight glass reads low but visible. Customers on the transformer "
                     "were not interrupted. Area around the pole base is gravel, no watercourse within 50 m.",
    },
    {
        "file": "Field-Report-2026-10-02-Watertown.docx",
        "Report ID": "FR-NYN-2026-0442",
        "Date": "2026-10-02",
        "Region": "New York North Country",
        "Crew ID": "CREW-NY-13",
        "Crew leader": "Colleen Brady",
        "Asset ID": "",
        "Asset description": "Pole-mount transformer on the east side of Arsenal Street near the Coffeen Street "
                             "intersection, Watertown NY. Asset tag missing from the tank.",
        "Issue": "Squirrel guard missing on the primary bushing; evidence of animal contact and a blown cutout fuse "
                 "that was refused once overnight by the trouble crew.",
        "Severity": "Low",
        "Recommended action": "Install a new wildlife guard and replace the asset tag at the next planned visit; "
                              "records team to confirm the asset ID from GIS.",
        "narrative": "Trouble crew restored service to 9 customers at 02:40 by replacing the cutout fuse. No asset "
                     "tag could be read on the tank, so the asset ID field is left blank for records to complete.",
    },
    {
        "file": "Field-Report-2026-10-03-Ashtabula.docx",
        "Report ID": "FR-OHN-2026-0918",
        "Date": "2026-10-03",
        "Region": "Northeast Ohio",
        "Crew ID": "CREW-OH-06",
        "Crew leader": "Hannah Voss",
        "Asset ID": "TX-OH-20871",
        "Asset description": "50 kVA pole-mount transformer, pole P-51260, Lake Avenue at West 5th Street, "
                             "Ashtabula OH",
        "Issue": "Primary bushing flashover with fire damage to the tank and crossarm; pole cracked at the "
                 "transformer bracket and leaning toward the roadway. 64 customers out.",
        "Severity": "Critical",
        "Recommended action": "Immediate replacement of pole and transformer. Section isolated and locked out under "
                              "permit LOTO-OHN-2026-00388. Request oil spill response for approximately 40 L "
                              "released to soil.",
        "narrative": "Fire department on scene at arrival, fire out. Police holding traffic on Lake Avenue. Line "
                     "crew and a second crew with a digger derrick requested. ETR for customers is 18:30.",
    },
]


def field_reports(outdir):
    for fr in FIELD_REPORTS:
        doc = _docx_base()
        doc.core_properties.title = f"Field Report {fr['Report ID']}"
        p = doc.add_paragraph()
        r = p.add_run(COMPANY)
        r.bold = True
        doc.add_heading("Field Report", 0)
        fields = ["Report ID", "Date", "Region", "Crew ID", "Crew leader", "Asset ID", "Asset description",
                  "Issue", "Severity", "Recommended action"]
        _docx_table(doc, ["Field", "Value"], [[f, fr[f]] for f in fields], widths=[1.8, 4.7])
        doc.add_heading("Crew notes", 1)
        doc.add_paragraph(fr["narrative"])
        doc.add_heading("Submission", 1)
        doc.add_paragraph("Submitted from the mobile workforce application to the Procedures/Incoming folder for "
                          "triage by the Operations Triage agent.")
        doc.save(os.path.join(outdir, fr["file"]))


# ---------------------------------------------------------------------------

def main():
    ap = argparse.ArgumentParser()
    here = os.path.dirname(os.path.abspath(__file__))
    ap.add_argument("--root", default=os.path.abspath(os.path.join(here, "..", "..")))
    a = ap.parse_args()
    proc = os.path.join(a.root, "data", "sharepoint", "Harbourline-Operations", "Procedures")
    drops = os.path.join(a.root, "data", "sharepoint", "lab-09-drops")
    os.makedirs(proc, exist_ok=True)
    os.makedirs(drops, exist_ok=True)
    outage_response(os.path.join(proc, "Outage-Response-Procedure.docx"))
    storm_playbook(os.path.join(proc, "Storm-Restoration-Playbook.pdf"))
    loto_manual(os.path.join(proc, "Lockout-Tagout-Safety-Manual.pdf"))
    substation_switching(os.path.join(proc, "Substation-Switching-Procedure.docx"))
    vegetation(os.path.join(proc, "Vegetation-Management.docx"))
    crew_deck(os.path.join(proc, "Crew-Briefing-Deck.pptx"))
    transformer_specs(os.path.join(proc, "Transformer-Equipment-Specs.xlsx"))
    field_reports(drops)
    print("Wrote", proc)
    print("Wrote", drops)


if __name__ == "__main__":
    main()
