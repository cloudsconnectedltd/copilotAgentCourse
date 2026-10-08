#!/usr/bin/env python3
"""Generate the Archive-Bulk library for the Harbourline-Operations site.

Creates --count small files (default 1,200) spread evenly across 12 year/month
folders (<out>/2025/01 .. <out>/2025/12). File types are a deterministic mix of
.docx, .txt and .md: meter reading notes (MR) and field inspection memos (IM).

Exactly one file carries a planted fact used by Labs 2 and 3 to test retrieval
depth in a large library:
    <out>/2025/09/Inspection-Memo-IM-2025-0873.docx
    "Inspection memo IM-2025-0873: pole P-44817 on Concession Road 9 flagged
     for woodpecker damage."
No other file mentions pole P-44817, Concession Road 9, or woodpecker damage.

The output folder is gitignored and generated at setup time.
Usage: python3 generate_archive_bulk.py [--out <folder>] [--count 1200]
"""

import argparse
import datetime as dt
import os
import random

from docx import Document
from docx.shared import Pt

SEED = 20250101
YEAR = 2025
PLANT_INDEX = 872  # zero-based; becomes IM-2025-0873 in 2025/09 when count is 1200
PLANT_POLE = "P-44817"
PLANT_ROAD = "Concession Road 9"
FIXED_DT = dt.datetime(2026, 1, 5, 8, 0, 0)

REGIONS = [
    ("Eastern Ontario", "EON", ["Sydenham Road", "Bath Road", "Unity Road", "Perth Road", "County Road 2",
                                "Concession Road 4", "Taylor Kidd Boulevard", "Montreal Street"]),
    ("Central Ontario", "CON", ["Welham Road", "Essa Road", "Line 7 North", "Concession Road 11",
                                "Horseshoe Valley Road", "Penetanguishene Road", "Old Barrie Road"]),
    ("Durham", "DUR", ["Thickson Road North", "Simcoe Street North", "Concession Road 6", "Brock Road",
                       "Taunton Road East", "Scugog Line 8", "Winchester Road East"]),
    ("New York North Country", "NYN", ["Arsenal Street", "Coffeen Street", "State Route 12", "County Route 16",
                                       "Outer Washington Street", "Route 3"]),
    ("Western New York", "NYW", ["Harrison Street", "Fairmount Avenue", "Route 60", "Lake Shore Drive",
                                 "Foote Avenue", "Hunt Road"]),
    ("Northeast Ohio", "OHN", ["Lake Avenue", "State Road", "Route 20", "Main Avenue", "Bunker Hill Road",
                               "North Ridge Road"]),
    ("North Central Ohio", "OHC", ["Park Avenue West", "Lexington Avenue", "Route 42", "Ashland Road",
                                   "Trimble Road", "Springmill Street"]),
]
INSPECTORS = ["K. Mahabir", "L. Fontaine", "J. Parrish", "T. Kowalczyk", "M. Delaney", "R. Szabo", "D. Oyelaran",
              "S. McAllister", "B. Nguyen", "H. Leclerc", "P. Galloway", "E. Whitcombe"]
FINDINGS = [
    ("crossarm split at the insulator pin", "Replace crossarm within 90 days"),
    ("ground rod wire cut at the base of the pole", "Reattach ground wire on next visit"),
    ("guy wire slack, anchor rod exposed", "Retension guy and backfill anchor"),
    ("pole top rot noted on sounding", "Schedule pole replacement in next capital plan"),
    ("insulator chipped on the centre phase", "Replace insulator within 60 days"),
    ("transformer tank rust on lower seam, no leak", "Monitor at next inspection cycle"),
    ("vegetation within 1 m of the primary conductor", "Refer to Forestry for hot spot trim"),
    ("missing pole number tag", "Install new pole tag"),
    ("cutout door cracked", "Replace cutout within 60 days"),
    ("lightning arrester disconnector operated", "Replace arrester"),
    ("pole leaning about 5 degrees toward the road", "Straighten and add guy"),
    ("no defects found", "No action required"),
    ("no defects found", "No action required"),
    ("no defects found", "No action required"),
]
METER_ISSUES = [
    "read completed, no issues",
    "read completed, no issues",
    "read completed, no issues",
    "dog on premises, read taken from lane",
    "meter glass fogged, read estimated from interval data",
    "meter seal missing, referred to revenue protection",
    "access blocked by snowbank, card left",
    "meter base loose on wall, referred to metering technician",
    "register reading lower than previous read, flagged for review",
    "customer requested a callback about a high bill",
]


def pole_id(rnd):
    while True:
        p = f"P-{rnd.randint(10000, 69999)}"
        if p != PLANT_POLE:
            return p


def meter_note(rnd, seq, date, region):
    name, code, roads = region
    n = rnd.randint(4, 7)
    lines = []
    for _ in range(n):
        mtr = f"M{rnd.randint(1000000, 9999999)}"
        addr = f"{rnd.randint(10, 4999)} {rnd.choice(roads)}"
        kwh = rnd.randint(250, 2400)
        lines.append((mtr, addr, kwh, rnd.choice(METER_ISSUES)))
    title = f"Meter Reading Notes MR-{YEAR}-{seq:04d}"
    head = [f"Date: {date.isoformat()}", f"Region: {name}", f"Route: {code}-R{rnd.randint(1, 60):02d}",
            f"Reader: {rnd.choice(INSPECTORS)}"]
    return title, head, lines


def inspection_memo(rnd, seq, date, region, planted=False):
    name, code, roads = region
    title = f"Inspection Memo IM-{YEAR}-{seq:04d}"
    feeder = f"{code}-F{rnd.randint(1, 48):02d}"
    head = [f"Date: {date.isoformat()}", f"Region: {name}", f"Feeder: {feeder}",
            f"Inspector: {rnd.choice(INSPECTORS)}"]
    items = []
    for _ in range(rnd.randint(2, 5)):
        f, a = rnd.choice(FINDINGS)
        items.append((pole_id(rnd), f"{rnd.randint(10, 9999)} {rnd.choice(roads)}", f, a))
    summary = ("Routine distribution pole inspection under the 10 year visual and sounding program. "
               f"{len(items)} poles inspected on this memo.")
    if planted:
        items.insert(1, (PLANT_POLE, f"1180 {PLANT_ROAD}", "woodpecker damage: three cavities in the upper third "
                         "of the pole, largest about 12 cm deep, shell thickness below minimum",
                         "Flagged for replacement within 30 days; install woodpecker wrap on adjacent poles"))
        summary = (f"Inspection memo IM-{YEAR}-{seq:04d}: pole {PLANT_POLE} on {PLANT_ROAD} flagged for woodpecker "
                   "damage. Routine distribution pole inspection under the 10 year visual and sounding program. "
                   f"{len(items)} poles inspected on this memo.")
    return title, head, items, summary


def write_txt(path, title, head, body_lines):
    with open(path, "w", encoding="utf-8", newline="\n") as fh:
        fh.write(f"Harbourline Energy Co.\n{title}\n\n")
        fh.write("\n".join(head) + "\n\n")
        fh.write("\n".join(body_lines) + "\n")


def write_md(path, title, head, table_head, rows, summary=None):
    with open(path, "w", encoding="utf-8", newline="\n") as fh:
        fh.write(f"# {title}\n\nHarbourline Energy Co.\n\n")
        for h in head:
            fh.write(f"- {h}\n")
        fh.write("\n")
        if summary:
            fh.write(summary + "\n\n")
        fh.write("| " + " | ".join(table_head) + " |\n")
        fh.write("|" + "---|" * len(table_head) + "\n")
        for r in rows:
            fh.write("| " + " | ".join(str(v) for v in r) + " |\n")


def write_docx(path, title, head, table_head, rows, summary=None):
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)
    cp = doc.core_properties
    cp.author = "Harbourline Energy Co."
    cp.title = title
    cp.created = FIXED_DT
    cp.modified = FIXED_DT
    doc.add_paragraph("Harbourline Energy Co.")
    doc.add_heading(title, 1)
    for h in head:
        doc.add_paragraph(h)
    if summary:
        doc.add_paragraph(summary)
    t = doc.add_table(rows=1, cols=len(table_head))
    t.style = "Table Grid"
    for i, h in enumerate(table_head):
        t.rows[0].cells[i].text = h
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = str(v)
    doc.save(path)


def main():
    here = os.path.dirname(os.path.abspath(__file__))
    default_out = os.path.abspath(os.path.join(here, "..", "..", "data", "sharepoint", "Harbourline-Operations",
                                               "Archive-Bulk"))
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--out", default=default_out)
    ap.add_argument("--count", type=int, default=1200)
    a = ap.parse_args()
    rnd = random.Random(SEED)
    plant = PLANT_INDEX if a.count > PLANT_INDEX else int(a.count * 0.73)
    planted_path = None
    for i in range(a.count):
        month = i * 12 // a.count + 1
        folder = os.path.join(a.out, str(YEAR), f"{month:02d}")
        os.makedirs(folder, exist_ok=True)
        day = rnd.randint(1, 28)
        date = dt.date(YEAR, month, day)
        region = rnd.choice(REGIONS)
        seq = i + 1
        roll = rnd.random()
        fmt_roll = rnd.random()
        is_plant = i == plant
        if is_plant:
            kind, ext = "IM", "docx"
        else:
            kind = "MR" if roll < 0.45 else "IM"
            ext = "docx" if fmt_roll < 0.6 else ("txt" if fmt_roll < 0.85 else "md")
        if kind == "MR":
            title, head, lines = meter_note(rnd, seq, date, region)
            fname = f"Meter-Reading-Notes-MR-{YEAR}-{seq:04d}.{ext}"
            th = ["Meter", "Service address", "kWh since last read", "Note"]
            rows = lines
            summary = None
        else:
            title, head, rows, summary = inspection_memo(rnd, seq, date, region, planted=is_plant)
            fname = f"Inspection-Memo-IM-{YEAR}-{seq:04d}.{ext}"
            th = ["Pole", "Location", "Finding", "Action"]
        path = os.path.join(folder, fname)
        if ext == "docx":
            write_docx(path, title, head, th, rows, summary)
        elif ext == "md":
            write_md(path, title, head, th, rows, summary)
        else:
            body = ([summary, ""] if summary else []) + ["; ".join(str(v) for v in r) for r in rows]
            write_txt(path, title, head, body)
        if is_plant:
            planted_path = path
    print(f"Wrote {a.count} files under {a.out}")
    print(f"Planted fact file: {planted_path}")


if __name__ == "__main__":
    main()
