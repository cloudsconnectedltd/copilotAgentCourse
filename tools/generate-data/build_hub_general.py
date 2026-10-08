#!/usr/bin/env python3
"""Generate Harbourline-Hub general content: Getting-Started, Finance, Vendors list.

Outputs (relative to repo root):
  data/sharepoint/Harbourline-Hub/Getting-Started/*.docx   (5 clean docs, Lab 1)
  data/sharepoint/getting-started.zip                      (same 5 docs, flat)
  data/sharepoint/Harbourline-Hub/Finance/*                (expense, approval matrix docx+xlsx, card pdf, capital)
  data/sharepoint/Harbourline-Hub/Vendors.csv              (2,600 rows, planted rows for evals)

Deterministic: fixed random seed, fixed document metadata dates.
Run: python3 tools/generate-data/build_hub_general.py
"""
import csv
import datetime as dt
import os
import random
import zipfile

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
HUB = os.path.join(ROOT, "data", "sharepoint", "Harbourline-Hub")
GS = os.path.join(HUB, "Getting-Started")
FIN = os.path.join(HUB, "Finance")
FIXED_DATE = dt.datetime(2025, 9, 15, 9, 0, 0)
SEED = 20250915

BANNED = (chr(0x2013), chr(0x2014))


def check_text(s):
    for b in BANNED:
        if b in s:
            raise ValueError("Banned dash character in text: %r" % s[:80])
    return s


# ---------------------------------------------------------------------------
# docx helpers
# ---------------------------------------------------------------------------
class Doc:
    def __init__(self, title, subtitle, owner, effective, version, doc_id):
        self.d = Document()
        st = self.d.styles["Normal"]
        st.font.name = "Calibri"
        st.font.size = Pt(11)
        cp = self.d.core_properties
        cp.title = check_text(title)
        cp.author = check_text(owner)
        cp.last_modified_by = check_text(owner)
        cp.created = FIXED_DATE
        cp.modified = FIXED_DATE
        cp.revision = 1
        cp.subject = check_text(subtitle)
        self.d.add_heading(check_text(title), level=0)
        p = self.d.add_paragraph()
        r = p.add_run(check_text(subtitle))
        r.italic = True
        r.font.color.rgb = RGBColor(0x40, 0x40, 0x40)
        self.table(
            ["Field", "Value"],
            [
                ["Document ID", doc_id],
                ["Owner", owner],
                ["Effective date", effective],
                ["Version", version],
                ["Applies to", "All Harbourline Energy Co. employees unless stated otherwise"],
            ],
        )

    def h(self, text, level=1):
        self.d.add_heading(check_text(text), level=level)

    def p(self, text, bold_lead=None):
        para = self.d.add_paragraph()
        if bold_lead:
            r = para.add_run(check_text(bold_lead) + " ")
            r.bold = True
        para.add_run(check_text(text))
        return para

    def bullets(self, items, numbered=False):
        style = "List Number" if numbered else "List Bullet"
        for it in items:
            self.d.add_paragraph(check_text(it), style=style)

    def table(self, header, rows, widths=None):
        t = self.d.add_table(rows=1, cols=len(header))
        t.style = "Light Grid Accent 1"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(header):
            cell = t.rows[0].cells[i]
            cell.text = ""
            run = cell.paragraphs[0].add_run(check_text(str(h)))
            run.bold = True
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = check_text(str(v))
        self.d.add_paragraph()
        return t

    def revisions(self, rows):
        self.h("Revision history", 1)
        self.table(["Version", "Date", "Author", "Change"], rows)

    def save(self, path):
        os.makedirs(os.path.dirname(path), exist_ok=True)
        self.d.save(path)
        return path


# ---------------------------------------------------------------------------
# Getting-Started (Lab 1, clean, one home per fact)
# ---------------------------------------------------------------------------
def gs_welcome():
    d = Doc(
        "Welcome to Harbourline",
        "Your first-week guide to Harbourline Energy Co.",
        "People and Culture, Onboarding Team",
        "2025-09-01",
        "2.0",
        "HLE-GS-001",
    )
    d.h("1. About Harbourline Energy Co.")
    d.p(
        "Harbourline Energy Co. is a regulated electric utility founded in 1924. We deliver "
        "electricity to about 1.2 million customers across Ontario, New York and Ohio, and we "
        "employ about 2,000 people."
    )
    d.p(
        "Our Ontario operations are regulated by the Ontario Energy Board. Our US operations are "
        "regulated by the New York and Ohio public service commissions, and our transmission "
        "assets follow NERC and FERC reliability requirements."
    )
    d.h("2. Our mission and values")
    d.p("Our mission is to keep the lights on safely, reliably and affordably for every community we serve.")
    d.p("Harbourline has four core values, known internally as SAFE:")
    d.bullets(
        [
            "Safety first: no job is so urgent that it cannot be done safely.",
            "Accountability: we own our commitments to customers and to each other.",
            "Fairness: we treat customers, colleagues and suppliers with respect.",
            "Excellence: we measure our work and keep improving it.",
        ]
    )
    d.h("3. Leadership team")
    d.table(
        ["Name", "Role"],
        [
            ["Eleanor Voss", "President and Chief Executive Officer"],
            ["Rajiv Menon", "Chief Financial Officer"],
            ["Dana Okafor", "Chief Operating Officer"],
            ["Colette Brisebois", "Chief People Officer"],
            ["Martin Szabo", "Chief Information Officer"],
        ],
    )
    d.h("4. Your first week")
    d.p("Complete these steps during your first five working days:")
    d.bullets(
        [
            "Day 1: collect your badge and laptop, and meet your onboarding buddy.",
            "Day 1: sign in to your laptop and set up multi-factor authentication (see the IT Help Desk FAQ).",
            "Day 2: complete the mandatory Safety Orientation course in the learning portal.",
            "Day 3: review your benefits options (see Benefits at a Glance).",
            "Day 4: read the Vacation Policy Quick Guide and confirm your time-off balance in Workday.",
            "Day 5: attend the New Employee Welcome Session with the leadership team.",
        ],
        numbered=True,
    )
    d.p(
        "The New Employee Welcome Session runs on the first Friday of every month from 10:00 a.m. "
        "to 12:00 p.m. Eastern Time, in person at head office and streamed on Microsoft Teams."
    )
    d.h("5. Your onboarding buddy")
    d.p(
        "Every new employee is paired with an onboarding buddy for their first 90 days. Your buddy "
        "is a colleague from a different team who can answer everyday questions. Your manager "
        "introduces you to your buddy on your first day."
    )
    d.h("6. Where to find things")
    d.table(
        ["Topic", "Where to look"],
        [
            ["Time off", "Vacation Policy: Quick Guide for New Employees"],
            ["Laptops, passwords, Wi-Fi and IT support", "IT Help Desk FAQ"],
            ["Office addresses and building hours", "Office Locations"],
            ["Health, dental, retirement and wellness", "Benefits at a Glance"],
        ],
    )
    d.revisions([["1.0", "2023-02-01", "Onboarding Team", "First edition."],
                 ["2.0", "2025-09-01", "Onboarding Team", "Updated leadership team and first-week steps."]])
    return d.save(os.path.join(GS, "Welcome-to-Harbourline.docx"))


def gs_vacation():
    d = Doc(
        "Vacation Policy: Quick Guide for New Employees",
        "A short guide to earning, booking and carrying over vacation days",
        "People and Culture, Total Rewards",
        "2025-01-01",
        "4.0",
        "HLE-GS-002",
    )
    d.h("1. Purpose")
    d.p(
        "This quick guide explains how vacation works for new full-time employees at Harbourline "
        "Energy Co. It summarizes the current (version 4) vacation entitlements."
    )
    d.h("2. How many vacation days you receive")
    d.p("Full-time employees receive 15 vacation days in their first year.", bold_lead="Key fact:")
    d.p("Your annual vacation entitlement is based on your years of continuous service:")
    d.table(
        ["Years of service", "Vacation days per year"],
        [
            ["Years 1 to 4", "15 days"],
            ["Years 5 to 9", "20 days"],
            ["Year 10 and beyond", "25 days"],
        ],
    )
    d.p(
        "Your entitlement increases on the anniversary of your start date. Part-time employees "
        "receive vacation in proportion to their scheduled hours."
    )
    d.h("3. Booking vacation")
    d.bullets(
        [
            "Request vacation through Workday at least 2 weeks before your first day off.",
            "Your manager approves or declines the request in Workday.",
            "Check your team calendar before requesting time off during peak storm season.",
        ]
    )
    d.h("4. Carrying over unused days")
    d.p(
        "You may carry over up to 5 unused vacation days into the next calendar year. Carried-over "
        "days must be used by the end of the first quarter (March 31) of the next year. Any "
        "carried-over days not used by March 31 are forfeited."
    )
    d.h("5. Questions")
    d.p("Contact your manager or the Total Rewards team through the People and Culture service portal in Workday.")
    d.revisions([["4.0", "2025-01-01", "Total Rewards", "Quick guide aligned with version 4 entitlements."]])
    return d.save(os.path.join(GS, "Vacation-Policy.docx"))


def gs_it_faq():
    d = Doc(
        "IT Help Desk FAQ",
        "Answers to the most common technology questions from new employees",
        "Information Technology, Service Desk",
        "2025-06-01",
        "3.1",
        "HLE-GS-003",
    )
    d.h("1. How do I contact the IT Help Desk?")
    d.table(
        ["Channel", "Details"],
        [
            ["Phone", "1-800-555-0142"],
            ["Self-service portal", "https://helpdesk.harbourline.example"],
            ["Email", "helpdesk@harbourline.example"],
            ["Hours", "Monday to Friday, 7:00 a.m. to 7:00 p.m. Eastern Time"],
            ["After-hours support", "24 hours a day, 7 days a week for outage-critical systems (press 1 when calling)"],
        ],
    )
    d.h("2. How quickly will my ticket be answered?")
    d.table(
        ["Priority", "Example", "First response target"],
        [
            ["P1 Critical", "Control room or outage management system unavailable", "15 minutes"],
            ["P2 High", "A whole team cannot work", "1 hour"],
            ["P3 Normal", "One person cannot use an application", "4 business hours"],
            ["P4 Request", "New software or accessory request", "2 business days"],
        ],
    )
    d.h("3. How do I reset my password?")
    d.p(
        "Use the Reset Password link on the self-service portal. Passwords must be at least 14 "
        "characters long and expire every 90 days. You cannot reuse any of your last 10 passwords."
    )
    d.h("4. How do I set up multi-factor authentication?")
    d.p(
        "Install the Microsoft Authenticator app on your phone and follow the prompts the first "
        "time you sign in. Multi-factor authentication is required for all Harbourline accounts."
    )
    d.h("5. How do I connect from home?")
    d.p(
        "Use the HarbourLink VPN client that is preinstalled on your laptop. Microsoft 365 "
        "applications such as Outlook and Teams work without the VPN."
    )
    d.h("6. What equipment will I receive?")
    d.p(
        "Office employees receive a standard laptop, a docking station, two monitors and a headset. "
        "Field employees receive a rugged tablet. Equipment requests are submitted as P4 requests "
        "through the self-service portal."
    )
    d.h("7. How do I connect to office Wi-Fi?")
    d.p(
        "Company laptops connect automatically to the HLE-Corp network. Visitors and personal "
        "devices use the HLE-Guest network, which requires a daily access code from reception."
    )
    d.h("8. How do I report a suspicious email?")
    d.p(
        "Select the Report Phishing button in Outlook. Do not forward the suspicious email to "
        "colleagues. The Cyber Security team reviews every report."
    )
    d.revisions([["3.0", "2024-11-01", "Service Desk", "Added priority targets."],
                 ["3.1", "2025-06-01", "Service Desk", "Updated password length to 14 characters."]])
    return d.save(os.path.join(GS, "IT-Help-Desk-FAQ.docx"))


def gs_offices():
    d = Doc(
        "Office Locations",
        "Addresses and building information for Harbourline offices",
        "Corporate Real Estate and Facilities",
        "2025-07-01",
        "2.2",
        "HLE-GS-004",
    )
    d.h("1. Head office")
    d.p("Harbourline Energy Co. head office is located at 1450 Harbourfront Drive, Toronto, ON M5J 2X9, Canada.")
    d.p("Reception is open Monday to Friday, 8:00 a.m. to 6:00 p.m. Eastern Time. The main reception phone number is 416-555-0100.")
    d.h("2. All office locations")
    d.table(
        ["Office", "Type", "Address"],
        [
            ["Toronto", "Head office", "1450 Harbourfront Drive, Toronto, ON M5J 2X9"],
            ["Hamilton", "Ontario System Control Centre", "88 Escarpment Road, Hamilton, ON L8N 3T4"],
            ["Sudbury", "Northern Ontario Service Centre", "310 Nickel Ridge Way, Sudbury, ON P3A 4R7"],
            ["Albany", "New York Regional Office", "25 State Pier Street, Albany, NY 12207"],
            ["Syracuse", "New York Operations Centre", "740 Onondaga Parkway, Syracuse, NY 13202"],
            ["Columbus", "Ohio Regional Office", "1200 Scioto Commons Blvd, Columbus, OH 43215"],
            ["Akron", "Ohio Operations Centre", "455 Cuyahoga Falls Road, Akron, OH 44308"],
        ],
    )
    d.h("3. Building access")
    d.bullets(
        [
            "Your photo badge opens every Harbourline office.",
            "Office buildings are open to badge holders from 6:00 a.m. to 10:00 p.m., seven days a week.",
            "The Hamilton Ontario System Control Centre and the Syracuse New York Operations Centre are staffed 24 hours a day and require an additional control room access approval.",
            "Visitors must sign in at reception and be escorted at all times.",
        ]
    )
    d.h("4. Parking")
    d.p(
        "Head office has underground parking for employees with a parking permit, which is "
        "requested from Facilities. All other offices have free on-site parking."
    )
    d.h("5. Booking a desk or meeting room")
    d.p("Desks and meeting rooms at every office are booked through the Harbourline Spaces app in Microsoft Teams.")
    d.revisions([["2.2", "2025-07-01", "Facilities", "Added Akron Ohio Operations Centre."]])
    return d.save(os.path.join(GS, "Office-Locations.docx"))


def gs_benefits():
    d = Doc(
        "Benefits at a Glance",
        "A summary of the benefits available to new full-time employees",
        "People and Culture, Total Rewards",
        "2025-01-01",
        "5.0",
        "HLE-GS-005",
    )
    d.h("1. When your benefits start")
    d.p("Benefits coverage starts on your first day of employment. You have 31 days from your start date to enrol.")
    d.p("After your first year, you can change your elections during annual open enrolment, which runs from November 1 to November 15 each year.")
    d.h("2. Benefits providers")
    d.table(
        ["Employees in", "Health and dental provider", "Enrolment website"],
        [
            ["Ontario", "Maplecrest Benefits", "https://maplecrest.example/harbourline"],
            ["New York and Ohio", "Keystone Health Partners", "https://keystonehealth.example/harbourline"],
        ],
    )
    d.h("3. Benefits summary")
    d.table(
        ["Benefit", "Ontario employees", "New York and Ohio employees"],
        [
            ["Health and prescription drugs", "Covered at 90 percent", "Covered at 85 percent after deductible"],
            ["Dental", "Covered at 80 percent, basic and major", "Covered at 80 percent, basic and major"],
            ["Vision", "CAD 400 every 24 months", "USD 300 every 24 months"],
            ["Health spending account", "CAD 750 per year", "Not applicable"],
            ["Retirement savings", "Group RRSP, company matches up to 5 percent of salary", "401(k), company matches up to 6 percent of salary"],
            ["Wellness allowance", "CAD 500 per year", "USD 400 per year"],
            ["Life insurance", "2 times annual salary", "2 times annual salary"],
        ],
    )
    d.h("4. Employee Assistance Program")
    d.p(
        "The Employee Assistance Program is provided by ClearPath EAP. It offers free, "
        "confidential counselling for you and your family, 24 hours a day, at 1-800-555-0199."
    )
    d.h("5. How to enrol")
    d.bullets(
        [
            "Open Workday and select the Benefits worker task.",
            "Choose your coverage level and add any dependants.",
            "Submit your elections within 31 days of your start date. If you do not enrol, you receive single coverage only.",
        ],
        numbered=True,
    )
    d.revisions([["5.0", "2025-01-01", "Total Rewards", "Updated wellness allowance and vision amounts."]])
    return d.save(os.path.join(GS, "Benefits-At-a-Glance.docx"))


# ---------------------------------------------------------------------------
# Finance
# ---------------------------------------------------------------------------
LEVELS = ["Supervisor", "Manager", "Director", "VP", "CFO", "Board"]
# Maximum amount (CAD for Ontario, USD for US operations) each level can approve.
MATRIX = [
    ("Operating expense", ["5,000", "25,000", "100,000", "500,000", "2,000,000", "Above 2,000,000"]),
    ("Capital project", ["Not authorized", "50,000", "250,000", "1,000,000", "5,000,000", "Above 5,000,000"]),
    ("Vendor contract", ["10,000", "50,000", "250,000", "1,000,000", "3,000,000", "Above 3,000,000"]),
    ("Emergency restoration procurement", ["25,000", "100,000", "500,000", "2,000,000", "10,000,000", "Above 10,000,000"]),
    ("Consulting services", ["Not authorized", "25,000", "75,000", "300,000", "1,000,000", "Above 1,000,000"]),
    ("IT software", ["2,500", "15,000", "75,000", "250,000", "1,000,000", "Above 1,000,000"]),
]
XLSX_CONSULTING_DIRECTOR = "150,000"  # deliberate difference: docx says 75,000


def fin_expense():
    d = Doc(
        "Expense Policy",
        "Reimbursement of business travel and other business expenses",
        "Finance, Office of the Controller",
        "2025-04-01",
        "6.0",
        "HLE-FIN-101",
    )
    d.h("1. Purpose and scope")
    d.p(
        "This policy sets out how Harbourline Energy Co. reimburses employees for reasonable "
        "business expenses. It applies to all employees in Ontario, New York and Ohio. Expenses "
        "must be necessary for Harbourline business, reasonable in amount and supported by "
        "documentation."
    )
    d.p(
        "Ontario employees are reimbursed in Canadian dollars (CAD). New York and Ohio employees "
        "are reimbursed in US dollars (USD). Costs of service that are recovered through regulated "
        "rates must also meet the prudence standard expected by our regulators."
    )
    d.h("2. Roles and responsibilities")
    d.table(
        ["Role", "Responsibility"],
        [
            ["Employee", "Incurs only necessary expenses, keeps receipts and submits claims on time."],
            ["Approving manager", "Reviews every claim for business purpose and policy compliance within 5 business days."],
            ["Accounts Payable", "Audits claims, processes reimbursement and reports exceptions monthly."],
            ["Controller", "Owns this policy and approves exceptions below VP level."],
        ],
    )
    d.h("3. Meals and per diem rates")
    d.p(
        "When travelling overnight on company business, employees may claim a daily meal per "
        "diem instead of actual meal costs. Receipts are not required for per diem claims. "
        "Per diem is not paid for meals provided by a hotel, conference or client."
    )
    d.table(
        ["Meal", "Ontario (CAD)", "New York and Ohio (USD)"],
        [
            ["Breakfast", "20.00", "18.00"],
            ["Lunch", "25.00", "22.00"],
            ["Dinner", "45.00", "40.00"],
            ["Full day total", "90.00", "80.00"],
            ["Incidentals per day", "10.00", "8.00"],
        ],
    )
    d.p(
        "Employees travelling on the first and last day of a trip receive 75 percent of the full "
        "day total. Alcohol is not reimbursable, except at client events pre-approved by a Director."
    )
    d.h("4. Accommodation")
    d.table(
        ["Location", "Nightly maximum before taxes"],
        [
            ["Toronto", "CAD 275"],
            ["Other Ontario locations", "CAD 200"],
            ["New York City", "USD 350"],
            ["Other US locations", "USD 225"],
        ],
    )
    d.p("Rooms above the nightly maximum require written approval from a Director before booking.")
    d.h("5. Transportation")
    d.h("5.1 Air and rail", 2)
    d.p(
        "Book economy class through the Harbourline travel desk at least 14 days before travel "
        "where possible. Business class is permitted only for flights longer than 6 hours and "
        "requires VP approval."
    )
    d.h("5.2 Personal vehicle mileage", 2)
    d.table(
        ["Jurisdiction", "Mileage rate"],
        [
            ["Ontario", "CAD 0.72 per kilometre"],
            ["New York and Ohio", "USD 0.70 per mile"],
        ],
    )
    d.p(
        "Mileage is paid for business distance only. Deduct your normal commute distance. Mileage "
        "claims must include the date, start and end locations, purpose and distance."
    )
    d.h("5.3 Fleet vehicles", 2)
    d.p("When a Harbourline fleet vehicle is available, it must be used instead of a personal vehicle for trips longer than 200 kilometres (125 miles).")
    d.h("6. Receipts")
    d.bullets(
        [
            "An original or photographed receipt is required for every expense of CAD 25 or more (Ontario) or USD 25 or more (New York and Ohio).",
            "An itemized receipt is required for any single meal of CAD 75 or more or USD 75 or more, together with the names of all attendees.",
            "Credit card slips alone are not accepted as receipts.",
            "A missing receipt declaration may be used once per quarter for an expense below CAD 100 or USD 100.",
        ]
    )
    d.h("7. Submission deadlines")
    d.table(
        ["Rule", "Deadline"],
        [
            ["Submit expense claims in Workday", "Within 30 days of the expense date"],
            ["Late claims", "Claims 31 to 90 days old require Director approval"],
            ["Very late claims", "Claims older than 90 days are not reimbursed unless approved by a VP"],
            ["Fiscal year end", "All expenses dated in December must be submitted by January 10"],
            ["Reimbursement", "Paid by direct deposit within 10 business days of final approval"],
        ],
    )
    d.h("8. Non-reimbursable expenses")
    d.bullets(
        [
            "Personal entertainment, including in-room movies and minibar charges.",
            "Traffic and parking fines.",
            "Upgrades to air, rail, hotel or car rental without approval.",
            "Gifts to public officials or regulator staff.",
            "Membership fees for personal clubs.",
        ]
    )
    d.h("9. Approval of expense claims")
    d.p(
        "Expense claims are approved by the employee's direct manager. Claims submitted by a VP "
        "are approved by the CFO, and claims submitted by the CFO are approved by the CEO. No "
        "one may approve their own expenses. Corporate card spending follows the Corporate Card "
        "Guidelines, and purchasing authority follows the Approval Matrix."
    )
    d.h("10. Policy exceptions")
    d.p(
        "Exceptions must be requested in writing before the expense is incurred. Violations of "
        "this policy may result in the claim being denied, repayment of amounts already paid, "
        "and disciplinary action."
    )
    d.revisions(
        [
            ["5.0", "2024-01-15", "Office of the Controller", "Added US per diem rates."],
            ["5.1", "2024-09-01", "Office of the Controller", "Updated accommodation maximums."],
            ["6.0", "2025-04-01", "Office of the Controller", "Mileage rates set to CAD 0.72 per km and USD 0.70 per mile; submission deadline set to 30 days."],
        ]
    )
    return d.save(os.path.join(FIN, "Expense-Policy.docx"))


def fin_matrix_docx():
    d = Doc(
        "Approval Matrix",
        "Delegation of financial authority by spend category and approver level",
        "Finance, Office of the CFO",
        "2025-04-01",
        "3.2",
        "HLE-FIN-102",
    )
    d.h("1. Purpose")
    d.p(
        "This approval matrix sets the maximum amount each approver level may authorize for "
        "each spend category. A commitment must be approved by the lowest level whose limit is "
        "equal to or greater than the total commitment value."
    )
    d.h("2. Approval thresholds")
    d.p(
        "Amounts are maximum commitment values per transaction. Ontario amounts are in CAD. For "
        "New York and Ohio operations, the same numeric amounts apply in USD."
    )
    d.table(["Spend category"] + LEVELS, [[cat] + vals for cat, vals in MATRIX])
    d.h("3. Notes")
    d.bullets(
        [
            "Note 1: Commitments must not be split into smaller amounts to avoid a higher approval level.",
            "Note 2: Capital project amounts are total project cost including contingency, as defined in the Capital Project Approval procedure.",
            "Note 3: Vendor contract amounts are the total value over the full contract term, including renewal options.",
            "Note 4: Emergency restoration procurement applies only during a declared emergency. All emergency commitments above 500,000 must be ratified by the CFO within 10 business days after the emergency is closed.",
            "Note 5: Consulting services includes management consulting, legal advisory and regulatory advisory work. Engineering design services are treated as Vendor contract.",
            "Note 6: IT software includes subscriptions and licences. Amounts are the total annual cost.",
            "Note 7: Supervisors cannot approve capital projects or consulting services.",
            "Note 8: Temporary delegations of authority must be recorded by Finance before use and may not exceed the delegator's own limit.",
        ]
    )
    d.h("4. Worked examples")
    d.table(
        ["Scenario", "Required approver"],
        [
            ["Operating expense of CAD 80,000", "Director"],
            ["Vendor contract of USD 1,200,000 over three years", "CFO"],
            ["Capital project of CAD 6,500,000", "Board"],
            ["IT software subscription of CAD 12,000 per year", "Manager"],
        ],
    )
    d.revisions(
        [
            ["3.0", "2024-04-01", "Office of the CFO", "Added Emergency restoration procurement category."],
            ["3.1", "2024-10-01", "Office of the CFO", "Added IT software category."],
            ["3.2", "2025-04-01", "Office of the CFO", "Consulting services Director limit reduced from 150,000 to 75,000."],
        ]
    )
    return d.save(os.path.join(FIN, "Approval-Matrix.docx"))


DELEGATIONS = [
    ("Rajiv Menon", "Chief Financial Officer", "Helena Park", "VP, Financial Planning", "Operating expense", "2,000,000", "2025-07-14", "2025-07-25", "Vacation"),
    ("Dana Okafor", "Chief Operating Officer", "Grant Albrecht", "VP, Ontario Operations", "Emergency restoration procurement", "10,000,000", "2025-08-01", "2025-08-15", "Business travel"),
    ("Grant Albrecht", "VP, Ontario Operations", "Simone Tremblay", "Director, Distribution Operations", "Operating expense", "500,000", "2025-06-02", "2025-06-13", "Vacation"),
    ("Priscilla Adeyemi", "VP, US Operations", "Kevin Rourke", "Director, New York Field Services", "Vendor contract", "1,000,000", "2025-09-02", "2025-09-12", "Medical leave"),
    ("Martin Szabo", "Chief Information Officer", "Aisha Rahman", "Director, IT Infrastructure", "IT software", "1,000,000", "2025-05-19", "2025-05-30", "Conference"),
    ("Helena Park", "VP, Financial Planning", "Owen Gallagher", "Director, Treasury", "Capital project", "1,000,000", "2025-10-06", "2025-10-17", "Vacation"),
    ("Simone Tremblay", "Director, Distribution Operations", "Lucas Ferreira", "Manager, Line Operations", "Operating expense", "100,000", "2025-07-28", "2025-08-08", "Vacation"),
    ("Kevin Rourke", "Director, New York Field Services", "Maya Lindqvist", "Manager, Albany Crews", "Emergency restoration procurement", "500,000", "2025-11-03", "2025-11-14", "Training"),
    ("Colette Brisebois", "Chief People Officer", "Nathan Oduya", "VP, Talent", "Consulting services", "300,000", "2025-08-18", "2025-08-29", "Vacation"),
    ("Aisha Rahman", "Director, IT Infrastructure", "Deepak Sethi", "Manager, Cloud Platforms", "IT software", "75,000", "2025-06-16", "2025-06-27", "Vacation"),
    ("Owen Gallagher", "Director, Treasury", "Julia Mancuso", "Manager, Cash Management", "Operating expense", "100,000", "2025-12-22", "2026-01-02", "Holiday coverage"),
    ("Theresa Nakamura", "VP, Regulatory Affairs", "Benoit Lavigne", "Director, Rate Applications", "Consulting services", "300,000", "2025-09-15", "2025-09-26", "OEB hearing preparation"),
    ("Marcus Oyelaran", "Director, Ohio Field Services", "Carla Jennings", "Manager, Akron Crews", "Vendor contract", "250,000", "2025-10-20", "2025-10-31", "Vacation"),
    ("Eleanor Voss", "President and CEO", "Rajiv Menon", "Chief Financial Officer", "Capital project", "5,000,000", "2025-11-17", "2025-11-21", "Business travel"),
    ("Ingrid Solberg", "Director, Procurement", "Tobias Wren", "Manager, Strategic Sourcing", "Vendor contract", "250,000", "2025-08-04", "2025-08-15", "Parental leave coverage"),
]


def fin_matrix_xlsx():
    wb = Workbook()
    wb.properties.creator = "Office of the CFO"
    wb.properties.title = "Approval Matrix"
    wb.properties.created = FIXED_DATE
    wb.properties.modified = FIXED_DATE
    ws = wb.active
    ws.title = "Approval Matrix"
    hdr_fill = PatternFill("solid", fgColor="1F4E79")
    hdr_font = Font(bold=True, color="FFFFFF")
    ws["A1"] = "Harbourline Energy Co. Approval Matrix (maximum commitment per transaction; CAD for Ontario, USD for US operations)"
    ws["A1"].font = Font(bold=True, size=12)
    ws.append([])
    ws.append(["Spend category"] + LEVELS)
    for c in ws[3]:
        c.fill = hdr_fill
        c.font = hdr_font
        c.alignment = Alignment(horizontal="center", wrap_text=True)

    def num(v):
        return int(v.replace(",", "")) if v[0].isdigit() else v

    for cat, vals in MATRIX:
        vals = list(vals)
        if cat == "Consulting services":
            vals[2] = XLSX_CONSULTING_DIRECTOR
        ws.append([cat] + [num(v) for v in vals])
    for row in ws.iter_rows(min_row=4, max_row=3 + len(MATRIX), min_col=2, max_col=7):
        for c in row:
            if isinstance(c.value, int):
                c.number_format = "#,##0"
    ws.append([])
    ws.append(["Source: Approval Matrix working copy maintained by Finance. See Approval-Matrix.docx notes for rules."])
    ws.column_dimensions["A"].width = 36
    for col in "BCDEFG":
        ws.column_dimensions[col].width = 18

    ws2 = wb.create_sheet("Delegations")
    ws2.append(["Delegation ID", "Delegator", "Delegator title", "Delegate", "Delegate title",
                "Spend category", "Delegated limit", "Start date", "End date", "Reason"])
    for c in ws2[1]:
        c.fill = hdr_fill
        c.font = hdr_font
    for i, row in enumerate(DELEGATIONS, 1):
        r = list(row)
        ws2.append(["DEL-2025-%03d" % i, r[0], r[1], r[2], r[3], r[4], int(r[5].replace(",", "")),
                    dt.date.fromisoformat(r[6]), dt.date.fromisoformat(r[7]), r[8]])
    for row in ws2.iter_rows(min_row=2, min_col=7, max_col=9):
        row[0].number_format = "#,##0"
        row[1].number_format = "yyyy-mm-dd"
        row[2].number_format = "yyyy-mm-dd"
    widths = [14, 20, 34, 20, 36, 32, 16, 12, 12, 26]
    for i, w in enumerate(widths):
        ws2.column_dimensions[chr(65 + i)].width = w
    for wsx in (ws, ws2):
        for row in wsx.iter_rows():
            for c in row:
                if isinstance(c.value, str):
                    check_text(c.value)
    path = os.path.join(FIN, "Approval-Matrix.xlsx")
    wb.save(path)
    return path


def fin_card_pdf():
    path = os.path.join(FIN, "Corporate-Card-Guidelines.pdf")
    doc = SimpleDocTemplate(path, pagesize=letter, title="Corporate Card Guidelines",
                            author="Finance, Office of the Controller", invariant=1,
                            leftMargin=0.9 * inch, rightMargin=0.9 * inch)
    ss = getSampleStyleSheet()
    body = ParagraphStyle("b", parent=ss["BodyText"], fontSize=10.5, leading=14)
    h1 = ss["Heading2"]
    story = []

    def P(t, s=body):
        story.append(Paragraph(check_text(t), s))

    def T(rows, widths):
        rows = [[Paragraph(check_text(str(c)), body) for c in r] for r in rows]
        t = Table(rows, colWidths=widths)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E79")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        story.append(t)
        story.append(Spacer(1, 10))

    P("Corporate Card Guidelines", ss["Title"])
    P("Document ID: HLE-FIN-103 | Owner: Finance, Office of the Controller | Effective date: 2025-03-01 | Version 2.1")
    story.append(Spacer(1, 10))
    P("1. Purpose", h1)
    P("These guidelines explain who may hold a Harbourline Energy Co. corporate card, how the card may be used, "
      "and how card transactions are reconciled. Corporate cards are issued by Northstar Commercial Card and "
      "are the property of Harbourline Energy Co.")
    P("2. Eligibility", h1)
    P("A corporate card may be issued to an employee who travels for business at least four times a year, "
      "or who regularly buys low-value goods and services for their team. The request must be approved by "
      "the employee's Director and submitted to Finance through the Workday service catalogue.")
    P("3. Card types and limits", h1)
    T([["Card type", "Single transaction limit", "Monthly limit", "Typical holder"],
       ["Travel card", "CAD 5,000 / USD 5,000", "CAD 15,000 / USD 15,000", "Employees who travel"],
       ["Purchasing card", "CAD 2,500 / USD 2,500", "CAD 10,000 / USD 10,000", "Administrators and supervisors"],
       ["Storm response card", "CAD 10,000 / USD 10,000", "CAD 50,000 / USD 50,000", "Designated field leaders, active only during a declared emergency"]],
      [1.3 * inch, 1.6 * inch, 1.6 * inch, 2.2 * inch])
    P("Temporary limit increases require Director approval and are valid for a maximum of 30 days.")
    P("4. Permitted use", h1)
    P("The card may be used only for business expenses that comply with the Expense Policy, including travel, "
      "accommodation, registration fees, and low-value supplies that are not available through a Harbourline "
      "purchase order or catalogue.")
    P("5. Prohibited use", h1)
    for t in ["Personal purchases of any kind, even if you intend to repay them.",
              "Cash advances and cash withdrawals.",
              "Gift cards, prepaid cards or cash equivalents.",
              "Splitting a purchase into smaller transactions to stay under a limit.",
              "Purchases from a vendor that holds an active Harbourline contract, which must go through a purchase order.",
              "Payments to individuals for services, which must go through Accounts Payable."]:
        P("&bull; " + t)
    P("6. Reconciliation", h1)
    P("The card statement cycle closes on the 25th day of each month. Cardholders must reconcile every "
      "transaction in Workday within 10 business days after the statement closes. Each transaction needs a "
      "business purpose, a cost centre, and a receipt that meets the Expense Policy receipt rules. The "
      "cardholder's manager approves the reconciliation within 5 business days.")
    P("7. Lost or stolen cards", h1)
    P("Report a lost or stolen card immediately to Northstar Commercial Card at 1-800-555-0177, which is "
      "available 24 hours a day. Notify the Finance card administrator at cards@harbourline.example within "
      "24 hours.")
    P("8. Non-compliance", h1)
    T([["Occurrence", "Consequence"],
       ["First", "Written reminder from Finance, copied to the cardholder's manager"],
       ["Second within 12 months", "Card suspended for 60 days"],
       ["Third within 12 months", "Card cancelled and referral to People and Culture"],
       ["Any fraudulent use", "Card cancelled immediately, repayment, and disciplinary action up to termination"]],
      [2.0 * inch, 4.7 * inch])
    P("Unreconciled transactions older than 60 days result in automatic card suspension.")
    P("9. Leaving Harbourline or changing roles", h1)
    P("Cardholders must reconcile all open transactions and return the card to their manager, who destroys it "
      "and confirms cancellation with Finance, before their last day or role change.")
    P("10. Revision history", h1)
    T([["Version", "Date", "Change"],
       ["2.0", "2024-06-01", "Introduced storm response card."],
       ["2.1", "2025-03-01", "Reconciliation deadline set to 10 business days."]],
      [0.8 * inch, 1.2 * inch, 4.7 * inch])
    doc.build(story)
    return path


def fin_capital():
    d = Doc(
        "Capital Project Approval Procedure",
        "Stage-gate process for approving and governing capital projects",
        "Finance, Capital Planning and Investment",
        "2025-02-01",
        "4.0",
        "HLE-FIN-104",
    )
    d.h("1. Purpose and scope")
    d.p(
        "This procedure governs how capital projects are proposed, justified, approved and closed "
        "at Harbourline Energy Co. It applies to every project that creates or extends the life of "
        "a long-lived asset with a total cost of CAD 50,000 or more (USD 50,000 or more for US "
        "operations). Smaller items are expensed or handled as minor capital by the business unit."
    )
    d.h("2. Governance")
    d.p(
        "The Capital Investment Committee (CIC) reviews capital projects. The CIC is chaired by the "
        "CFO and includes the COO, the VP Regulatory Affairs, the VP Asset Management and the "
        "Director of Capital Planning. The CIC meets monthly on the second Tuesday."
    )
    d.p("Approval authority by project value follows the Capital project row of the Approval Matrix.")
    d.h("3. Stage gates")
    d.table(
        ["Gate", "Name", "Purpose", "Decision by", "Funding released"],
        [
            ["G0", "Idea", "Record the need and confirm strategic fit", "Director", "None"],
            ["G1", "Concept", "Compare options and prepare a Class 4 estimate", "CIC", "Up to 5 percent of estimate for studies"],
            ["G2", "Business case", "Approve preferred option and Class 3 estimate", "CIC, plus approver per Approval Matrix", "Design and long-lead materials, up to 20 percent"],
            ["G3", "Execution approval", "Confirm Class 2 estimate and schedule", "Approver per Approval Matrix", "Full budget"],
            ["G4", "Close-out", "Confirm in-service, final cost and lessons learned", "CIC", "Not applicable"],
        ],
    )
    d.p(
        "A project that exceeds its approved budget by more than 10 percent, or by more than CAD "
        "500,000, whichever is lower, must return to the CIC for re-approval before further spending."
    )
    d.h("4. Business case requirements")
    d.p("Every G2 business case must include:")
    d.bullets(
        [
            "Problem statement, including asset condition data and reliability impact (SAIDI and SAIFI).",
            "At least three options, including a do-nothing or deferral option.",
            "Cost estimate with contingency: 20 percent at Class 4, 15 percent at Class 3 and 10 percent at Class 2.",
            "Net present value using the corporate discount rate of 6.5 percent over the asset's useful life.",
            "Risk assessment scored on the Harbourline 5 by 5 risk matrix.",
            "Rate impact analysis prepared by Regulatory Affairs.",
            "Resourcing plan, including internal crews and external vendors.",
            "Environmental, Indigenous consultation and permitting requirements.",
        ]
    )
    d.h("5. Ontario Energy Board rate-base considerations")
    d.p(
        "Ontario capital spending is recovered from customers only when the Ontario Energy Board "
        "(OEB) accepts it into rate base. Project sponsors must therefore show that the investment "
        "is prudent and that the asset will be used and useful for customers."
    )
    d.bullets(
        [
            "Distribution System Plan alignment: Ontario projects must be included in, or reconciled to, the current five-year Distribution System Plan filed with the OEB.",
            "Investment category: each project is classified as System Access, System Renewal, System Service or General Plant, as required for OEB filings.",
            "Materiality: Ontario projects above CAD 1,000,000 require a project summary suitable for inclusion in a rate application.",
            "Incremental capital: projects outside the approved plan may require an Incremental Capital Module request, which Regulatory Affairs must confirm before G2.",
            "In-service timing: assets are added to rate base only when in service. Construction work in progress does not earn a return until then.",
            "Records: keep all decisions, estimates and variance explanations for at least 7 years to support regulatory review.",
        ]
    )
    d.p(
        "US projects follow the rate-case requirements of the New York or Ohio public service "
        "commission, and transmission projects follow FERC formula-rate requirements. Regulatory "
        "Affairs advises on these requirements at G1."
    )
    d.h("6. Post-implementation review")
    d.p(
        "A post-implementation review is completed 12 months after the in-service date for every "
        "project above CAD 1,000,000. The review compares actual cost, schedule and benefits with "
        "the approved business case and is presented to the CIC."
    )
    d.revisions(
        [
            ["3.0", "2023-05-01", "Capital Planning", "Introduced five-gate process."],
            ["4.0", "2025-02-01", "Capital Planning", "Added OEB rate-base section and re-approval trigger."],
        ]
    )
    return d.save(os.path.join(FIN, "Capital-Project-Approval.docx"))


# ---------------------------------------------------------------------------
# Vendors list
# ---------------------------------------------------------------------------
CATEGORIES = ["Line Construction", "Vegetation Management", "IT Services", "Metering", "Fleet",
              "Engineering Consulting", "Safety Equipment", "Transformers", "Facilities", "Professional Services"]
CAT_WORDS = {
    "Line Construction": ["Line Builders", "Powerline Contractors", "Overhead Services", "Utility Construction", "Pole Line", "Grid Constructors"],
    "Vegetation Management": ["Tree Care", "Arborist Services", "Right-of-Way Clearing", "Forestry Services", "Brush Control", "Vegetation Solutions"],
    "IT Services": ["Digital Systems", "Data Solutions", "Software Group", "Networks", "Cloud Partners", "Technology Services"],
    "Metering": ["Meter Services", "Metering Solutions", "Meter Works", "Smart Meter Group", "Measurement Systems", "AMI Services"],
    "Fleet": ["Fleet Services", "Truck Equipment", "Fleet Upfitters", "Vehicle Leasing", "Bucket Truck Repair", "Motor Works"],
    "Engineering Consulting": ["Engineering", "Engineering Associates", "Power Consultants", "Grid Engineering", "Design Group", "Technical Consultants"],
    "Safety Equipment": ["Safety Supply", "Protective Equipment", "Safety Products", "Arc Flash Gear", "Rescue Systems", "Workwear"],
    "Transformers": ["Transformer Works", "Transformer Services", "Electrical Apparatus", "Power Equipment", "Substation Supply", "Coil Works"],
    "Facilities": ["Facility Services", "Building Maintenance", "Janitorial Services", "HVAC Services", "Property Services", "Groundskeeping"],
    "Professional Services": ["Advisory", "Consulting Partners", "Legal Services", "Audit Associates", "Staffing Solutions", "Communications Group"],
}
PLACES = {
    "Ontario": ["Algonquin", "Muskoka", "Niagara", "Georgian Bay", "Kawartha", "Haliburton", "Ottawa Valley", "Rideau",
                "Temiskaming", "Huron", "Grey Bruce", "Quinte", "Thousand Islands", "Superior North", "Nipissing",
                "Durham", "Halton", "Oxford", "Elgin", "Lambton", "Frontenac", "Lanark", "Simcoe", "Wellington"],
    "New York": ["Adirondack", "Catskill", "Finger Lakes", "Hudson Valley", "Mohawk", "Genesee", "Chautauqua",
                 "Oneida", "Seneca", "Tioga", "Saratoga", "Champlain", "Susquehanna", "Allegany", "Cayuga",
                 "Herkimer", "Otsego", "Schoharie", "Tompkins", "Chemung", "Oswego", "Madison", "Delaware", "Ulster"],
    "Ohio": ["Buckeye", "Maumee", "Scioto", "Muskingum", "Hocking", "Cuyahoga", "Miami Valley", "Western Reserve",
             "Erie Shore", "Licking", "Tuscarawas", "Sandusky", "Mahoning", "Wayne", "Portage", "Ashtabula",
             "Knox", "Delaware County", "Fairfield", "Pickaway", "Ross", "Athens", "Marietta", "Wooster"],
}
PREFIXES = ["", "North", "South", "East", "West", "Central", "Summit", "Keystone", "Pioneer", "Heritage", "Evergreen",
            "Ironwood", "Bluewater", "Redstone", "Silverline", "Clearview", "Granite", "Cedar", "Beacon", "Lakeside"]
SUFFIX = {"Ontario": ["Inc.", "Ltd.", "Corp.", "Limited"], "New York": ["LLC", "Inc.", "Corp.", "Co."],
          "Ohio": ["LLC", "Inc.", "Co.", "Group LLC"]}
FIRST = ["Aaron", "Abigail", "Adrian", "Alicia", "Amir", "Ana", "Andre", "Beatrice", "Bianca", "Brandon", "Caleb",
         "Camille", "Chloe", "Colin", "Dalia", "Damien", "Elena", "Elliot", "Farah", "Felix", "Gabriel", "Gemma",
         "Hannah", "Hugo", "Imani", "Isaac", "Jasmine", "Jonah", "Kara", "Kenji", "Laila", "Liam", "Lucia",
         "Malik", "Maren", "Nadia", "Noah", "Olivia", "Omar", "Paige", "Pavel", "Quinn", "Rosa", "Ruben",
         "Sahana", "Sean", "Tamsin", "Theo", "Uma", "Victor", "Wendy", "Xavier", "Yara", "Zoe"]
LAST = ["Abernathy", "Bajwa", "Castellano", "Dubois", "Eriksen", "Fitzgerald", "Gauthier", "Hollister", "Ibrahim",
        "Jablonski", "Kerrigan", "Lachance", "Macintyre", "Novak", "Osei", "Pelletier", "Quintero", "Rasmussen",
        "Sandoval", "Thibodeau", "Underwood", "Vasquez", "Whitlock", "Yamamoto", "Zielinski", "Ashcroft",
        "Bergeron", "Chowdhury", "Delgado", "Fontaine", "Grewal", "Haddad", "Iverson", "Kowal", "Lindgren",
        "Moreau", "Nakagawa", "Okonkwo", "Petrakis", "Rahimi", "Sokolov", "Tanaka", "Vandermeer", "Weiss"]
STATUSES = ["Active", "Expired", "Suspended", "Pending Review"]
STATUS_W = [70, 15, 5, 10]
RISKS = ["Low", "Medium", "High"]
RISK_W = [55, 35, 10]
VALUE_RANGE = {  # (min, max) contract values, whole units of currency
    "Line Construction": (250_000, 4_500_000), "Vegetation Management": (80_000, 2_400_000),
    "IT Services": (20_000, 1_800_000), "Metering": (50_000, 3_000_000), "Fleet": (30_000, 1_500_000),
    "Engineering Consulting": (25_000, 1_200_000), "Safety Equipment": (5_000, 400_000),
    "Transformers": (150_000, 4_800_000), "Facilities": (10_000, 900_000), "Professional Services": (15_000, 750_000),
}

# Planted rows: data row number (1-based, header excluded) -> record (without VendorId)
PLANTED = {
    12: dict(Title="Lakeshore Arborist Collective Inc.", Category="Vegetation Management", Region="Ontario",
             ContractValue=2_975_000, Status="Suspended", RenewalDate="2026-05-31",
             PrimaryContact="Genevieve Marchetti", RiskRating="High"),
    777: dict(Title="Mohawk Valley Meter Works LLC", Category="Metering", Region="New York",
              ContractValue=4_850_000, Status="Active", RenewalDate="2026-11-30",
              PrimaryContact="Desmond Achterberg", RiskRating="Low"),
    1500: dict(Title="Tri-County Fleet Upfitters Inc.", Category="Fleet", Region="Ohio",
               ContractValue=1_735_500, Status="Pending Review", RenewalDate="2026-02-15",
               PrimaryContact="Rowena Castellanos", RiskRating="Medium"),
    2049: dict(Title="Buckeye Substation Engineering Group LLC", Category="Engineering Consulting", Region="Ohio",
               ContractValue=3_310_000, Status="Active", RenewalDate="2027-01-31",
               PrimaryContact="Ignatius Pemberton", RiskRating="High"),
    2501: dict(Title="Northgate Pole & Crossarm Ltd.", Category="Line Construction", Region="Ontario",
               ContractValue=7_425_000, Status="Active", RenewalDate="2027-03-31",
               PrimaryContact="Renata Kowalczyk", RiskRating="High"),
    2600: dict(Title="Adirondack Transformer Rebuild Co.", Category="Transformers", Region="New York",
               ContractValue=5_960_000, Status="Expired", RenewalDate="2025-06-30",
               PrimaryContact="Philippa Vandersloot", RiskRating="Medium"),
}
N_VENDORS = 2600


def slug(name):
    out = []
    for ch in name.lower():
        if ch.isalnum():
            out.append(ch)
        elif out and out[-1] != "-":
            out.append("-")
    s = "".join(out).strip("-")
    for suf in ("-inc", "-ltd", "-llc", "-corp", "-co", "-limited", "-group-llc"):
        if s.endswith(suf):
            s = s[: -len(suf)]
            break
    return s


def build_vendors():
    rng = random.Random(SEED)
    used = {p["Title"] for p in PLANTED.values()}
    planted_contacts = {p["PrimaryContact"] for p in PLANTED.values()}
    rows = []
    regions = ["Ontario", "New York", "Ohio"]
    for n in range(1, N_VENDORS + 1):
        if n in PLANTED:
            rec = dict(PLANTED[n])
        else:
            region = rng.choices(regions, weights=[50, 30, 20])[0]
            cat = rng.choice(CATEGORIES)
            while True:
                name = " ".join(x for x in [rng.choice(PREFIXES), rng.choice(PLACES[region]),
                                            rng.choice(CAT_WORDS[cat]), rng.choice(SUFFIX[region])] if x)
                if name not in used and "Northgate" not in name:
                    break
            lo, hi = VALUE_RANGE[cat]
            val = int(round(rng.uniform(lo, hi) / 500.0) * 500)
            year = rng.choice([2024, 2025, 2025, 2026, 2026, 2027, 2028])
            date = dt.date(year, 1, 1) + dt.timedelta(days=rng.randint(0, 364))
            while True:
                contact = rng.choice(FIRST) + " " + rng.choice(LAST)
                if contact not in planted_contacts:
                    break
            rec = dict(Title=name, Category=cat, Region=region, ContractValue=val,
                       Status=rng.choices(STATUSES, weights=STATUS_W)[0], RenewalDate=date.isoformat(),
                       PrimaryContact=contact, RiskRating=rng.choices(RISKS, weights=RISK_W)[0])
        used.add(rec["Title"])
        rec["VendorId"] = "V-%05d" % n
        rec["Currency"] = "CAD" if rec["Region"] == "Ontario" else "USD"
        first, last = rec["PrimaryContact"].lower().split(" ", 1)
        rec["ContactEmail"] = "%s.%s@%s.example" % (first, last.replace(" ", ""), slug(rec["Title"]))
        rows.append(rec)
    cols = ["Title", "VendorId", "Category", "Region", "ContractValue", "Currency", "Status",
            "RenewalDate", "PrimaryContact", "ContactEmail", "RiskRating"]
    path = os.path.join(HUB, "Vendors.csv")
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=cols)
        w.writeheader()
        for r in rows:
            for v in r.values():
                check_text(str(v))
            w.writerow({c: r[c] for c in cols})
    assert len({r["Title"] for r in rows}) == N_VENDORS
    return path, rows


def vendor_stats(rows):
    """Print aggregate facts for the answer key (full list vs first 2,048 rows)."""
    from collections import Counter
    def summary(sub):
        return dict(
            n=len(sub),
            status=Counter(r["Status"] for r in sub),
            region=Counter(r["Region"] for r in sub),
            risk=Counter(r["RiskRating"] for r in sub),
            high_active=sum(1 for r in sub if r["RiskRating"] == "High" and r["Status"] == "Active"),
            line_on=sum(1 for r in sub if r["Category"] == "Line Construction" and r["Region"] == "Ontario"),
            maxv=max(sub, key=lambda r: r["ContractValue"]),
        )
    for label, sub in (("ALL", rows), ("FIRST_2048", rows[:2048]), ("AFTER_2048", rows[2048:])):
        s = summary(sub)
        print(label, s["n"], dict(s["status"]), dict(s["region"]), dict(s["risk"]),
              "high_active=%d" % s["high_active"], "line_on=%d" % s["line_on"],
              "max=%s %s %s" % (s["maxv"]["VendorId"], s["maxv"]["Title"], s["maxv"]["ContractValue"]))
    cat_counts = Counter(r["Category"] for r in rows)
    print("CATEGORY", dict(cat_counts))


def build_zip(files):
    path = os.path.join(ROOT, "data", "sharepoint", "getting-started.zip")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for fp in files:
            info = zipfile.ZipInfo(os.path.basename(fp), date_time=(2025, 9, 15, 9, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            with open(fp, "rb") as fh:
                z.writestr(info, fh.read())
    return path


def main():
    os.makedirs(GS, exist_ok=True)
    os.makedirs(FIN, exist_ok=True)
    gs = [gs_welcome(), gs_vacation(), gs_it_faq(), gs_offices(), gs_benefits()]
    out = gs + [build_zip(gs)]
    out += [fin_expense(), fin_matrix_docx(), fin_matrix_xlsx(), fin_card_pdf(), fin_capital()]
    vpath, rows = build_vendors()
    out.append(vpath)
    for p in out:
        print("wrote", os.path.relpath(p, ROOT))
    vendor_stats(rows)


if __name__ == "__main__":
    main()
